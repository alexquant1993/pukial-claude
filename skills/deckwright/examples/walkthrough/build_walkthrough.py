"""The walkthrough deck: two slides through every Phase 3 tool.

Ported from examples/walkthrough/spec.md (route B for both slides); there is
no HTML draft for this deck, the spec's section 5 is the design. Exercises
notes.emit_notes in every mode (notes are opt-in: the default is a deck with
no notes pane and no teleprompter, and the gate runs `--notes both` to prove
the feature), pagenums.renumber on a machine-built slide, and gives
lint_deck.py and compose_qa.py a subject.

Deltas from the spec: none.

Geometry and type: 1280x720 canvas, brands/relay type scale, one pill. It
draws its own two slides rather than calling `style.header()`, so it places no
wordmark and no rule: one auto-shape (the pill), zero pictures.
Shape arithmetic for assert_native.py: ONE auto-shape (the pill's rounded
rectangle), zero pictures, zero freeforms; text boxes are not counted.

Run: PYTHONIOENCODING=utf-8 uv run --offline --no-project --with python-pptx \
       --with pillow python examples/walkthrough/build_walkthrough.py \
       --template out/template.pptx --out out/walk/deck.pptx \
       --teleprompter out/walk/teleprompter.html
"""

import argparse
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from deck_kit.brand import load_brand
from deck_kit.deck import add_slide, open_deck
from deck_kit.merge import refuse_in_place
from deck_kit.metrics import WarnRegister
from deck_kit.notes import NOTES_MODES, emit_notes, read_notes
from deck_kit.pagenums import MarkerStyle, renumber
from deck_kit.primitives import rrect, textbox
from notes_data import NOTES, TITLE


def build(template, out, mode, teleprompter=None):
    brand = load_brand("relay")
    s = brand.style()
    m = brand.metrics()
    warn = WarnRegister()
    prs = open_deck(template, brand)

    cover = add_slide(prs, brand)
    textbox(cover, s.L, 280, s.W, 60, "Walkthrough", s.T_COVER, s.TITLE, brand.sans)
    textbox(cover, s.L, 350, s.W, 40, "One deck through the whole pipeline", s.T_SUB,
            s.MUTED, brand.sans)

    loop = add_slide(prs, brand)
    textbox(loop, s.L, 40, s.W, 24, "01. THE LOOP", s.T_EYEBROW, s.BLUE, brand.sans, upper=True)
    title = "One deck, every stage"
    if not m.fits(title, s.T_TITLE, True, s.W):
        warn.add("OVERFLOW", "title: " + title)
    textbox(loop, s.L, 70, s.W, 40, [(title, True)], s.T_TITLE, s.TITLE, brand.sans)
    body = ["Four stages and one loop.", "Three kinds of check: file, visual, content."]
    for k, line in enumerate(body):
        if not m.fits(line, s.T_BODY, False, s.W):
            warn.add("OVERFLOW", "body: " + line)
        textbox(loop, s.L, 140 + 34 * k, s.W, 30, line, s.T_BODY, s.GREY, brand.sans)
    pill_w = int(m.width("Route B", s.T_PILL, True)) + 24
    rrect(loop, s.L, 230, pill_w, 28, fill=s.PANEL, rad=14)
    textbox(loop, s.L, 230, pill_w, 28, [("Route B", True)], s.T_PILL, s.TITLE, brand.sans,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    textbox(loop, s.L, s.FOOT_Y, 600, 14, s.COPYRIGHT, s.T_META, s.MUTED, brand.sans)
    # s.FOOT_Y, not a literal: renumber() puts the page marker at the brand's
    # PAGE_Y, which is FOOT_Y, and a copyright line on a different baseline
    # from the number it sits opposite reads as a mistake in the export.
    # T_META, not T_FOOT: T_FOOT is in ALL_CAPS_MICRO and licensed only with
    # upper=True; the brand's own header() writes COPYRIGHT at T_META.

    style = MarkerStyle(s.PAGE_X, s.PAGE_Y, s.PAGE_W, s.PAGE_H, s.PAGE_SIZE, s.PAGE_COLOR,
                        brand.sans)
    report = renumber(prs, style, skip=(1,))
    if report != [(2, 2, "added")]:
        raise SystemExit("renumber report %r, expected [(2, 2, 'added')]" % (report,))

    refuse_in_place(out, template)
    # Into the package before the save; the teleprompter file is written here
    # too, or not at all, according to the mode.
    emit_notes(prs, NOTES, mode, font=brand.sans, size_pt=s.T_BODY, gap_pt=6, title=TITLE,
               teleprompter=teleprompter, skip=(1,))
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    warn.exit_if_dirty()
    return out


def check(out, mode, teleprompter=None):
    """Reopen what the mode asked for and assert the rest was NOT written.

    These are the gate's teeth for D5's 'generated once' under --notes both,
    and under --notes none they prove the default really is a deck without
    a notes pane and without a teleprompter file.
    """
    prs = Presentation(str(out))
    if len(prs.slides) != 2:
        raise SystemExit("expected 2 slides, found %d" % len(prs.slides))
    got = read_notes(prs)
    if mode in ("pptx", "both"):
        if got != {1: ["Welcome. This is one deck through the whole pipeline."],
                   2: ["Four stages and one loop: spec, HTML draft, native build, integration.",
                       "Three kinds of check, kept distinct: file, visual, content.",
                       "> If asked: the HTML draft is disposable once the builder exists."]}:
            raise SystemExit("notes did not round-trip: %r" % (got,))
    elif got != {}:
        raise SystemExit("--notes %s must leave the notes pane empty, found %r" % (mode, got))
    if mode in ("teleprompter", "both"):
        html_text = teleprompter.read_text(encoding="utf-8")
        for needle in ('<section id="s1"', '<section id="s2"', "Slide 2 [2]",
                       "Total spoken time: 3 min", 'class="aside"'):
            if needle not in html_text:
                raise SystemExit("teleprompter lacks %r" % needle)
        if "Slide 1 [" in html_text:
            raise SystemExit("the cover must carry no printed number")
    elif teleprompter is not None and teleprompter.exists():
        raise SystemExit("--notes %s must not write a teleprompter, but %s exists (this run, "
                         "or a previous one at the same path)" % (mode, teleprompter))


def parse_args(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", type=Path, required=True)
    ap.add_argument("--out", type=Path, required=True)
    ap.add_argument("--notes", choices=NOTES_MODES, default="none",
                    help="which speaker-notes outputs to emit (default: none)")
    ap.add_argument("--teleprompter", type=Path,
                    help="where to write the teleprompter page; needed with "
                         "--notes teleprompter or both")
    args = ap.parse_args(argv)
    if args.notes in ("teleprompter", "both") and args.teleprompter is None:
        raise SystemExit("--notes %s writes a teleprompter: pass --teleprompter PATH"
                         % args.notes)
    return args


def main():
    args = parse_args()
    out = build(args.template, args.out, args.notes, args.teleprompter)
    check(out, args.notes, args.teleprompter)
    print("walkthrough OK ->", out, "notes:", args.notes,
          *([args.teleprompter] if args.notes in ("teleprompter", "both") else []))


if __name__ == "__main__":
    main()
