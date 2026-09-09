"""The synthetic 'author's deck': what a human hands back after editing.

It carries, deliberately, everything Phase 2 has to survive:
  - a paragraph split into three runs with the middle one bold, which is what
    PowerPoint produces after a few edits and what naive retexting destroys;
  - all three page-number marker flavours across three slides;
  - a slide with no marker at all, as a machine-built slide arrives;
  - an animation targeting one shape, so the deletion guard has something to
    guard.

Run: PYTHONIOENCODING=utf-8 uv run --offline --no-project --with python-pptx \
       --with pillow python examples/integration/build_host.py \
       --template out/template.pptx --out out/host.pptx
"""

import argparse
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))

from lxml import etree
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from deck_kit.brand import load_brand
from deck_kit.deck import add_slide, open_deck
from deck_kit.metrics import WarnRegister
from deck_kit.pagenums import MarkerStyle
from deck_kit.primitives import rrect, textbox

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def _field_marker(slide, brand, style_mod):
    """Flavour one: a text box whose paragraph holds an <a:fld> slidenum.

    Lives here, not in tests/test_pagenums.py, so the round-trip host and the
    unit-test deck build the same shape from one construction rather than two
    that can drift; tests/test_pagenums.py imports this back.
    """
    box = textbox(slide, style_mod.x, style_mod.y, style_mod.w, style_mod.h,
                  "0", style_mod.size, style_mod.color, style_mod.font)
    # Deliberately NOT named "Slide Number...": a name that matches
    # LOCALE_NAMES would be returned by find_marker's name branch before the
    # field branch is ever consulted, and deleting the field handling
    # altogether would still leave this test green.
    box._element.find(".//" + "{%s}cNvPr" % P).set("name", "TextBox 7")
    para = box.text_frame.paragraphs[0]._p
    run = para.find("{%s}r" % A)
    rpr = run.find("{%s}rPr" % A)
    fld = etree.SubElement(para, "{%s}fld" % A)
    fld.set("id", "{7B2A4C1E-0000-4000-8000-0000000000FF}")
    fld.set("type", "slidenum")
    if rpr is not None:
        fld.append(etree.fromstring(etree.tostring(rpr)))
    t = etree.SubElement(fld, "{%s}t" % A)
    t.text = "0"
    para.remove(run)
    return box


def build(template, out):
    brand = load_brand("relay")
    s = brand.style()
    warn = WarnRegister()
    # PAGE_FONT, not brand.sans: the marker is a page number and the brand
    # says which face a page number is set in. examples/integration/version.py
    # builds the identical MarkerStyle, and renumber() has to recognise the
    # boxes this file writes.
    style = MarkerStyle(x=s.PAGE_X, y=s.PAGE_Y, w=s.PAGE_W, h=s.PAGE_H,
                        size=s.PAGE_SIZE, color=s.PAGE_COLOR, font=s.PAGE_FONT)

    prs = open_deck(template, brand)

    # 1: cover - a title box only, no page-number marker.
    slide = add_slide(prs, brand)
    textbox(slide, 80, 80, 900, 40, [("Deckwright Phase 2 round trip", True)],
           24, s.INK, brand.sans)

    # 2: the mixed-bold paragraph a naive retext would flatten, plus a
    #    field-flavour marker (catalogue item 14, flavour one).
    slide = add_slide(prs, brand)
    textbox(slide, 80, 80, 900, 40,
           [("The gap is ", False), ("wide", True), (" in three domains.", False)],
           16, s.INK, brand.sans)
    _field_marker(slide, brand, style)

    # 3: a marker named for the locale, holding a literal "99" (flavour two).
    #    Positioned well above the marker floor, NOT at style.x/style.y: at
    #    the marker's usual position "99" would also satisfy the literal-
    #    digit fallback, so deleting the name branch would still pass.
    slide = add_slide(prs, brand)
    textbox(slide, 80, 80, 900, 40, "A locale-named marker box.", 16, s.INK, brand.sans)
    box = textbox(slide, 900, 50, style.w, style.h, "99",
                 style.size, style.color, style.font)
    box._element.find(".//" + "{%s}cNvPr" % P).set("name", "Numero de diapositiva 3")

    # 4: a plain literal marker text box at the footer position (flavour
    #    three), recognised only by being small, low and numeric.
    slide = add_slide(prs, brand)
    textbox(slide, 80, 80, 900, 40, "A plain literal marker.", 16, s.INK, brand.sans)
    textbox(slide, style.x, style.y, style.w, style.h, "99",
           style.size, style.color, style.font)

    # 5: no marker at all, as a machine-built slide arrives. The round trip's
    #    "delete" step removes this position, so its content is inert.
    slide = add_slide(prs, brand)
    textbox(slide, 80, 80, 900, 40, "To be deleted.", 16, s.INK, brand.sans)

    # 6: a text box on a pill, identical bounding boxes, with a p:timing tree
    #    targeting the text box - so the animation deletion guard (catalogue
    #    item 8) and delete_with_backing have a real pill/text pair to find,
    #    even though this round trip never deletes it.
    slide = add_slide(prs, brand)
    rrect(slide, 400, 300, 300, 60, fill=s.PANEL, rad=30)
    label = textbox(slide, 400, 300, 300, 60, "on a pill", 16, s.INK, brand.sans,
                    align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    timing = etree.SubElement(slide._element, "{%s}timing" % P)
    tgt = etree.SubElement(timing, "{%s}spTgt" % P)
    tgt.set("spid", str(label.shape_id))

    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print("wrote", out, "slides:", len(prs.slides))
    warn.exit_if_dirty()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", required=True, type=Path)
    ap.add_argument("--out", required=True, type=Path)
    args = ap.parse_args()
    build(args.template, args.out)


if __name__ == "__main__":
    main()
