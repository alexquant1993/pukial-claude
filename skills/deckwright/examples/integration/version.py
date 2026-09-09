"""The Phase 2 round trip, wired through the real version driver.

This is the gate: it is the whole of stage 4 run end to end on decks this
repository builds itself, so anyone can re-run it with no client file.

Run: PYTHONIOENCODING=utf-8 uv run --offline --no-project --with python-pptx \
       --with pillow python examples/integration/version.py \
       --host out/host.pptx --donor out/donor.pptx --fonts out/fontdonor.pptx \
       --work out/v2
"""

import argparse
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "scripts"))

from pptx import Presentation

from deck_kit.brand import load_brand
from deck_kit.deck import delete_positions
from deck_kit.fonts import font_parts
from deck_kit.geometry import px
from deck_kit.merge import dangling_refs, read_package
from deck_kit.pagenums import MarkerStyle, find_marker, renumber
from deck_kit.textedit import (iter_paragraphs, iter_shapes, paragraph_text,
                               patch_text)

from new_version import Version  # noqa: E402

# The brand supplies the marker's geometry, size, colour and font - the same
# rule pagenums.MarkerStyle enforces everywhere else - so the round trip
# derives these from brands/relay rather than hardcoding position 7's
# expectations against a second set of numbers. Identical to the MarkerStyle
# examples/integration/build_host.py builds, deliberately: renumber() has to
# recognise the boxes that file wrote.
_brand = load_brand("relay")
_s = _brand.style()
MARKER_STYLE = MarkerStyle(x=_s.PAGE_X, y=_s.PAGE_Y, w=_s.PAGE_W, h=_s.PAGE_H,
                           size=_s.PAGE_SIZE, color=_s.PAGE_COLOR, font=_s.PAGE_FONT)
MIN_TOP = px(_s.PAGE_Y) - 1


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--host", required=True, type=Path)
    ap.add_argument("--donor", required=True, type=Path)
    ap.add_argument("--fonts", required=True, type=Path,
                    help="the font donor for the re-embed stage - see "
                         "docs/decisions.md for why this is a "
                         "deck_kit.fonts.embed donor and not a COM one")
    ap.add_argument("--work", required=True, type=Path)
    args = ap.parse_args()

    v = Version(source=args.host, work=args.work, final=args.work / "final.pptx")
    v.expect(slide_count=6, fingerprint={2: "The gap is", 6: "on a pill"})

    # 1. retext through run attribution - the bold must survive.
    v.step("patch text", lambda prs: patch_text(prs, [
        (2, "The gap is wide in three domains.",
            "The gap is wide in four domains.")]), slides=6)

    # 2. delete the position being replaced.
    v.step("delete", lambda prs: delete_positions(prs, [5]), slides=5)

    # 3. transplant, one hidden, out of process.
    v.subprocess_step("merge", [sys.executable, str(HERE / "_merge_step.py"),
                                "--donor", str(args.donor)], slides=7)

    # 4. renumber - all three flavours, hidden slide skipped. `renumber`
    #    derives its own marker floor from MARKER_STYLE.y; it takes no
    #    separate min_top_emu parameter (see docs/decisions.md and the Task 9
    #    report - the brief's snippet passed one that the real signature does
    #    not accept).
    #
    #    The report is captured, not discarded: `renumber` self-heals when a
    #    marker flavour goes unrecognised. If find_marker ever misses slide
    #    3's locale-named box (flavour two), renumber concludes the slide has
    #    no marker at all and ADDS a brand-new, correctly-positioned "3" -
    #    leaving the stale "99" box behind. The final numbering dict below
    #    would still be exactly right, because the new box carries the right
    #    number; only "fixed" vs "added" tells the two cases apart.
    renumber_report = []
    v.step("renumber", lambda prs: renumber_report.extend(
        renumber(prs, MARKER_STYLE, skip=(1,))), slides=7)

    # 5. re-embed fonts, out of process. See the font decision point in
    #    docs/decisions.md: this donor is written by deck_kit.fonts.embed,
    #    measured on this machine to survive the doubled COM check, not the
    #    COM (embed_fonts.ps1) route, which is unavailable here.
    v.subprocess_step("fonts", [sys.executable, str(HERE / "_fonts_step.py"),
                                "--donor", str(args.fonts), "--expect", "2"],
                      slides=7)
    v.finalise()

    # ------------------------------------------------------- the gate's teeth
    final = Presentation(str(v.final))
    assert len(final.slides) == 7, len(final.slides)          # 6 - 1 + 2

    # the bold survived the retexting
    para = [p for p in iter_paragraphs(final.slides[1])
            if "four domains" in paragraph_text(p)][0]
    bold = [r.text for r in para.runs if r.font.bold]
    assert bold == ["wide"], bold

    # exactly one hidden slide, and it is the one asked for
    hidden = [i for i, sl in enumerate(final.slides, 1)
              if sl._element.get("show") == "0"]
    assert hidden == [6], hidden

    # numbering: hidden slide consumed no number, cover unnumbered. Asserted
    # as the whole dict, not three keys - a three-key check let a reviewer
    # delete LOCALE_NAMES entirely (flavour two, position 3's marker) and
    # stay green while position 3 carried a stale "99" and a duplicate added
    # marker at once. Position 6 is the hidden donor slide: renumber() never
    # visits a hidden slide at all, and build_donor.py gives it no marker to
    # begin with, so it stays None - not skipped-and-numbered, never visited.
    numbers = {i: (lambda m: m.text_frame.text.strip() if m else None)(
                  find_marker(sl, min_top_emu=MIN_TOP))
               for i, sl in enumerate(final.slides, 1)}
    assert numbers == {1: None, 2: "2", 3: "3", 4: "4", 5: "5", 6: None, 7: "6"}, numbers

    # renumber's MECHANISM, not just its outcome. Derived from what
    # build_host.py and build_donor.py actually build, not guessed: slides
    # 2, 3 and 4 each carry a marker of one flavour or another and must come
    # back "fixed"; slides 5 and 7 (donor slide 1, and the host's pill
    # slide) are built with no marker at all and must come back "added".
    # Position 1 (the cover, skipped) and position 6 (hidden, never
    # visited) never enter the report at all. If flavour two's name match
    # regresses, slide 3 comes back "added" instead of "fixed" even though
    # the numbers dict above stays identical - this is what actually catches
    # that, not the dict.
    assert renumber_report == [
        (2, 2, "fixed"), (3, 3, "fixed"), (4, 4, "fixed"),
        (5, 5, "added"), (7, 6, "added"),
    ], renumber_report

    # The other side of the same defect: a slide that both "fixed" a stale
    # marker AND "added" a new one ends up with two page-number-shaped boxes,
    # which neither the numbers dict nor the report alone would catch if
    # renumber ever did both to the same slide. Count every text-bearing
    # shape at or below the marker floor, per slide, and refuse more than
    # one - what a human looking at the slide would notice immediately.
    for i, sl in enumerate(final.slides, 1):
        marker_shaped = [sh for sh in iter_shapes(sl.shapes)
                         if sh.has_text_frame and sh.text_frame.text.strip()
                         and sh.top is not None and sh.top >= MIN_TOP]
        assert len(marker_shaped) <= 1, (
            i, [sh.text_frame.text for sh in marker_shaped])

    # the package is internally consistent and carries exactly two font parts
    parts = read_package(v.final)
    assert dangling_refs(parts) == []
    assert len(font_parts(parts)) == 2
    print("round trip OK ->", v.final)


if __name__ == "__main__":
    main()
