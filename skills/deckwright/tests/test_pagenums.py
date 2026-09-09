"""The three flavours, in one deck, asserted in the saved package.

Catalogue item 14: printed page number is not file position, and the markers
come in three incompatible shapes - a real slide-number placeholder holding a
field, an inherited text box named for the locale, and a plain literal text
box on machine-built slides.
"""

import re
import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import PP_PLACEHOLDER

from deck_kit import merge, notes, pagenums
from deck_kit.brand import load_brand
from deck_kit.deck import add_slide, open_deck
from deck_kit.geometry import px
from deck_kit.primitives import textbox

# `template` arrives from the session-scoped fixture in conftest.py. Never a
# module constant pointing at out/: the gate runs the unit tests BEFORE it
# builds the template, and out/ is gitignored, so a hardcoded path fails on a
# clean checkout and, worse, silently uses a stale template on a dirty one.
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"

# _field_marker moved to examples/integration/build_host.py (Task 9) so the
# round-trip host and this test's deck build the same shape from one
# construction rather than two that can drift.
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "examples" / "integration"))
from build_host import _field_marker  # noqa: E402


@pytest.fixture(scope="module")
def brand():
    return load_brand("relay")


@pytest.fixture(scope="module")
def style(brand):
    s = brand.style()
    return pagenums.MarkerStyle(x=1172, y=690, w=60, h=14, size=8,
                                color=s.MUTED, font=brand.sans)


@pytest.fixture(scope="module")
def deck(tmp_path_factory, template, brand, style):
    """Five slides: 1 cover (unnumbered), 2 field marker, 3 locale-named box,
    4 literal box, 5 nothing at all. Slide 3 will be hidden."""
    prs = open_deck(template, brand)
    s = brand.style()
    for i in range(5):
        slide = add_slide(prs, brand)
        textbox(slide, 80, 80, 600, 40, "SLIDE %d" % (i + 1), 20, s.INK, brand.sans)
    slides = list(prs.slides)

    _field_marker(slides[1], brand, style)

    # Positioned well above the marker floor (min_top_emu), NOT at
    # style.x/style.y: at the marker's usual position the box's "99" text
    # would also satisfy the literal-digit fallback, so deleting the name
    # branch entirely would still leave test_find_marker_recognises_the_
    # locale_named_box green. Above the floor, only the name branch can
    # return this shape.
    box = textbox(slides[2], 900, 50, style.w, style.h, "99",
                  style.size, style.color, style.font)
    box._element.find(".//" + "{%s}cNvPr" % P).set("name", "Numero de diapositiva 3")
    slides[2]._element.set("show", "0")

    textbox(slides[3], style.x, style.y, style.w, style.h, "99",
            style.size, style.color, style.font)

    out = tmp_path_factory.mktemp("p") / "nums.pptx"
    prs.save(str(out))
    return out


def test_find_marker_recognises_the_field_placeholder(deck, style):
    slide = Presentation(str(deck)).slides[1]
    assert pagenums.find_marker(slide, min_top_emu=px(style.y) - 1) is not None


def test_find_marker_recognises_the_locale_named_box(deck, style):
    slide = Presentation(str(deck)).slides[2]
    marker = pagenums.find_marker(slide, min_top_emu=px(style.y) - 1)
    assert marker is not None and "diapositiva" in marker.name


def test_find_marker_recognises_the_plain_literal_box(deck, style):
    slide = Presentation(str(deck)).slides[3]
    assert pagenums.find_marker(slide, min_top_emu=px(style.y) - 1) is not None


def test_find_marker_ignores_body_text_that_happens_to_be_a_number(tmp_path, template, brand, style):
    """A literal marker is recognised by being small, low on the slide and
    numeric. A '5' in a headline at the top is not a page number."""
    s = brand.style()
    prs = open_deck(template, brand)
    slide = add_slide(prs, brand)
    textbox(slide, 80, 80, 200, 60, "5", 40, s.INK, brand.sans)
    src = tmp_path / "headline.pptx"
    prs.save(str(src))
    got = Presentation(str(src)).slides[0]
    assert pagenums.find_marker(got, min_top_emu=px(style.y) - 1) is None


def test_find_marker_recognises_a_real_slide_number_placeholder(tmp_path, template, brand, style):
    """Flavour one, for real: the content layout's own SLIDE_NUMBER
    placeholder, kept via add_slide(..., keep_placeholders=...) - separate
    from _field_marker's synthetic text box above, which stays as its own
    fixture.

    The cloned slide-level shape carries NO field of its own: python-pptx's
    clone_placeholder() produces a bare <p:sp> with no <p:txBody> at all -
    the <a:fld type="slidenum"> lives only on the layout, and PowerPoint
    renders the slide's number by inheriting from it. This test exercises
    find_marker's placeholder-type branch, not the field branch, and
    set_number later overrides the inherited number by writing a literal
    into what starts as an empty text body.

    python-pptx also names a cloned SLIDE_NUMBER placeholder "Slide Number
    Placeholder N", which itself matches LOCALE_NAMES - so deleting the
    placeholder-type branch would still leave this green by falling through
    to the name branch. Renamed to something neutral so only
    shape.placeholder_format.type can match.
    """
    prs = open_deck(template, brand)
    slide = add_slide(prs, brand,
                      keep_placeholders=(PP_PLACEHOLDER.SLIDE_NUMBER,))
    marker_shape = next(ph for ph in slide.placeholders
                        if ph.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER)
    marker_shape._element.find(".//" + "{%s}cNvPr" % P).set("name", "TextBox 9")
    out = tmp_path / "placeholder.pptx"
    prs.save(str(out))
    got = Presentation(str(out)).slides[0]
    marker = pagenums.find_marker(got, min_top_emu=px(style.y) - 1)
    assert marker is not None
    assert marker.is_placeholder
    assert marker.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER
    assert "Slide Number" not in marker.name and "slidenum" not in marker.name


def test_set_number_collapses_the_field_and_keeps_its_run_properties(tmp_path, deck, style):
    """The rPr lives inside the <a:fld>. Replacing the field with a plain run
    and dropping it leaves the number in the theme's default face and size."""
    prs = Presentation(str(deck))
    marker = pagenums.find_marker(prs.slides[1], min_top_emu=px(style.y) - 1)
    pagenums.set_number(marker, "2")
    out = tmp_path / "fixed.pptx"
    prs.save(str(out))

    parts = merge.read_package(out)
    xml = parts[merge.slide_order(parts)[1]].decode("utf-8")
    assert "slidenum" not in xml
    body = re.search(r'<p:sp>(?:(?!</p:sp>).)*?>2</a:t>.*?</p:sp>', xml, re.S).group(0)
    assert "<a:rPr" in body, "run properties were lost with the field"
    assert 'sz="%d"' % (style.size * 100) in body


def test_renumber_handles_all_three_flavours_and_skips_hidden(tmp_path, deck, style):
    """The whole of catalogue item 14 in one assertion.

    Visible order: 1 cover, 2 field, [3 hidden, not counted], 4 literal -> 3,
    5 no marker -> 4. Slide 1 is skipped as the cover.
    """
    prs = Presentation(str(deck))
    report = pagenums.renumber(prs, style, skip=(1,))
    out = tmp_path / "renumbered.pptx"
    prs.save(str(out))

    got = Presentation(str(out))
    numbers = {}
    for i, slide in enumerate(got.slides, 1):
        marker = pagenums.find_marker(slide, min_top_emu=px(style.y) - 1)
        numbers[i] = marker.text_frame.text.strip() if marker else None

    assert numbers[1] is None, "the cover must stay unnumbered"
    assert numbers[2] == "2"
    assert numbers[4] == "3", "the hidden slide must not consume a number"
    assert numbers[5] == "4", "a slide with no marker must get one"
    assert ("added" in [r[2] for r in report]) and ("fixed" in [r[2] for r in report])
    assert merge.dangling_refs(merge.read_package(out)) == []


def test_renumber_leaves_the_hidden_slide_alone(tmp_path, deck, style):
    prs = Presentation(str(deck))
    pagenums.renumber(prs, style, skip=(1,))
    out = tmp_path / "renumbered.pptx"
    prs.save(str(out))
    marker = pagenums.find_marker(Presentation(str(out)).slides[2],
                                  min_top_emu=px(style.y) - 1)
    assert marker.text_frame.text.strip() == "99", "hidden slide was renumbered"


def test_find_marker_ranks_by_flavour_not_z_order(tmp_path, template, brand, style):
    """C1: all three flavours on one slide, in adversarial z-order - the
    literal heuristic box FIRST, the real placeholder LAST. A single loop
    that returns on the first shape matching any branch would return the
    literal box; the fix ranks by flavour and the real placeholder must win
    regardless of where it sits in the shape tree."""
    prs = open_deck(template, brand)
    slide = add_slide(prs, brand, keep_placeholders=(PP_PLACEHOLDER.SLIDE_NUMBER,))
    placeholder = next(ph for ph in slide.placeholders
                       if ph.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER)

    literal = textbox(slide, style.x, style.y, style.w, style.h, "7",
                      style.size, style.color, style.font)
    named = textbox(slide, style.x, style.y, style.w, style.h, "named",
                    style.size, style.color, style.font)
    named._element.find(".//" + "{%s}cNvPr" % P).set("name", "Numero de diapositiva")

    # Move the real placeholder to the END of the shape tree, so z-order
    # alone would pick the literal box first, the named box second, and the
    # placeholder last - the exact reverse of the intended rank order.
    tree = placeholder._element.getparent()
    tree.remove(placeholder._element)
    tree.append(placeholder._element)

    out = tmp_path / "z_order.pptx"
    prs.save(str(out))
    got = Presentation(str(out)).slides[0]
    marker = pagenums.find_marker(got, min_top_emu=px(style.y) - 1)
    assert marker is not None
    assert marker.is_placeholder
    assert marker.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER


def test_renumber_does_not_overwrite_a_numeric_footnote_that_precedes_the_real_marker(
        tmp_path, template, brand, style):
    """The reproduced defect: a small numeric footnote at or below the
    marker floor, added BEFORE the real locale-named marker in z-order.
    Before the fix, find_marker returned the footnote (first hit in the
    loop) and renumber overwrote it with the page number, leaving the real
    marker stale and the footnote destroyed."""
    prs = open_deck(template, brand)
    slide = add_slide(prs, brand)
    footnote = textbox(slide, style.x, style.y, style.w, style.h, "1",
                       style.size, style.color, style.font)
    marker_box = textbox(slide, style.x, style.y + 20, style.w, style.h, "99",
                         style.size, style.color, style.font)
    marker_box._element.find(".//" + "{%s}cNvPr" % P).set(
        "name", "Numero de diapositiva")

    out = tmp_path / "footnote.pptx"
    prs.save(str(out))

    prs = Presentation(str(out))
    report = pagenums.renumber(prs, style, skip=())
    out2 = tmp_path / "footnote_renumbered.pptx"
    prs.save(str(out2))

    got = Presentation(str(out2)).slides[0]
    shapes = {sh.shape_id: sh for sh in got.shapes}
    assert shapes[footnote.shape_id].text_frame.text.strip() == "1", (
        "the footnote must be untouched by renumber")
    assert shapes[marker_box.shape_id].text_frame.text.strip() == "1", (
        "the real marker must carry the new page number")
    assert report and report[0][2] == "fixed"


def test_renumber_skips_by_visible_index_not_file_position(tmp_path, template, brand, style):
    """`skip` names VISIBLE positions, not file positions.

    Slide 1 is hidden; slides 2-4 are visible. File position 2 is therefore
    visible index 1. Calling renumber(..., skip=(2,)) tests file position 2
    against `skip` under the brief's original
    `if visible in skip or position in skip`, and visible index 1 against
    the controller's `if visible in skip`. The two disagree for exactly this
    slide: under `if visible in skip`, 1 is not in (2,), so it is NOT
    skipped and gets numbered "1"; under the disjunction, its file position
    2 IS in skip, so it stays unnumbered. Asserting it carries "1" pins the
    ruling, not just its cover-case restatement - reverting to the
    disjunction is checked separately, below.
    """
    prs = open_deck(template, brand)
    s = brand.style()
    hidden = add_slide(prs, brand)
    textbox(hidden, 80, 80, 600, 40, "HIDDEN", 20, s.INK, brand.sans)
    hidden._element.set("show", "0")
    for label in ("SLIDE 2", "SLIDE 3", "SLIDE 4"):
        slide = add_slide(prs, brand)
        textbox(slide, 80, 80, 600, 40, label, 20, s.INK, brand.sans)

    pagenums.renumber(prs, style, skip=(2,))
    out = tmp_path / "visible_skip.pptx"
    prs.save(str(out))

    got = Presentation(str(out))
    # File position 2 is got.slides[1] (0-based).
    marker = pagenums.find_marker(got.slides[1], min_top_emu=px(style.y) - 1)
    assert marker is not None, (
        "file position 2 (visible index 1) is not in skip=(2,) under "
        "visible-index-only semantics and must be numbered")
    assert marker.text_frame.text.strip() == "1"


def test_marker_style_colour_reaches_the_saved_xml(tmp_path, template, brand):
    """`MarkerStyle` carries the marker's colour; the core stores it as
    `color: object` and applies whatever the caller hands it. Renumbering one
    deck under two different colours must produce two different `srgbClr`
    values in the saved package - proof the colour actually travelled,
    which an `rglob` + `re.search(r"RGBColor\\(")` grep over src/ cannot show
    because pagenums.py never spells RGBColor at all."""
    s = brand.style()
    colours = (RGBColor(0x11, 0x22, 0x33), RGBColor(0xAA, 0xBB, 0xCC))
    seen = []
    for colour in colours:
        prs = open_deck(template, brand)
        slide = add_slide(prs, brand)
        textbox(slide, 80, 80, 600, 40, "SLIDE", 20, s.INK, brand.sans)
        style = pagenums.MarkerStyle(x=1172, y=690, w=60, h=14, size=8,
                                     color=colour, font=brand.sans)
        pagenums.renumber(prs, style, skip=())
        out = tmp_path / ("colour_%s.pptx" % str(colour))
        prs.save(str(out))

        parts = merge.read_package(out)
        xml = parts[merge.slide_order(parts)[0]].decode("utf-8")
        needle = 'srgbClr val="%s"' % str(colour)
        assert needle in xml
        seen.append(str(colour))
    assert seen[0] != seen[1]


def test_visible_numbers_skips_hidden_slides_and_the_cover(template, brand, style):
    prs = open_deck(template, brand)
    s = brand.style()
    for i in range(5):
        slide = add_slide(prs, brand)
        textbox(slide, 80, 80, 400, 40, "slide %d" % (i + 1), 14, s.GREY, brand.sans)
    prs.slides[2]._element.set("show", "0")
    got = pagenums.visible_numbers(prs, skip=(1,))
    assert got == {1: None, 2: 2, 3: None, 4: 3, 5: 4}
    # the walk is pure: nothing was written
    assert all(shape.text_frame.text != "2" for shape in prs.slides[1].shapes
               if shape.has_text_frame)


def test_renumber_report_agrees_with_visible_numbers(template, brand, style):
    prs = open_deck(template, brand)
    s = brand.style()
    for i in range(4):
        slide = add_slide(prs, brand)
        textbox(slide, 80, 80, 400, 40, "slide %d" % (i + 1), 14, s.GREY, brand.sans)
    prs.slides[1]._element.set("show", "0")
    expected = pagenums.visible_numbers(prs, skip=(1,))
    report = pagenums.renumber(prs, style, skip=(1,))
    assert {pos: vis for pos, vis, _ in report} == {
        pos: vis for pos, vis in expected.items() if vis is not None}


def test_show_false_is_hidden_everywhere_the_predicate_is_consumed(template, brand):
    """`show` is xsd:boolean, so a tool other than PowerPoint may write
    "false". `pagenums.is_hidden` is the single predicate, and both the footer
    counter and the teleprompter's numbering read it - if only one of them did,
    a deck hidden that way would show one number in the footer and another on
    the teleprompter."""
    prs = open_deck(template, brand)
    s = brand.style()
    for i in range(3):
        slide = add_slide(prs, brand)
        textbox(slide, 80, 80, 400, 40, "slide %d" % (i + 1), 14, s.GREY, brand.sans)
    prs.slides[1]._element.set("show", "false")
    assert pagenums.is_hidden(prs.slides[1])
    assert pagenums.visible_numbers(prs, skip=(1,)) == {1: None, 2: None, 3: 2}
    assert notes.numbering(prs, skip=(1,)) == {
        1: (None, False), 2: (None, True), 3: (2, False)}
