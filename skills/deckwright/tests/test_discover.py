"""discover.py - every test runs over a tree this test built under tmp_path."""

import json
import subprocess
import sys
from pathlib import Path

import pytest
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "discover.py"
sys.path.insert(0, str(ROOT / "scripts"))

import discover  # noqa: E402

_BRAND_PY = '''
from pathlib import Path
from deck_kit.brand import BrandSpec
HERE = Path(__file__).resolve().parent
BRAND = BrandSpec(name=%r, master_name=%r, layout_name=%r, font_dir=HERE / "fonts")
'''


def _brand(brands, name, master, layout, *, tokens=False, fonts=False, icons=False):
    pkg = brands / name
    pkg.mkdir(parents=True)
    (pkg / "brand.py").write_text(_BRAND_PY % (name, master, layout), encoding="utf-8")
    if tokens:
        (pkg / "tokens.css").write_text(":root { --ink: #000; }", encoding="utf-8")
    if fonts:
        (pkg / "fonts").mkdir()
    if icons:
        (pkg / "icons").mkdir()
    return pkg


def _system(path, *, tokens_dir=False, styles=False, readme=False, skill=False,
            manifest=False, extras=()):
    """A downloaded design system on disk, in whichever shape is asked for."""
    path.mkdir(parents=True, exist_ok=True)
    if tokens_dir:
        (path / "tokens").mkdir()
        (path / "tokens" / "colors.css").write_text(":root{--ink:#0A0A0A}", encoding="utf-8")
    if styles:
        (path / "styles.css").write_text('@import url("tokens/colors.css");', encoding="utf-8")
    if readme:
        (path / "README.md").write_text("# A system", encoding="utf-8")
    if skill:
        (path / "SKILL.md").write_text("---\nname: a\n---\n", encoding="utf-8")
    if manifest:
        (path / "_ds_manifest.json").write_text("{}", encoding="utf-8")
    for extra in extras:
        (path / extra).mkdir(parents=True)
    return path


@pytest.fixture
def tree(tmp_path):
    """A working tree with two brands, one readable master, one unreadable
    file with a deck suffix, a tokens.css, one downloaded design system that
    has no brand package, and decoys under the skipped dirs.

    Two brands, so the report's multi-brand branch ("choose a brand: ...") is
    exercised here. The single-brand branch is exercised against the real tree
    by test_the_repository_recommends_its_one_brand, and directly on
    `recommend` below.
    """
    root = tmp_path / "work"
    brands = tmp_path / "brands"
    _brand(brands, "acme", "Acme Master", "Acme Content", tokens=True, fonts=True, icons=True)
    _brand(brands, "beta", "Beta Master", "Beta Content")
    (brands / "not-a-brand").mkdir()          # no brand.py: not a package
    # A downloaded system next to the brand packages, in the shape Relay
    # arrived in: a tokens/ directory, a styles.css with a readme, its own
    # manifest, and no brand.py anywhere in it.
    _system(brands / "vendor_design_system", tokens_dir=True, styles=True,
            readme=True, manifest=True, extras=("guidelines", "assets/icons"))

    (root / "masters").mkdir(parents=True)
    prs = Presentation()
    prs.slide_masters[0].name = "Corp Master"
    prs.save(str(root / "masters" / "corporate.pptx"))
    (root / "masters" / "broken.pptx").write_text("not a deck", encoding="utf-8")
    (root / "tokens.css").write_text(":root {}", encoding="utf-8")

    # Decoys the search must never enter.
    for skipped in ("out", ".git", ".venv", ".superpowers", "node_modules"):
        d = root / skipped / "deep"
        d.mkdir(parents=True)
        prs.save(str(d / "built.pptx"))
        (d / "tokens.css").write_text(":root {}", encoding="utf-8")
        _system(d / "hidden_system", tokens_dir=True)
    return root, brands


def test_the_report_lists_brands_masters_tokens_and_a_recommendation(tree):
    root, brands = tree
    report = discover.discover(root, brands)

    assert [b["name"] for b in report["brands"]] == ["acme", "beta"]
    acme, beta = report["brands"]
    assert acme["master_name"] == "Acme Master" and acme["layout_name"] == "Acme Content"
    assert (acme["tokens_css"], acme["fonts"], acme["icons"]) == (True, True, True)
    assert (beta["tokens_css"], beta["fonts"], beta["icons"]) == (False, False, False)

    paths = {m["path"]: m for m in report["masters"]}
    assert set(paths) == {str(Path("masters/broken.pptx")), str(Path("masters/corporate.pptx"))}
    corporate = paths[str(Path("masters/corporate.pptx"))]
    assert corporate["masters"][0]["name"] == "Corp Master"
    assert "Title Slide" in corporate["masters"][0]["layouts"]

    assert report["tokens"] == ["tokens.css"]
    assert report["recommendation"] == "choose a brand: acme, beta"


def test_the_unreadable_deck_is_a_note_not_a_crash(tree):
    root, brands = tree
    broken = next(m for m in discover.find_masters(root) if m["path"].endswith("broken.pptx"))
    assert "masters" not in broken
    assert broken["error"].startswith("PackageNotFoundError")
    text = discover.render(discover.discover(root, brands))
    assert "broken.pptx: skipped, not readable - PackageNotFoundError" in text
    assert "Traceback" not in text


def test_the_text_report_carries_every_section(tree):
    root, brands = tree
    text = discover.render(discover.discover(root, brands))
    assert "brands (2)" in text
    assert "  acme: master 'Acme Master', layout 'Acme Content'" in text
    assert "    tokens.css yes, fonts/ yes, icons/ yes" in text
    assert "    tokens.css no, fonts/ no, icons/ no" in text
    assert "candidate masters (2)" in text
    assert "    master 'Corp Master': layouts 'Title Slide', " in text
    assert "tokens.css (1)" in text
    assert text.rstrip().endswith("recommendation: choose a brand: acme, beta")


def test_skipped_directories_are_never_searched(tree):
    root, _ = tree
    masters = [m["path"] for m in discover.find_masters(root)]
    assert not any("built.pptx" in p for p in masters)
    assert discover.find_tokens(root) == ["tokens.css"]


def test_recommendation_walks_one_brand_then_a_system_then_the_default(tree):
    """Every branch of `recommend`, in order of precedence. A candidate
    master is NOT one of them any more: this pipeline starts from a design
    system and HTML examples, so a .pptx on the machine is inventory and
    recommending one would offer a route that does not exist."""
    root, brands = tree
    records = discover.find_brands(brands)
    masters = discover.find_masters(root)
    one = [b for b in records if b["name"] == "acme"]
    assert discover.recommend(one, masters) == "use brand acme"
    assert discover.recommend(records, masters) == "choose a brand: acme, beta"
    assert discover.recommend([], masters) == (
        "no design system of your own found: use the default, brands/relay")
    assert discover.recommend([], []) == (
        "no design system of your own found: use the default, brands/relay")


def test_a_downloaded_design_system_is_reported_with_what_it_holds(tree):
    """The Relay shape: tokens/ rather than tokens.css, no brand.py, no
    master. Before this the report said "no design system found" and sent
    the user to the default brand with their own system sitting on disk."""
    root, brands = tree
    [system] = discover.discover(root, brands)["systems"]
    assert system["name"] == "vendor_design_system"
    assert system["path"] == str(brands / "vendor_design_system")
    assert system["holds"] == ["tokens/*.css", "styles.css", "_ds_manifest.json",
                               "guidelines/", "assets/", "assets/icons/"]


def test_each_of_the_three_shapes_is_evidence_on_its_own(tmp_path):
    """Any one marker is enough; a directory with none of them is not a
    design system, however many CSS files it happens to carry."""
    assert discover.system_holdings(_system(tmp_path / "a", tokens_dir=True)) == ["tokens/*.css"]
    assert discover.system_holdings(
        _system(tmp_path / "b", styles=True, readme=True)) == ["styles.css"]
    assert discover.system_holdings(
        _system(tmp_path / "c", styles=True, skill=True)) == ["styles.css"]
    assert discover.system_holdings(_system(tmp_path / "d", manifest=True)) == ["_ds_manifest.json"]
    # styles.css alone is not evidence: a single stylesheet is a stylesheet.
    assert discover.system_holdings(_system(tmp_path / "e", styles=True)) == []
    # An empty tokens/ is not evidence either, and neither is a bare dir.
    (tmp_path / "f" / "tokens").mkdir(parents=True)
    assert discover.system_holdings(tmp_path / "f") == []
    assert discover.system_holdings(tmp_path / "nowhere") == []


def test_a_directory_with_a_brand_py_is_a_brand_package_not_a_loose_system(tmp_path):
    """A brand package ships a tokens.css rather than a tokens/ directory,
    but a brand that grew one must still be reported once, as a brand."""
    brands = tmp_path / "brands"
    pkg = _brand(brands, "acme", "M", "L")
    _system(pkg, tokens_dir=True, styles=True, readme=True)
    assert discover.find_systems(tmp_path, brands) == []
    assert [b["name"] for b in discover.find_brands(brands)] == ["acme"]


def test_a_system_is_reported_once_even_when_both_trees_reach_it(tmp_path):
    """The brands dir is usually under --root. Walking both must not list
    the same directory twice."""
    brands = tmp_path / "brands"
    _system(brands / "dup_system", tokens_dir=True)
    found = discover.find_systems(tmp_path, brands)
    assert [s["name"] for s in found] == ["dup_system"]
    assert found[0]["path"] == str(Path("brands/dup_system"))


def test_skipped_directories_hide_a_design_system_too(tree):
    root, brands = tree
    names = [s["name"] for s in discover.find_systems(root, brands)]
    assert "hidden_system" not in names


def test_the_recommendation_sets_a_brand_up_from_a_system_ahead_of_a_master(tree):
    """A system that ships its own tokens outranks a .pptx somebody dropped
    in, and both are outranked by a brand package that already exists."""
    root, brands = tree
    systems = discover.find_systems(root, brands)
    masters = discover.find_masters(root)
    assert discover.recommend([], masters, systems) == (
        "set up a brand from design system vendor_design_system (%s): "
        "SKILL.md, Set up a brand" % (brands / "vendor_design_system"))
    assert discover.recommend([], [], systems) == discover.recommend([], masters, systems)
    # A brand package still wins over both.
    one = [b for b in discover.find_brands(brands) if b["name"] == "acme"]
    assert discover.recommend(one, masters, systems) == "use brand acme"


def test_the_text_report_carries_the_design_system_section(tree):
    root, brands = tree
    text = discover.render(discover.discover(root, brands))
    assert "design systems without a deckwright brand package (1)" in text
    assert ("  vendor_design_system: tokens/*.css, styles.css, _ds_manifest.json, "
            "guidelines/, assets/, assets/icons/") in text
    assert "    path: %s" % (brands / "vendor_design_system") in text


# --- the shapes an earlier version of this script did not recognise -------
# Reproduced from an unattended run of the skill against a real design
# system: `design_system/` held README.md, SKILL.md, colors_and_type.css,
# assets/, preview/, slides/ and ui_kits/. None of the three shapes the
# script knew matched, so the system went unreported - and the ONE directory
# it did name was a sample UI kit two levels inside it, which happened to
# carry a styles.css and a readme. The child was recommended and the parent
# was invisible.


def _pointed_at_system(base):
    """The corporate design-system shape: a named token sheet, a readme, a SKILL.md, and a
    sample UI kit inside that qualifies on the old rules all by itself."""
    d = base / "design_system"
    kit = d / "ui_kits" / "analytics-dashboard"
    (kit / "tokens").mkdir(parents=True)
    (d / "colors_and_type.css").write_text(
        ":root {\n  --ink-900: #111111;\n  --type-body: 16px;\n}\n", encoding="utf-8")
    (d / "README.md").write_text("# The design system", encoding="utf-8")
    (d / "SKILL.md").write_text("---\nname: baa\n---\n", encoding="utf-8")
    (d / "assets").mkdir()
    (kit / "styles.css").write_text(".card { color: #111; }", encoding="utf-8")
    (kit / "README.md").write_text("# A sample kit", encoding="utf-8")
    (kit / "tokens" / "kit.css").write_text(":root{--kit:1}", encoding="utf-8")
    return d


def test_a_named_token_sheet_beside_a_readme_is_a_design_system(tmp_path):
    """`colors_and_type.css` is a token sheet; the old evidence test knew
    only `styles.css`, `tokens/` and `_ds_manifest.json` and reported "no
    design system found" for a directory that unmistakably is one."""
    d = _pointed_at_system(tmp_path)
    holds = discover.system_holdings(d)
    assert holds == ["colors_and_type.css", "ui_kits/", "assets/"]


def test_the_outermost_directory_is_the_system_not_the_kit_inside_it(tmp_path):
    """The failure this rewrite is named for. The kit qualifies on its own -
    styles.css beside a README - and must not be reported once the system
    above it has been."""
    _pointed_at_system(tmp_path)
    found = discover.find_systems(tmp_path, tmp_path / "no-brands")
    assert [s["name"] for s in found] == ["design_system"]
    assert "analytics-dashboard" not in discover.render(
        discover.discover(tmp_path, tmp_path / "no-brands"))


def test_a_stylesheet_that_declares_root_custom_properties_is_evidence_alone(tmp_path):
    """No readme, no conventional filename: a sheet that sets `--x` on
    `:root` is a token sheet by construction."""
    d = tmp_path / "whatever"
    d.mkdir()
    (d / "site.css").write_text("@media print { a { color: red } }", encoding="utf-8")
    assert discover.system_holdings(d) == []
    (d / "site.css").write_text(":root, .dark {\n  --accent: #ff4f00;\n}\n",
                                encoding="utf-8")
    assert discover.system_holdings(d) == ["site.css (:root tokens)"]


def test_a_user_design_system_outranks_the_repositorys_own_default_brand(tree, tmp_path):
    """"The repository ships one brand" is not a fact about the user's
    identity. Before this, a tree with relay and the user's own system on it
    recommended relay, which is how a deck gets built in the wrong palette."""
    root, _ = tree
    _pointed_at_system(tmp_path)
    systems = discover.find_systems(tmp_path, tmp_path / "no-brands")
    relay_only = [{"name": "relay", "master_name": "M", "layout_name": "L"}]
    assert discover.recommend(relay_only, [], systems).startswith(
        "set up a brand from design system design_system")
    # With nothing of the user's found, the default is still the answer.
    assert discover.recommend(relay_only, [], []) == "use brand relay"
    # A brand package the user built themselves still outranks a loose system.
    theirs = relay_only + [{"name": "acme", "master_name": "M", "layout_name": "L"}]
    assert discover.recommend(theirs, [], systems) == "use brand acme"


def test_a_declared_system_is_recommended_over_a_directory_that_merely_has_root_vars(tmp_path):
    """Observed on the real tree: a deck directory carrying `theme.css` and
    `dya-matrix.css` was reported as a design system and, sorted by name
    alone, came out ahead of the actual one. Both are still listed; the
    stronger evidence is what the recommendation names."""
    _pointed_at_system(tmp_path)
    deck = tmp_path / "aaa-deck"          # sorts first by name
    deck.mkdir()
    (deck / "theme.css").write_text(":root{--x:1}", encoding="utf-8")
    found = discover.find_systems(tmp_path, tmp_path / "no-brands")
    assert [s["name"] for s in found] == ["design_system", "aaa-deck"]
    assert discover.recommend([], [], found).startswith(
        "set up a brand from design system design_system")


def test_the_pointed_design_system_is_reported_first_verified_and_recommended(tmp_path):
    d = _pointed_at_system(tmp_path)
    report = discover.discover(tmp_path, tmp_path / "no-brands", d)
    assert report["pointed"]["path"] == str(d.resolve())
    assert report["pointed"]["exists"] is True
    assert report["pointed"]["holds"] == ["colors_and_type.css", "ui_kits/", "assets/"]
    assert report["recommendation"] == (
        "use the design system you pointed at: %s - set a brand up from it "
        "(SKILL.md, Set up a brand)" % d.resolve())
    text = discover.render(report)
    assert text.splitlines()[0] == "recommendation: " + report["recommendation"]
    assert "design system you pointed at" in text
    # First: before the brands section, which is where every other report starts.
    assert text.index("design system you pointed at") < text.index("brands (")


def test_a_pointed_directory_with_no_markers_is_still_taken_as_given(tmp_path):
    """The user pointing at it is the strongest evidence there is. The check
    buys the sentence that follows, not a veto."""
    d = tmp_path / "bare"
    d.mkdir()
    pointed = discover.verify_system(d)
    assert pointed["exists"] is True and pointed["holds"] == []
    assert "check it" in pointed["note"]
    assert discover.recommend([], [], [], pointed).startswith(
        "use the design system you pointed at")


def test_the_cli_exits_two_on_a_missing_design_system(tmp_path):
    proc = _run("--root", str(tmp_path), "--design-system", str(tmp_path / "nope"))
    assert proc.returncode == 2
    assert "no such design system" in proc.stderr
    assert "Traceback" not in proc.stderr


# --- the report a human reads --------------------------------------------

def test_the_recommendation_is_printed_first_as_well_as_last(tree):
    """563 lines of report with the one actionable line at the bottom is a
    report nobody reads to the end."""
    root, brands = tree
    text = discover.render(discover.discover(root, brands))
    lines = text.splitlines()
    assert lines[0] == "recommendation: choose a brand: acme, beta"
    assert lines[-1] == lines[0]


def test_a_directory_of_many_masters_is_one_line_with_a_count(tmp_path):
    """80 of the 90 candidate masters in the run that prompted this were
    copies of one template under one deck's build directory."""
    import shutil
    work = tmp_path / "work"
    many = work / "archive"
    many.mkdir(parents=True)
    prs = Presentation()
    prs.slide_masters[0].name = "Corp Master"
    first = many / "a01.pptx"
    prs.save(str(first))
    for i in range(2, 8):
        shutil.copy(first, many / ("a%02d.pptx" % i))
    lines = discover.render_masters(discover.find_masters(work))
    collapsed = [l for l in lines if "collapsed" in l]
    assert len(collapsed) == 1, lines
    assert "7 files" in collapsed[0]
    assert "Corp Master" not in "\n".join(lines), "the group prints no layout list"


def test_a_long_layout_list_is_truncated_with_a_count(tmp_path):
    work = tmp_path / "work"
    work.mkdir()
    Presentation().save(str(work / "one.pptx"))   # 11 default layouts
    text = "\n".join(discover.render_masters(discover.find_masters(work)))
    assert "(+5 more)" in text, text


def test_build_and_archive_directories_are_never_searched(tmp_path):
    work = tmp_path / "work"
    for skipped in ("build", "dist", "_archive"):
        d = work / skipped
        d.mkdir(parents=True)
        Presentation().save(str(d / "built.pptx"))
        (d / "tokens.css").write_text(":root{}", encoding="utf-8")
        _system(d / "hidden_system", tokens_dir=True)
    assert discover.find_masters(work) == []
    assert discover.find_tokens(work) == []
    assert discover.find_systems(work, work / "no-brands") == []


def test_the_repository_reports_a_downloaded_design_system_when_one_is_there():
    """Against the real tree.

    `brands/relay_design_system/` is the system `brands/relay` was derived
    from. It is NOT committed - it is a third-party download, and what this
    repository carries is the derivation, not the input - so on a clean
    checkout there is nothing to find and this test says exactly that. An
    earlier version called `next(...)` on the search and raised StopIteration
    on any tree without it, which is every tree anyone else clones.
    """
    systems = discover.find_systems(ROOT, discover.brands_dir())
    names = [s["name"] for s in systems]
    if not (ROOT / "brands" / "relay_design_system").is_dir():
        assert "relay_design_system" not in names, names
        return
    relay = next(s for s in systems if s["name"] == "relay_design_system")
    assert "tokens/*.css" in relay["holds"]
    assert relay["path"] == str(Path("brands/relay_design_system"))


def test_the_repository_recommends_its_one_brand():
    """Against the real tree. `brands/relay` is the only brand package, so the
    intake's first answer is decided by looking rather than by asking.

    This replaces an assertion that `brands/example` and `brands/example2`
    were marked "test fixture, not a design system" and skipped by the
    recommendation. Both packages are gone, and with them the special case in
    discover.py: every brand on disk is now a candidate, and the report says
    "use brand relay" because relay is the only one - not because two others
    were filtered out.
    """
    records = discover.find_brands(discover.brands_dir())
    assert [b["name"] for b in records] == ["relay"], records
    assert all("fixture" not in b for b in records), \
        "the fixture flag went with the fixture brands"
    text = discover.render(discover.discover(ROOT, discover.brands_dir()))
    assert "test fixture, not a design system" not in text
    assert text.rstrip().endswith("recommendation: use brand relay")


def test_a_brand_that_declares_a_font_it_does_not_have_is_reported(tmp_path):
    """The intake has to say "run fetch_fonts.py" before a builder measures
    against a face that is not on disk - the build would otherwise fail
    mid-deck, after the template is already written."""
    brands = tmp_path / "brands"
    pkg = _brand(brands, "acme", "M", "L", fonts=True)
    (pkg / "fonts.toml").write_text(
        '[[font]]\n'
        'family = "Some Face"\n'
        'role = "display"\n'
        'weight = "bold"\n'
        'file = "SomeFace-Bold.ttf"\n'
        'url = "https://example.invalid/SomeFace-Bold.ttf"\n'
        'sha256 = ""\n'
        'licence = "SIL Open Font License 1.1"\n'
        'licence_url = "https://example.invalid/OFL.txt"\n'
        'licence_file = "OFL.txt"\n', encoding="utf-8")
    [record] = [b for b in discover.find_brands(brands) if b["name"] == "acme"]
    assert record["missing_fonts"] == [["Some Face", "bold", "SomeFace-Bold.ttf"]]
    text = discover.render(discover.discover(tmp_path, brands))
    assert "MISSING FONT Some Face bold (SomeFace-Bold.ttf)" in text
    assert "scripts/fetch_fonts.py --brand acme" in text


def test_a_brand_with_no_font_manifest_reports_no_missing_fonts(tmp_path):
    brands = tmp_path / "brands"
    _brand(brands, "acme", "M", "L")
    [record] = [b for b in discover.find_brands(brands) if b["name"] == "acme"]
    assert record["missing_fonts"] == []


def test_a_brand_that_does_not_load_is_reported_not_fatal(tmp_path):
    brands = tmp_path / "brands"
    pkg = brands / "half"
    pkg.mkdir(parents=True)
    (pkg / "brand.py").write_text("BRAND = BrandSpec(", encoding="utf-8")
    [half] = discover.find_brands(brands)
    assert half["error"].startswith("SyntaxError")
    assert "does not load - SyntaxError" in discover.render(discover.discover(tmp_path, brands))
    assert discover.recommend([half], []) == (
        "no design system of your own found: use the default, brands/relay")


def _run(*args):
    return subprocess.run([sys.executable, str(SCRIPT), *args], capture_output=True,
                          text=True, encoding="utf-8")


def test_the_cli_prints_json_and_exits_zero(tree):
    root, brands = tree
    proc = _run("--root", str(root), "--brands", str(brands), "--json")
    assert proc.returncode == 0, proc.stdout + proc.stderr
    got = json.loads(proc.stdout)
    assert got["recommendation"] == "choose a brand: acme, beta"
    assert [b["name"] for b in got["brands"]] == ["acme", "beta"]
    assert got["tokens"] == ["tokens.css"]
    assert got == discover.discover(root.resolve(), brands.resolve())


def test_the_cli_text_report_matches_the_renderer(tree):
    root, brands = tree
    proc = _run("--root", str(root), "--brands", str(brands))
    assert proc.returncode == 0, proc.stdout + proc.stderr
    assert proc.stdout.strip() == discover.render(
        discover.discover(root.resolve(), brands.resolve()))


def test_the_cli_exits_two_on_a_missing_root(tmp_path):
    proc = _run("--root", str(tmp_path / "nowhere"))
    assert proc.returncode == 2
    assert "no such root" in proc.stderr
    assert "Traceback" not in proc.stderr


def test_the_repository_itself_reports_its_own_brands():
    """Against the real tree: one brand package, and it loads."""
    brands = discover.find_brands(discover.brands_dir())
    assert {b["name"] for b in brands} == {"relay"}
    assert all("error" not in b for b in brands)
    [relay] = brands
    assert relay["master_name"] == "Relay Master"
    assert (relay["tokens_css"], relay["fonts"], relay["icons"]) == (True, True, True)
    assert relay["missing_fonts"] == [], relay["missing_fonts"]
