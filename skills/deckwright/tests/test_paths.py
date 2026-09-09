import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

from deck_kit.paths import arrow_head, bezier_points, polyline

BLUE = RGBColor(0x25, 0x63, 0xEB)


@pytest.fixture
def slide():
    prs = Presentation()
    return prs.slides.add_slide(prs.slide_layouts[6])


def test_move_command_seeds_the_first_point():
    assert bezier_points("M10 20") == [(10.0, 20.0)]


def test_cubic_curve_starts_and_ends_on_its_endpoints():
    pts = bezier_points("M0 0 C0 50 100 50 100 0", steps=8)
    assert pts[0] == (0.0, 0.0)
    assert pts[-1] == pytest.approx((100.0, 0.0))
    assert len(pts) == 9


def test_cubic_curve_bulges_between_the_endpoints():
    pts = bezier_points("M0 0 C0 50 100 50 100 0", steps=8)
    assert max(y for _, y in pts) > 0


def test_smooth_command_reflects_the_previous_control_point():
    """S must mirror the previous C's second control point; a parser that
    ignores the reflection produces a visible kink."""
    pts = bezier_points("M0 0 C0 40 40 40 40 0 S80 -40 80 0", steps=6)
    assert pts[-1] == pytest.approx((80.0, 0.0))
    assert min(y for _, y in pts[7:]) < 0, "the reflected segment bulges the other way"


def test_leading_dot_decimals_are_parsed():
    """Every SVG optimiser emits ".5". A number pattern that requires a digit
    before the dot reads that as 5 - a ten times coordinate error, silently."""
    assert bezier_points("M.5 .5") == [(0.5, 0.5)]
    pts = bezier_points("M.5 .5 C.5 10.5 20.5 10.5 20.5 .5", steps=4)
    assert pts[0] == (0.5, 0.5)
    assert pts[-1] == pytest.approx((20.5, 0.5))


def test_negative_and_exponent_numbers_are_parsed():
    assert bezier_points("M-1.5 -2") == [(-1.5, -2.0)]
    assert bezier_points("M1e2 0") == [(100.0, 0.0)]


def test_unsupported_command_raises_instead_of_being_skipped():
    """A skipped command does not vanish: its numbers get consumed as extra
    parameters of whichever command is still active, and the geometry comes
    out wrong with no error anywhere."""
    with pytest.raises(ValueError, match="unsupported path command 'Q'"):
        bezier_points("M0 0 Q50 50 100 0")
    with pytest.raises(ValueError, match="unsupported path command 'c'"):
        bezier_points("M0 0 c60 0 60 -40 120 -40")


def test_curve_before_any_move_raises():
    with pytest.raises(ValueError, match="before any M"):
        bezier_points("C0 50 100 50 100 0")


def test_coordinates_without_a_command_raise():
    with pytest.raises(ValueError, match="not a command"):
        bezier_points("10 20")


def test_truncated_coordinate_run_raises():
    with pytest.raises(ValueError, match="missing coordinates"):
        bezier_points("M0 0 C0 50 100")


def test_empty_path_raises():
    with pytest.raises(ValueError, match="no points"):
        bezier_points("")


def test_polyline_builds_a_freeform_with_no_fill_and_the_given_line(slide):
    sp = polyline(slide, [(0, 0), (10, 10), (20, 0)], 100, 200, BLUE, 1.5)
    assert sp.line.color.rgb == BLUE
    assert sp._element.spPr.find(qn("a:noFill")) is not None
    assert sp.line.width.pt == pytest.approx(1.5)


def test_polyline_offsets_every_point_by_the_origin(slide):
    a = polyline(slide, [(0, 0), (10, 0)], 0, 0, BLUE, 1.0)
    b = polyline(slide, [(0, 0), (10, 0)], 100, 0, BLUE, 1.0)
    assert b.left > a.left


def test_polyline_can_be_dashed(slide):
    plain = polyline(slide, [(0, 0), (10, 0)], 0, 0, BLUE, 1.0)
    dashed_ = polyline(slide, [(0, 0), (10, 0)], 0, 0, BLUE, 1.0, dash=True)
    assert plain.line.dash_style != dashed_.line.dash_style


def test_arrow_head_rotates_per_direction(slide):
    assert arrow_head(slide, 50, 50, "up", BLUE).rotation == 0
    assert arrow_head(slide, 50, 50, "right", BLUE).rotation == 90
    assert arrow_head(slide, 50, 50, "down", BLUE).rotation == 180
    assert arrow_head(slide, 50, 50, "left", BLUE).rotation == 270


def test_arrow_head_is_centred_on_the_point_it_is_given(slide):
    from deck_kit.geometry import px
    sp = arrow_head(slide, 100, 200, "right", BLUE, w=10, h=8)
    assert sp.left == px(95)
    assert sp.top == px(196)


def test_arrow_head_rejects_an_unknown_direction(slide):
    with pytest.raises(KeyError):
        arrow_head(slide, 0, 0, "sideways", BLUE)
