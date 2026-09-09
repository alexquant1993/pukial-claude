import re
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation

from deck_kit.brand import load_brand

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from make_template import variant_spec  # noqa: E402


def _run(*args):
    return subprocess.run([sys.executable, *args], capture_output=True, text=True)


@pytest.fixture(scope="module")
def template(tmp_path_factory):
    out = tmp_path_factory.mktemp("tpl") / "template.pptx"
    r = _run("scripts/make_template.py", "--brand", "relay", "--out", str(out))
    assert r.returncode == 0, r.stdout + r.stderr
    return out


def test_template_has_the_declared_slide_size(template):
    spec = load_brand("relay")
    prs = Presentation(str(template))
    assert (prs.slide_width, prs.slide_height) == spec.slide_size_emu


def test_template_has_a_master_and_layout_with_the_declared_names(template):
    spec = load_brand("relay")
    prs = Presentation(str(template))
    assert spec.master_name in [m.name for m in prs.slide_masters]
    master = next(m for m in prs.slide_masters if m.name == spec.master_name)
    assert spec.layout_name in [l.name for l in master.slide_layouts]


def test_template_keeps_exactly_one_layout(template):
    spec = load_brand("relay")
    prs = Presentation(str(template))
    master = next(m for m in prs.slide_masters if m.name == spec.master_name)
    assert [l.name for l in master.slide_layouts] == [spec.layout_name]


def test_template_ships_with_no_slides(template):
    assert len(Presentation(str(template)).slides) == 0


def test_template_layout_carries_a_background_shape(template):
    spec = load_brand("relay")
    prs = Presentation(str(template))
    master = next(m for m in prs.slide_masters if m.name == spec.master_name)
    layout = next(l for l in master.slide_layouts if l.name == spec.layout_name)
    covering = [sh for sh in layout.shapes
                if sh.width == prs.slide_width and sh.height == prs.slide_height]
    assert covering, "the layout must supply the background"


def test_inspect_passes_on_a_generated_template(template):
    r = _run("scripts/inspect_template.py", "--brand", "relay", str(template))
    assert r.returncode == 0, r.stdout + r.stderr
    assert "PASS" in r.stdout
    assert "FAIL" not in r.stdout


def test_inspect_fails_and_names_the_clause_on_a_foreign_deck(tmp_path):
    foreign = tmp_path / "foreign.pptx"
    Presentation().save(str(foreign))          # default 10x7.5in, unnamed master
    r = _run("scripts/inspect_template.py", "--brand", "relay", str(foreign))
    assert r.returncode != 0
    assert "FAIL" in r.stdout
    assert "slide size" in r.stdout
    assert "master" in r.stdout


def test_inspect_names_the_remedy_for_an_unnamed_master(tmp_path):
    """python-pptx's default presentation has an unnamed master, which is
    also what a real corporate template looks like. "Fix what it reports"
    had no documented fix for that case, and the obvious workaround -
    master_name = "" - passed the verifier while binding by index."""
    foreign = tmp_path / "unnamed.pptx"
    Presentation().save(str(foreign))
    r = _run("scripts/inspect_template.py", "--brand", "relay", str(foreign))
    assert r.returncode != 0
    assert "remedy:" in r.stdout
    assert "Slide Master view" in r.stdout
    assert "make_template.py" in r.stdout


def test_inspect_prints_no_remedy_when_every_clause_passes(template):
    r = _run("scripts/inspect_template.py", "--brand", "relay", str(template))
    assert "remedy:" not in r.stdout


# --- the variant master: what makes the Phase 2 transplant cross-master ------

@pytest.fixture(scope="module")
def variant_template(tmp_path_factory):
    out = tmp_path_factory.mktemp("tpl2") / "template2.pptx"
    r = _run("scripts/make_template.py", "--brand", "relay",
             "--variant", "donor", "--out", str(out))
    assert r.returncode == 0, r.stdout + r.stderr
    return out


def test_a_variant_spec_still_refuses_an_index_shaped_selector():
    """`variant_spec` goes through `dataclasses.replace`, which re-runs
    BrandSpec's `__post_init__` - so the by-name-never-by-index rule survives
    the rename rather than becoming a property of the original object only.

    A variant name is appended to a string, so `variant_spec(spec, 3)` still
    yields a NAME ("Relay Master (3)"), which is correct and is asserted here
    rather than being confused with the rule. The rule is what stops a
    caller reaching past `variant_spec` and putting an index in directly.
    """
    from dataclasses import replace

    spec = load_brand("relay")
    v = variant_spec(spec, "donor")
    assert v.master_name == "Relay Master (donor)"
    assert v.layout_name == "Relay Content (donor)"
    assert spec.master_name == "Relay Master", "the loaded brand was mutated"
    assert variant_spec(spec, 3).master_name == "Relay Master (3)"
    with pytest.raises(TypeError, match="by name"):
        replace(spec, layout_name=0)


def _master_xml(path):
    with zipfile.ZipFile(str(path)) as z:
        names = [n for n in z.namelist()
                 if n.startswith("ppt/slideMasters/") and n.endswith(".xml")]
        return {n: z.read(n).decode("utf-8") for n in names}


def test_a_variant_template_is_a_second_master_of_the_same_brand(template,
                                                                variant_template):
    """The Phase 2 round trip transplants slides ACROSS masters - catalogue
    item 5, the reason deck_kit.merge exists. Two templates from one brand
    with no variant flag are the same master twice, and the whole transplant
    path runs against the easiest possible input. This asserts the three
    things that make the pair a real cross-master merge, against the files:

      - the two masters occupy the SAME part name, so merge has to mint a
        fresh one in the destination;
      - they carry the SAME p:sldLayoutId, so merge's re-mint resolves a
        genuine collision rather than being a no-op;
      - their master and layout names DIFFER, so deck._layout can still
        resolve a layout by name in the merged file.

    This is the coverage that `brands/example2` used to supply as a second
    brand package.
    """
    host, donor = _master_xml(template), _master_xml(variant_template)
    assert list(host) == list(donor) == ["ppt/slideMasters/slideMaster1.xml"]

    hx = host["ppt/slideMasters/slideMaster1.xml"]
    dx = donor["ppt/slideMasters/slideMaster1.xml"]
    ids = re.compile(r'<p:sldLayoutId id="(\d+)"')
    assert ids.findall(hx) == ids.findall(dx), "no p:sldLayoutId collision to re-mint"

    names = re.compile(r'<p:cSld name="([^"]*)"')
    assert names.findall(hx) == ["Relay Master"]
    assert names.findall(dx) == ["Relay Master (donor)"]

    spec = load_brand("relay")
    for path, want in ((template, spec.layout_name),
                       (variant_template, variant_spec(spec, "donor").layout_name)):
        prs = Presentation(str(path))
        assert [l.name for m in prs.slide_masters for l in m.slide_layouts] == [want]


def test_the_variant_ground_is_visibly_not_the_brands_own(template, variant_template):
    """The names are invisible to anyone looking at an exported PNG. The
    variant's ground is the brand's PANEL and the brand's own is BG_TOP, so
    the two masters can be told apart by eye as well as by XML - which the
    two fixture brands' pastel pairs could not manage."""
    style = load_brand("relay").style()

    def ground(path):
        with zipfile.ZipFile(str(path)) as z:
            layout = next(n for n in z.namelist()
                          if n.startswith("ppt/slideLayouts/") and n.endswith(".xml"))
            xml = z.read(layout).decode("utf-8")
        return sorted(set(re.findall(r'srgbClr val="([0-9A-Fa-f]{6})"', xml)))

    assert ground(template) == [str(style.BG_TOP)]
    assert ground(variant_template) == [str(style.PANEL)]
    assert ground(template) != ground(variant_template)


def test_inspect_checks_a_variant_against_the_variant_names(variant_template):
    """Without --variant the verifier is asked for the brand's own master and
    reports the variant as a failure, which is the honest answer."""
    ok = _run("scripts/inspect_template.py", "--brand", "relay",
              "--variant", "donor", str(variant_template))
    assert ok.returncode == 0, ok.stdout + ok.stderr
    assert "FAIL" not in ok.stdout
    bad = _run("scripts/inspect_template.py", "--brand", "relay", str(variant_template))
    assert bad.returncode != 0
    assert "FAIL" in bad.stdout and "Relay Master" in bad.stdout
