"""lint_deck.py - every test lints a deck this test built and saved."""

import subprocess
import sys
from pathlib import Path

import pytest

from deck_kit.brand import load_brand
from deck_kit.deck import add_slide, open_deck
from deck_kit.notes import Note, write_notes
from deck_kit.primitives import add_para, textbox

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "lint_deck.py"
sys.path.insert(0, str(ROOT / "scripts"))

import lint_deck  # noqa: E402


@pytest.fixture(scope="module")
def brand():
    return load_brand("relay")


@pytest.fixture
def deck(tmp_path, template, brand):
    """Three slides: a cover, a clean content slide, a content slide with every defect."""
    prs = open_deck(template, brand)
    s = brand.style()

    cover = add_slide(prs, brand)
    textbox(cover, 80, 80, 800, 60, "Walkthrough", 24, s.GREY, brand.sans)

    good = add_slide(prs, brand)
    textbox(good, 80, 40, 800, 30, "01. CONTEXT", 11, s.GREY, brand.sans)
    textbox(good, 80, 80, 800, 60, "Twelve people, one team.", 14, s.GREY, brand.sans)
    textbox(good, 80, 680, 800, 30, "(c) Test Organisation", 8, s.MUTED, brand.sans)

    bad = add_slide(prs, brand)
    textbox(bad, 80, 40, 800, 30, "02. FINDINGS", 11, s.GREY, brand.sans)
    textbox(bad, 80, 70, 800, 30, "03. NOT ALLOWED", 11, s.GREY, brand.sans)
    textbox(bad, 80, 120, 800, 60, "ASSESSEMENT of D2-07 for Acme client, if asked.", 14,
            s.GREY, brand.sans)

    write_notes(prs, [Note(1, "a", "x"), Note(2, "b", "Say this only if asked."),
                      Note(3, "c", "z")], font=brand.sans, size_pt=12, gap_pt=6)
    out = tmp_path / "lint_me.pptx"
    prs.save(str(out))
    return out


RULES = r"""
[typos]
needles = ["ASSESSEMENT"]

[codes]
patterns = ['\bD[123]-\d{2}\b']

[forbidden]
tokens = ["acme CLIENT"]

[required]
strings = ["(c) Test Organisation"]
exempt = [1]

[kickers]
pattern = '\b0[1-9]\. [A-Z][A-Z .&]{2,60}'
allowed = ["01. CONTEXT", "02. FINDINGS"]
exempt = [1]

[notes_only]
strings = ["if asked"]
"""


def _rules(tmp_path, text=RULES):
    p = tmp_path / "lint.toml"
    p.write_text(text, encoding="utf-8")
    return p


def _by_name(results):
    return {r.name: r for r in results}


def test_slide_texts_covers_every_slide_including_the_cover(deck):
    from pptx import Presentation
    texts = lint_deck.slide_texts(Presentation(str(deck)))
    assert sorted(texts) == [1, 2, 3]
    assert "Walkthrough" in texts[1]


def test_typos_are_reported_with_the_slide_and_the_needle(deck, tmp_path):
    r = _by_name(lint_deck.lint(deck, _rules(tmp_path)))["typos"]
    assert r.rules == 1
    assert [(f.position, f.detail) for f in r.findings] == [(3, "ASSESSEMENT")]


def test_codes_are_matched_as_regexes(deck, tmp_path):
    r = _by_name(lint_deck.lint(deck, _rules(tmp_path)))["codes"]
    assert [(f.position, f.detail) for f in r.findings] == [(3, "D2-07")]


def test_forbidden_tokens_match_case_insensitively(deck, tmp_path):
    r = _by_name(lint_deck.lint(deck, _rules(tmp_path)))["forbidden"]
    assert [(f.position, f.detail) for f in r.findings] == [(3, "acme CLIENT")]


def test_required_boilerplate_is_reported_missing_per_non_exempt_slide(deck, tmp_path):
    r = _by_name(lint_deck.lint(deck, _rules(tmp_path)))["required"]
    assert [(f.position, f.detail) for f in r.findings] == [(3, "(c) Test Organisation")]


def test_kickers_must_be_exactly_one_per_slide_and_from_the_allowed_list(deck, tmp_path):
    r = _by_name(lint_deck.lint(deck, _rules(tmp_path)))["kickers"]
    details = sorted((f.position, f.detail) for f in r.findings)
    assert details == [(3, "2 kickers on one slide: ['02. FINDINGS', '03. NOT ALLOWED']"),
                       (3, "kicker not allowed: '03. NOT ALLOWED'")]


# --- the kicker pattern and languages that are not English ----------------
# The pattern shipped in both examples and printed in references/05-qa.md was
# `[A-Z][A-Z .&]{2,60}`, which stops at the first accented capital. On a
# Spanish deck that is not a near miss: `01. DÓNDE ESTÁ LA PRESIÓN` matches
# nothing and is reported as "no kicker" - the report says the opposite of
# what is true - and `03. DISEÑO Y OPERACIÓN` matches the prefix `03. DISE`
# and reports a kicker that does not exist.

SHIPPED_PATTERN = r'\b0[1-9]\. [A-ZÁÉÍÓÚÜÑ][A-ZÁÉÍÓÚÜÑ .&]{2,60}'
ACCENTED = ["03. DISEÑO Y OPERACIÓN", "01. DÓNDE ESTÁ LA PRESIÓN"]


def test_an_accented_kicker_is_found_whole_and_allowed():
    """Whole: `03. DISEÑO Y OPERACIÓN`, not the `03. DISE` an ASCII class
    stops at. Allowed: the allowlist carries the real kicker, so a check
    that truncated the match would report it as not allowed."""
    texts = {1: "cover", 2: "03. DISEÑO Y OPERACIÓN\nSeis meses de trabajo"}
    rules = {"pattern": SHIPPED_PATTERN, "allowed": ACCENTED, "exempt": [1]}
    count, findings = lint_deck.check_kickers(texts, {}, rules, {})
    assert count == 2
    assert findings == []

    ascii_only = dict(rules, pattern=r'\b0[1-9]\. [A-Z][A-Z .&]{2,60}')
    _, broken = lint_deck.check_kickers(texts, {}, ascii_only, {})
    assert [f.detail for f in broken] == ["kicker not allowed: '03. DISE'"], \
        "the ASCII class must still be wrong, or this test proves nothing"


def test_the_accented_class_reports_a_missing_kicker_only_when_one_is_missing():
    """`01. DÓNDE ESTÁ LA PRESIÓN` used to match nothing at all, so the
    finding was "no kicker" on a slide that carried one."""
    texts = {2: "01. DÓNDE ESTÁ LA PRESIÓN\nEl coste de revisar a mano"}
    rules = {"pattern": SHIPPED_PATTERN, "allowed": ACCENTED}
    _, findings = lint_deck.check_kickers(texts, {}, rules, {})
    assert findings == []
    _, ascii_findings = lint_deck.check_kickers(
        texts, {}, dict(rules, pattern=r'\b0[1-9]\. [A-Z][A-Z .&]{2,60}'), {})
    assert [f.detail for f in ascii_findings] == ["no kicker"]


def test_the_shipped_kicker_pattern_accepts_accented_capitals():
    """Every file a user copies the pattern out of: both example rules files
    and the annotated block in the reference."""
    import re
    import tomllib

    for rules_file in (ROOT / "examples" / "walkthrough" / "lint.toml",
                       ROOT / "examples" / "ai-horizon" / "lint.toml",
                       ROOT / "templates" / "lint.toml"):
        with open(rules_file, "rb") as fh:
            pattern = tomllib.load(fh)["kickers"]["pattern"]
        for kicker in ACCENTED:
            assert re.search(pattern, kicker + "\nbody").group(0).strip() == kicker, \
                (rules_file, kicker)
    reference = (ROOT / "references" / "05-qa.md").read_text(encoding="utf-8")
    assert SHIPPED_PATTERN in reference


def test_notes_only_strings_must_be_in_the_notes_and_never_on_a_slide(deck, tmp_path):
    r = _by_name(lint_deck.lint(deck, _rules(tmp_path)))["notes_only"]
    assert [(f.position, f.detail) for f in r.findings] == [(3, "if asked")]
    absent = _rules(tmp_path, RULES.replace('strings = ["if asked"]', 'strings = ["never said"]'))
    r = _by_name(lint_deck.lint(deck, absent))["notes_only"]
    assert [(f.position, f.detail) for f in r.findings] == [(None, "never said: not in any note")]


def test_a_check_with_no_rules_is_skipped_not_passed(deck, tmp_path, capsys):
    results = lint_deck.lint(deck, _rules(tmp_path, "[typos]\nneedles = []\n"))
    by = _by_name(results)
    assert by["typos"].rules == 0 and by["typos"].findings == []
    code = lint_deck.report(results, expect_checks=1)
    out = capsys.readouterr().out
    assert code == 1
    assert "ran=0" in out and "expected 1 check(s) to run" in out


def test_the_clean_slide_alone_passes_with_every_check_run(tmp_path, template, brand):
    prs = open_deck(template, brand)
    s = brand.style()
    cover = add_slide(prs, brand)
    textbox(cover, 80, 80, 800, 60, "Walkthrough", 24, s.GREY, brand.sans)
    good = add_slide(prs, brand)
    textbox(good, 80, 40, 800, 30, "01. CONTEXT", 11, s.GREY, brand.sans)
    textbox(good, 80, 680, 800, 30, "(c) Test Organisation", 8, s.MUTED, brand.sans)
    write_notes(prs, [Note(1, "a", "x"), Note(2, "b", "only if asked")], font=brand.sans,
                size_pt=12, gap_pt=6)
    out = tmp_path / "clean.pptx"
    prs.save(str(out))
    proc = subprocess.run([sys.executable, str(SCRIPT), "--deck", str(out),
                           "--rules", str(_rules(tmp_path)), "--expect-checks", "6"],
                          capture_output=True, text=True)
    assert proc.returncode == 0, proc.stdout + proc.stderr
    assert "LINT PASSED ran=6" in proc.stdout   # skipped= is 2: exact and assets have no rules in RULES


def test_the_cli_exits_one_and_names_every_finding(deck, tmp_path):
    proc = subprocess.run([sys.executable, str(SCRIPT), "--deck", str(deck),
                           "--rules", str(_rules(tmp_path))],
                          capture_output=True, text=True)
    assert proc.returncode == 1
    assert "FAIL typos slide=3: ASSESSEMENT" in proc.stdout
    assert "FAIL codes slide=3: D2-07" in proc.stdout
    assert "check kickers: rules=2 findings=2" in proc.stdout


def test_a_scalar_where_a_rule_list_is_expected_raises(deck, tmp_path):
    bad = _rules(tmp_path, RULES.replace('needles = ["ASSESSEMENT"]', 'needles = "ASSESSEMENT"'))
    with pytest.raises(ValueError, match="needles must be a list"):
        lint_deck.lint(deck, bad)


def test_a_list_where_the_kicker_pattern_belongs_raises(deck, tmp_path):
    """`pattern` was the one scalar read with a bare `rules.get`: a list
    reached `re.compile` and came back out as `TypeError: unhashable type:
    'list'`, against references/05-qa.md's promise that no bad input
    produces a traceback."""
    bad = _rules(tmp_path, '[kickers]\npattern = ["a"]\nallowed = ["x"]\n')
    with pytest.raises(ValueError, match="pattern must be a"):
        lint_deck.lint(deck, bad)


def test_forbidden_matches_a_token_split_across_paragraphs(tmp_path, template, brand):
    prs = open_deck(template, brand)
    s = brand.style()
    cover = add_slide(prs, brand)
    tb = textbox(cover, 80, 80, 800, 120, "Acme", 14, s.GREY, brand.sans)
    add_para(tb.text_frame, "client", 14, s.GREY, brand.sans)
    write_notes(prs, [Note(1, "a", "x")], font=brand.sans, size_pt=12, gap_pt=6)
    out = tmp_path / "split.pptx"
    prs.save(str(out))
    rules = _rules(tmp_path, '[forbidden]\ntokens = ["acme client"]\n')
    r = _by_name(lint_deck.lint(out, rules))["forbidden"]
    assert [(f.position, f.detail) for f in r.findings] == [(1, "acme client")]


def test_the_cli_fails_cleanly_on_a_missing_rules_file(deck, tmp_path):
    proc = subprocess.run([sys.executable, str(SCRIPT), "--deck", str(deck),
                           "--rules", str(tmp_path / "does_not_exist.toml")],
                          capture_output=True, text=True)
    assert proc.returncode == 1
    assert "FAIL:" in proc.stdout


def test_the_cli_fails_cleanly_on_a_missing_deck(tmp_path):
    """python-pptx raises PackageNotFoundError, not OSError, on a missing deck.

    Without it in main's except tuple a mistyped --deck path comes out as a
    raw traceback - the one bad input that did not produce the FAIL: line
    every other one does, and that references/05-qa.md says it produces.
    """
    proc = subprocess.run([sys.executable, str(SCRIPT),
                           "--deck", str(tmp_path / "nope.pptx"),
                           "--rules", str(_rules(tmp_path))],
                          capture_output=True, text=True)
    assert proc.returncode == 1
    assert "FAIL:" in proc.stdout


def test_the_cli_fails_cleanly_on_an_invalid_regex(deck, tmp_path):
    bad = _rules(tmp_path, "[codes]\npatterns = ['(']\n")
    proc = subprocess.run([sys.executable, str(SCRIPT), "--deck", str(deck),
                           "--rules", str(bad)],
                          capture_output=True, text=True)
    assert proc.returncode == 1
    assert "FAIL: invalid regex in codes" in proc.stdout


import hashlib

SPEC_WITH_FACTS = """# SPEC - test

## 3. Verified facts

| Fact | Exact string | Source |
|---|---|---|
| Team size | `Twelve people, one team.` | interview notes |
| Missing figure | `Forty percent of nothing` | nowhere |
"""


def test_exact_strings_come_from_the_spec_table(tmp_path):
    spec = tmp_path / "spec.md"
    spec.write_text(SPEC_WITH_FACTS, encoding="utf-8")
    assert lint_deck.exact_strings_from_spec(spec) == [
        "Twelve people, one team.", "Forty percent of nothing"]
    spec.write_text("# no table here\n", encoding="utf-8")
    with pytest.raises(ValueError, match="Exact string"):
        lint_deck.exact_strings_from_spec(spec)


def test_exact_figures_must_appear_verbatim_on_some_slide(deck, tmp_path):
    (tmp_path / "spec.md").write_text(SPEC_WITH_FACTS, encoding="utf-8")
    rules = _rules(tmp_path, RULES + '\n[exact]\nspec = "spec.md"\n')
    r = _by_name(lint_deck.lint(deck, rules))["exact"]
    assert r.rules == 2
    assert [(f.position, f.detail) for f in r.findings] == [
        (None, "not found verbatim on any slide: 'Forty percent of nothing'")]


def _brand_dir(tmp_path, monkeypatch, manifest=None):
    root = tmp_path / "brands"
    b = root / "fake"
    b.mkdir(parents=True)
    (b / "logo.png").write_bytes(b"\x89PNG fake")
    (b / "brand.py").write_text("BRAND = None\n")
    if manifest is not None:
        (b / "assets.toml").write_text(manifest, encoding="utf-8")
    monkeypatch.setenv("DECKWRIGHT_BRANDS", str(root))
    return b


def _manifest_for(b, provenance="drawn for this test"):
    digest = hashlib.sha256((b / "logo.png").read_bytes()).hexdigest()
    return '[[asset]]\npath = "logo.png"\nsha256 = "%s"\nprovenance = "%s"\n' % (digest, provenance)


def test_assets_check_passes_when_every_binary_is_listed_with_provenance(deck, tmp_path, monkeypatch):
    b = _brand_dir(tmp_path, monkeypatch)
    (b / "assets.toml").write_text(_manifest_for(b), encoding="utf-8")
    rules = _rules(tmp_path, '[assets]\nbrand = "fake"\n')
    r = _by_name(lint_deck.lint(deck, rules))["assets"]
    assert r.rules == 1 and r.findings == []


def test_assets_check_fails_on_an_unlisted_binary(deck, tmp_path, monkeypatch):
    b = _brand_dir(tmp_path, monkeypatch)
    (b / "assets.toml").write_text(_manifest_for(b), encoding="utf-8")
    (b / "stray.png").write_bytes(b"\x89PNG stray")
    rules = _rules(tmp_path, '[assets]\nbrand = "fake"\n')
    r = _by_name(lint_deck.lint(deck, rules))["assets"]
    assert [f.detail for f in r.findings] == ["not in assets.toml: stray.png"]


def test_assets_check_is_not_evaded_by_an_unknown_or_missing_extension(deck, tmp_path, monkeypatch):
    """An extension-keyed check lets the person adding the file choose its
    blind spot. Anything that is not text needs provenance."""
    b = _brand_dir(tmp_path, monkeypatch)
    (b / "assets.toml").write_text(_manifest_for(b), encoding="utf-8")
    (b / "mark.webp").write_bytes(b"RIFF fake")
    (b / "noext").write_bytes(b"\x00\x01")
    (b / "notes.md").write_text("text is fine\n")
    (b / "__pycache__").mkdir()
    (b / "__pycache__" / "brand.cpython-311.pyc").write_bytes(b"\x00")
    rules = _rules(tmp_path, '[assets]\nbrand = "fake"\n')
    details = sorted(f.detail for f in _by_name(lint_deck.lint(deck, rules))["assets"].findings)
    assert details == ["not in assets.toml: mark.webp", "not in assets.toml: noext"]


def test_assets_check_fails_on_a_changed_hash_a_missing_file_and_empty_provenance(deck, tmp_path, monkeypatch):
    b = _brand_dir(tmp_path, monkeypatch)
    manifest = _manifest_for(b, provenance="") + \
        '\n[[asset]]\npath = "gone.png"\nsha256 = "00"\nprovenance = "x"\n'
    (b / "assets.toml").write_text(manifest, encoding="utf-8")
    (b / "logo.png").write_bytes(b"\x89PNG changed")
    rules = _rules(tmp_path, '[assets]\nbrand = "fake"\n')
    details = sorted(f.detail for f in _by_name(lint_deck.lint(deck, rules))["assets"].findings)
    assert details == ["empty provenance: logo.png", "listed but missing: gone.png",
                       "sha256 mismatch: logo.png"]


def test_assets_check_fails_when_the_manifest_itself_is_missing(deck, tmp_path, monkeypatch):
    _brand_dir(tmp_path, monkeypatch)
    rules = _rules(tmp_path, '[assets]\nbrand = "fake"\n')
    r = _by_name(lint_deck.lint(deck, rules))["assets"]
    assert r.rules == 1
    assert [f.detail for f in r.findings] == ["no assets.toml in fake"]


def test_the_manifest_subcommand_lists_every_binary_with_its_hash(tmp_path, monkeypatch):
    b = _brand_dir(tmp_path, monkeypatch)
    proc = subprocess.run([sys.executable, str(SCRIPT), "manifest", "--brand", "fake"],
                          capture_output=True, text=True)
    assert proc.returncode == 0, proc.stderr
    digest = hashlib.sha256((b / "logo.png").read_bytes()).hexdigest()
    assert 'path = "logo.png"' in proc.stdout and digest in proc.stdout
    assert 'provenance = ""' in proc.stdout
    assert "brand.py" not in proc.stdout


def test_the_shipped_brand_passes_its_own_asset_check(deck, tmp_path):
    """brands/relay is the repository's only brand and its assets.toml has to
    account for every binary in it: six faces, six icons and two wordmarks."""
    rules = _rules(tmp_path, '[assets]\nbrand = "relay"\n')
    r = _by_name(lint_deck.lint(deck, rules))["assets"]
    assert r.findings == [], r.findings


def test_a_manifest_entry_without_a_path_fails_cleanly(deck, tmp_path, monkeypatch):
    _brand_dir(tmp_path, monkeypatch)
    b = tmp_path / "brands" / "fake"
    (b / "assets.toml").write_text('[[asset]]\nsha256 = "00"\nprovenance = "x"\n', encoding="utf-8")
    rules = _rules(tmp_path, '[assets]\nbrand = "fake"\n')
    with pytest.raises(ValueError, match="needs string path"):
        lint_deck.lint(deck, rules)


def test_a_duplicated_manifest_path_is_refused(deck, tmp_path, monkeypatch):
    b = _brand_dir(tmp_path, monkeypatch)
    good = _manifest_for(b)
    (b / "assets.toml").write_text(good + good, encoding="utf-8")
    rules = _rules(tmp_path, '[assets]\nbrand = "fake"\n')
    proc = subprocess.run([sys.executable, str(SCRIPT), "--deck", str(deck), "--rules", str(rules)],
                          capture_output=True, text=True)
    assert proc.returncode == 1
    assert "FAIL:" in proc.stdout and "duplicate path" in proc.stdout


# --- --deck as a directory of HTML drafts ---------------------------------
# An HTML deck is a deliverable in its own right (the spec's **Output:**
# line), and a deliverable nobody lints is a deliverable nobody checked. The
# same rules file governs both outputs: the content constraints belong to the
# deck, not to the file format it ships as.

DRAFT_RULES = r"""
[typos]
needles = ["ASSESSEMENT"]

[codes]
patterns = ['\bD[123]-\d{2}\b']

[forbidden]
tokens = ["Acme Client"]

[required]
strings = ["(C) EXAMPLE - CONFIDENTIAL"]
exempt = [1]

[kickers]
pattern = '\b0[1-9]\. [A-ZÁÉÍÓÚÜÑ][A-ZÁÉÍÓÚÜÑ .&]{2,60}'
allowed = ["01. THE LOOP", "02. DISEÑO Y OPERACIÓN"]
exempt = [1]

[notes_only]
strings = ["If asked"]

[exact]
strings = ["Four stages and one loop."]
"""

DRAFT_COVER = """<!doctype html>
<!-- DISPOSABLE: the builder is the source of truth. -->
<html lang="en"><head><meta charset="utf-8"><title>S01 Cover</title>
<link rel="stylesheet" href="slide.css"></head><body>
<div class="slide"><div class="disp">A cover</div></div>
</body></html>
"""

DRAFT_BODY = """<!doctype html>
<html lang="en"><head><meta charset="utf-8"><title>S02 The loop</title>
<link rel="stylesheet" href="slide.css"></head><body>
<div class="slide">
  <div class="kicker">01. The loop</div>
  <div class="title">Four stages and one loop.</div>
  <div class="foot"><span class="meta">(C) EXAMPLE - CONFIDENTIAL</span></div>
</div>
</body></html>
"""

DRAFT_SPANISH = """<!doctype html>
<html lang="es"><head><meta charset="utf-8"><title>S03</title></head><body>
<div class="slide">
  <div class="kicker">02. Diseno y operacion</div>
  <div class="foot"><span class="meta">(C) EXAMPLE - CONFIDENTIAL</span></div>
</div>
</body></html>
""".replace("Diseno y operacion", "Diseño y operación")


@pytest.fixture
def drafts(tmp_path):
    d = tmp_path / "drafts"
    d.mkdir()
    (d / "s01-cover.html").write_text(DRAFT_COVER, encoding="utf-8")
    (d / "s02-loop.html").write_text(DRAFT_BODY, encoding="utf-8")
    (d / "s03-es.html").write_text(DRAFT_SPANISH, encoding="utf-8")
    (d / "slide.css").write_text(".slide{}", encoding="utf-8")
    return d


def test_a_drafts_directory_lints_one_slide_per_file_in_name_order(drafts):
    texts = lint_deck.draft_texts(drafts)
    assert sorted(texts) == [1, 2, 3]
    assert texts[1] == "A cover"
    assert texts[2].splitlines() == ["01. The loop", "Four stages and one loop.",
                                     "(C) EXAMPLE - CONFIDENTIAL"]


def test_the_harvest_reads_no_title_no_stylesheet_and_no_comment(drafts):
    text = lint_deck.draft_text(drafts / "s01-cover.html")
    assert "S01 Cover" not in text          # <title> is chrome
    assert "slide.css" not in text          # a <link> carries no text
    assert "DISPOSABLE" not in text         # a comment is not on the slide


def test_every_check_but_notes_only_runs_on_a_drafts_directory(drafts, tmp_path):
    rules = _rules(tmp_path, DRAFT_RULES)
    by = _by_name(lint_deck.lint(drafts, rules))
    assert by["notes_only"].rules == 0
    assert by["notes_only"].skip_reason == "a drafts directory carries no speaker notes"
    for name in ("typos", "codes", "forbidden", "required", "kickers", "exact"):
        assert by[name].rules > 0, name
        assert by[name].findings == [], (name, by[name].findings)


def test_the_drafts_report_says_why_notes_only_was_skipped(drafts, tmp_path, capsys):
    results = lint_deck.lint(drafts, _rules(tmp_path, DRAFT_RULES))
    assert lint_deck.report(results, expect_checks=6) == 0
    out = capsys.readouterr().out
    assert "check notes_only: skipped (a drafts directory carries no speaker notes)" in out
    assert "LINT PASSED ran=6 skipped=2" in out     # assets has no rules here


def test_a_draft_kicker_uppercased_by_css_is_still_matched(drafts, tmp_path):
    """`<div class="kicker">01. The loop</div>` reads `01. THE LOOP` on the
    slide, and that is the string the builder writes into the PPTX. The
    markup carries the source case, so [kickers] - and only [kickers] -
    folds case on a drafts directory."""
    by = _by_name(lint_deck.lint(drafts, _rules(tmp_path, DRAFT_RULES)))
    assert by["kickers"].findings == []
    # The accented Spanish kicker travels the same road: found whole, and
    # allowed, which an ASCII-only class could do neither of.
    assert "ño y operación" in lint_deck.draft_texts(drafts)[3]


def test_required_and_exact_do_not_fold_case_on_drafts(drafts, tmp_path):
    """They compare strings the spec fixed. A draft that writes its
    boilerplate in the wrong case is a draft that says the wrong thing."""
    lower = DRAFT_BODY.replace("(C) EXAMPLE - CONFIDENTIAL", "(c) example - confidential")
    (drafts / "s02-loop.html").write_text(lower, encoding="utf-8")
    by = _by_name(lint_deck.lint(drafts, _rules(tmp_path, DRAFT_RULES)))
    assert [(f.position, f.detail) for f in by["required"].findings] == \
        [(2, "(C) EXAMPLE - CONFIDENTIAL")]


def test_a_forbidden_token_on_a_draft_is_found(drafts, tmp_path):
    (drafts / "s02-loop.html").write_text(
        DRAFT_BODY.replace("Four stages and one loop.", "Written for Acme Client."),
        encoding="utf-8")
    by = _by_name(lint_deck.lint(drafts, _rules(tmp_path, DRAFT_RULES)))
    assert [(f.position, f.detail) for f in by["forbidden"].findings] == [(2, "Acme Client")]


def test_the_cli_lints_a_drafts_directory(drafts, tmp_path):
    rules = _rules(tmp_path, DRAFT_RULES)
    proc = subprocess.run([sys.executable, str(SCRIPT), "--deck", str(drafts),
                           "--rules", str(rules), "--expect-checks", "6"],
                          capture_output=True, text=True, encoding="utf-8")
    assert proc.returncode == 0, proc.stdout + proc.stderr
    assert "LINT PASSED ran=6 skipped=2" in proc.stdout


def test_the_cli_fails_cleanly_on_a_directory_with_no_drafts(tmp_path):
    empty = tmp_path / "empty"
    empty.mkdir()
    rules = _rules(tmp_path, DRAFT_RULES)
    proc = subprocess.run([sys.executable, str(SCRIPT), "--deck", str(empty),
                           "--rules", str(rules)],
                          capture_output=True, text=True, encoding="utf-8")
    assert proc.returncode == 1
    assert "FAIL: no sNN-*.html drafts" in proc.stdout
    assert "Traceback" not in proc.stdout + proc.stderr


def test_the_repositorys_own_drafts_lint_against_the_decks_own_rules():
    """The ai-horizon drafts and the ai-horizon rules file: the same seven
    checks the HTML run.ps1 asserts, on the artifacts the gate ships. One
    rules file for both outputs - the content constraints belong to the deck
    and not to the file format it ships as."""
    horizon = ROOT / "examples" / "ai-horizon"
    results = lint_deck.lint(horizon / "drafts", horizon / "lint.toml")
    by = _by_name(results)
    assert by["notes_only"].rules == 0 and by["notes_only"].skip_reason
    assert all(r.findings == [] for r in results), \
        [(r.name, r.findings) for r in results if r.findings]
    assert sum(1 for r in results if r.rules) == 7


# --- [forbidden] tokens_file: the half the check could not carry ----------

def test_forbidden_reads_tokens_from_a_file_outside_the_deck(deck, tmp_path):
    """A rules file travels with the deck, so listing a client's own name in
    it writes that name into the one file it must not be in. `tokens_file`
    points somewhere the deck directory does not have to carry."""
    secret = tmp_path / "not-in-the-deck.txt"
    secret.write_text("# one per line\nAcme Client\n\nClient Bank\n", encoding="utf-8")
    rules = _rules(tmp_path, '[forbidden]\ntokens = []\ntokens_file = %r\n' % str(secret))
    r = _by_name(lint_deck.lint(deck, rules))["forbidden"]
    assert r.rules == 2
    assert [(f.position, f.detail) for f in r.findings] == [(3, "Acme Client")]


def test_a_missing_tokens_file_fails_cleanly(deck, tmp_path):
    rules = _rules(tmp_path, '[forbidden]\ntokens_file = "nowhere.txt"\n')
    proc = subprocess.run([sys.executable, str(SCRIPT), "--deck", str(deck),
                           "--rules", str(rules)],
                          capture_output=True, text=True, encoding="utf-8")
    assert proc.returncode == 1
    assert proc.stdout.startswith("FAIL:")
    assert "Traceback" not in proc.stdout + proc.stderr


# --- --brand with no --deck: the assets check on its own ------------------

def test_the_assets_check_runs_against_a_brand_with_no_deck():
    """"Set up a brand" step 4 asks for this two stages before a deck
    exists. --deck and --rules used to be required, so it could not be run."""
    proc = subprocess.run([sys.executable, str(SCRIPT), "--brand", "relay"],
                          capture_output=True, text=True, encoding="utf-8")
    assert proc.returncode == 0, proc.stdout + proc.stderr
    assert "check assets: rules=1 findings=0" in proc.stdout
    assert "LINT PASSED ran=1 skipped=0" in proc.stdout


def test_neither_a_deck_nor_a_brand_is_a_clean_refusal():
    proc = subprocess.run([sys.executable, str(SCRIPT), "--rules", "x.toml"],
                          capture_output=True, text=True, encoding="utf-8")
    assert proc.returncode == 1
    assert "--deck and --rules are both required" in proc.stdout
