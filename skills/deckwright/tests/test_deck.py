import pytest
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from deck_kit.brand import load_brand
from deck_kit.deck import add_slide, delete_positions, open_deck


@pytest.fixture(scope="module")
def brand():
    return load_brand("relay")


def test_open_deck_returns_a_presentation_with_no_slides(template):
    assert len(open_deck(template, load_brand("relay")).slides) == 0


def test_open_deck_keeps_masters_and_layouts(template):
    spec = load_brand("relay")
    prs = open_deck(template, spec)
    assert spec.master_name in [m.name for m in prs.slide_masters]


def test_add_slide_uses_the_layout_named_by_the_brand(template):
    spec = load_brand("relay")
    prs = open_deck(template, spec)
    s = add_slide(prs, spec)
    assert s.slide_layout.name == spec.layout_name
    assert len(prs.slides) == 1


def test_add_slide_strips_inherited_placeholders(template):
    spec = load_brand("relay")
    s = add_slide(open_deck(template, spec), spec)
    assert list(s.placeholders) == []


def test_add_slide_purges_the_slide_number_placeholder_by_default(template):
    """The content layout carries a real SLIDE_NUMBER placeholder (Task 7),
    holding an <a:fld type="slidenum">. The default purge must still remove
    it - a builder that calls add_slide() without keep_placeholders must see
    exactly what it saw before that placeholder existed."""
    spec = load_brand("relay")
    s = add_slide(open_deck(template, spec), spec)
    types = [ph.placeholder_format.type for ph in s.placeholders]
    assert PP_PLACEHOLDER.SLIDE_NUMBER not in types
    assert list(s.placeholders) == []


def test_add_slide_keeps_the_slide_number_placeholder_when_asked(template):
    spec = load_brand("relay")
    s = add_slide(open_deck(template, spec), spec,
                  keep_placeholders=(PP_PLACEHOLDER.SLIDE_NUMBER,))
    types = [ph.placeholder_format.type for ph in s.placeholders]
    assert types == [PP_PLACEHOLDER.SLIDE_NUMBER]


def test_add_slide_raises_when_the_layout_name_is_absent(template):
    spec = load_brand("relay")
    spec.layout_name = "No Such Layout"
    with pytest.raises(LookupError, match="No Such Layout"):
        add_slide(open_deck(template, spec), spec)


def _ids(prs):
    rt = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    return [e.get(rt) for e in prs.slides._sldIdLst]


def test_delete_positions_removes_exactly_the_positions_named(template):
    """A count assertion alone passes even when the wrong two slides go."""
    spec = load_brand("relay")
    prs = open_deck(template, spec)
    for _ in range(4):
        add_slide(prs, spec)
    before = _ids(prs)
    delete_positions(prs, [2, 4])
    assert _ids(prs) == [before[0], before[2]]


def test_delete_positions_rejects_zero_rather_than_deleting_the_last_slide(template):
    """ids[pos - 1] with pos=0 is ids[-1]: the obvious 1-based slip silently
    destroys the wrong slide."""
    spec = load_brand("relay")
    prs = open_deck(template, spec)
    for _ in range(3):
        add_slide(prs, spec)
    with pytest.raises(IndexError, match="outside 1..3"):
        delete_positions(prs, [0])
    assert len(prs.slides) == 3


def test_delete_positions_rejects_out_of_range(template):
    spec = load_brand("relay")
    prs = open_deck(template, spec)
    add_slide(prs, spec)
    with pytest.raises(IndexError, match="outside 1..1"):
        delete_positions(prs, [5])


def test_delete_positions_tolerates_a_repeated_position(template):
    spec = load_brand("relay")
    prs = open_deck(template, spec)
    for _ in range(3):
        add_slide(prs, spec)
    delete_positions(prs, [2, 2])
    assert len(prs.slides) == 2


def test_open_deck_rejects_a_template_that_lacks_the_brand_layout(template):
    """Fail before any content is built, not at the first add_slide."""
    spec = load_brand("relay")
    spec.layout_name = "No Such Layout"
    with pytest.raises(LookupError, match="No Such Layout"):
        open_deck(template, spec)


def test_deleting_slides_leaves_a_file_that_reopens(template, tmp_path):
    spec = load_brand("relay")
    prs = open_deck(template, spec)
    for _ in range(3):
        add_slide(prs, spec)
    delete_positions(prs, [2])
    out = tmp_path / "after.pptx"
    prs.save(str(out))
    assert len(Presentation(str(out)).slides) == 2


def test_duplicate_slide_does_not_reuse_a_freed_part_name(tmp_path, template, brand):
    """Catalogue item 9: add_slide() picks the next slideN number naively, so
    after a deletion it reuses an occupied part name and one slide overwrites
    another in the package."""
    from deck_kit import merge
    from deck_kit.deck import duplicate_slide
    from deck_kit.primitives import textbox

    prs = open_deck(template, brand)
    style = brand.style()
    for label in ("A", "B", "C"):
        slide = add_slide(prs, brand)
        textbox(slide, 80, 80, 400, 40, label, 14, style.GREY, brand.sans)
    delete_positions(prs, [2])                 # frees slide2.xml
    duplicate_slide(prs, 0)                    # must not claim slide2.xml
    out = tmp_path / "dup.pptx"
    prs.save(str(out))

    parts = merge.read_package(out)
    order = merge.slide_order(parts)
    assert len(order) == len(set(order)) == 3
    texts = [s.shapes[0].text_frame.text for s in Presentation(str(out)).slides]
    assert texts == ["A", "A", "C"]
    assert merge.dangling_refs(parts) == []


def test_add_slide_does_not_reuse_a_freed_part_name(tmp_path, template, brand):
    """Catalogue item 9 names add_slide(), not duplicate_slide: the same
    part-naming defect lives in the function every builder calls, and fixing
    only duplicate_slide would leave it live there."""
    from deck_kit import merge
    from deck_kit.primitives import textbox

    prs = open_deck(template, brand)
    style = brand.style()
    for label in ("A", "B", "C"):
        slide = add_slide(prs, brand)
        textbox(slide, 80, 80, 400, 40, label, 14, style.GREY, brand.sans)
    delete_positions(prs, [2])                 # frees slide2.xml
    slide = add_slide(prs, brand)               # must not claim slide2.xml
    textbox(slide, 80, 80, 400, 40, "D", 14, style.GREY, brand.sans)
    out = tmp_path / "add.pptx"
    prs.save(str(out))

    parts = merge.read_package(out)
    order = merge.slide_order(parts)
    assert len(order) == len(set(order)) == 3
    texts = [s.shapes[0].text_frame.text for s in Presentation(str(out)).slides]
    assert texts == ["A", "C", "D"]
    assert merge.dangling_refs(parts) == []


def test_assert_fingerprint_accepts_the_file_it_describes(tmp_path, template, brand):
    from deck_kit.deck import assert_fingerprint
    from deck_kit.primitives import textbox

    prs = open_deck(template, brand)
    style = brand.style()
    for label in ("Overview", "Findings", "Roadmap"):
        slide = add_slide(prs, brand)
        textbox(slide, 80, 80, 600, 40, label, 20, style.GREY, brand.sans)
    assert_fingerprint(prs, 3, {2: "Find", 3: "Road"})


def test_assert_fingerprint_rejects_a_deck_whose_slide_moved(tmp_path, template, brand):
    """The author reorders slides between versions. A build that transplants
    at position 2 without checking what is there replaces the wrong slide."""
    from deck_kit.deck import assert_fingerprint
    from deck_kit.primitives import textbox

    prs = open_deck(template, brand)
    style = brand.style()
    for label in ("Overview", "Roadmap", "Findings"):
        slide = add_slide(prs, brand)
        textbox(slide, 80, 80, 600, 40, label, 20, style.GREY, brand.sans)
    with pytest.raises(AssertionError) as e:
        assert_fingerprint(prs, 3, {2: "Find", 3: "Road"})
    assert "position 2" in str(e.value)


def test_assert_fingerprint_rejects_a_wrong_slide_count(tmp_path, template, brand):
    from deck_kit.deck import assert_fingerprint
    from deck_kit.primitives import textbox

    prs = open_deck(template, brand)
    style = brand.style()
    slide = add_slide(prs, brand)
    textbox(slide, 80, 80, 600, 40, "Only", 20, style.GREY, brand.sans)
    with pytest.raises(AssertionError):
        assert_fingerprint(prs, 3, {})
