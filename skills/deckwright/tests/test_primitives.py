import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt

from deck_kit import primitives as P
from deck_kit.geometry import px

BLUE = RGBColor(0x25, 0x63, 0xEB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SANS = "Noto Sans"
LOGO = "brands/relay/logo_primary.png"
LOGO2 = "brands/relay/logo_inverse.png"
WHITE_RGBA = (255, 255, 255, 255)


@pytest.fixture
def slide():
    prs = Presentation()
    return prs.slides.add_slide(prs.slide_layouts[6])


def test_rrect_places_geometry_in_css_pixels(slide):
    sp = P.rrect(slide, 48, 100, 200, 30, fill=WHITE)
    assert (sp.left, sp.top, sp.width, sp.height) == (px(48), px(100), px(200), px(30))


def test_rrect_radius_zero_yields_a_plain_rectangle(slide):
    """Both are autoshapes; the preset geometry is what differs."""
    sharp = P.rrect(slide, 0, 0, 10, 10, fill=WHITE, rad=0)
    round_ = P.rrect(slide, 0, 0, 10, 10, fill=WHITE, rad=8)
    assert sharp.auto_shape_type == MSO_SHAPE.RECTANGLE
    assert round_.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE


def test_settext_writes_runs_with_size_font_and_colour(slide):
    sp = P.rrect(slide, 0, 0, 100, 20, fill=WHITE)
    P.settext(sp, [("Label", True)], 9.5, BLUE, SANS)
    run = sp.text_frame.paragraphs[0].runs[0]
    assert run.text == "Label"
    assert run.font.bold is True
    assert run.font.size == Pt(9.5)
    assert run.font.name == SANS
    assert run.font.color.rgb == BLUE


def test_settext_accepts_a_bare_string(slide):
    sp = P.rrect(slide, 0, 0, 100, 20, fill=WHITE)
    P.settext(sp, "plain", 9.0, BLUE, SANS)
    assert sp.text_frame.paragraphs[0].runs[0].text == "plain"


def test_settext_honours_a_per_run_colour_override(slide):
    sp = P.rrect(slide, 0, 0, 100, 20, fill=WHITE)
    P.settext(sp, [("a", False), ("b", True, WHITE)], 9.0, BLUE, SANS)
    runs = sp.text_frame.paragraphs[0].runs
    assert runs[0].font.color.rgb == BLUE
    assert runs[1].font.color.rgb == WHITE


def test_settext_uppercases_when_asked(slide):
    sp = P.rrect(slide, 0, 0, 100, 20, fill=WHITE)
    P.settext(sp, [("eyebrow", True)], 8.0, BLUE, SANS, upper=True)
    assert sp.text_frame.paragraphs[0].runs[0].text == "EYEBROW"


def test_add_para_appends_a_second_paragraph(slide):
    tb = P.textbox(slide, 0, 0, 200, 40, [("title", True)], 9.5, BLUE, SANS)
    P.add_para(tb.text_frame, [("detail", False)], 8.5, BLUE, SANS)
    assert len(tb.text_frame.paragraphs) == 2
    assert tb.text_frame.paragraphs[1].runs[0].text == "detail"


def test_settext_only_ever_writes_the_first_paragraph(slide):
    """This is why add_para exists; the behaviour is deliberate."""
    tb = P.textbox(slide, 0, 0, 200, 40, [("one", False)], 9.0, BLUE, SANS)
    P.add_para(tb.text_frame, [("two", False)], 9.0, BLUE, SANS)
    P.settext(tb, [("rewritten", False)], 9.0, BLUE, SANS)
    assert tb.text_frame.paragraphs[0].runs[-1].text == "rewritten"
    assert tb.text_frame.paragraphs[1].runs[0].text == "two"


def test_ensure_balls_renders_one_png_per_fraction(tmp_path):
    P.ensure_balls(tmp_path, [BLUE], ring=(0xB8, 0xBD, 0xC4, 255), body=WHITE_RGBA)
    names = sorted(p.name for p in tmp_path.iterdir() if p.suffix == ".png")
    assert names == ["ball_0_2563EB.png", "ball_100_2563EB.png",
                     "ball_25_2563EB.png", "ball_50_2563EB.png",
                     "ball_75_2563EB.png"]


def test_ensure_balls_regenerates_when_the_colour_changes(tmp_path):
    """The engagement's cache keyed on file existence, so a colour change
    silently produced stale assets."""
    P.ensure_balls(tmp_path, [BLUE], ring=(0, 0, 0, 255), body=WHITE_RGBA)
    before = (tmp_path / "ball_50_2563EB.png").read_bytes()
    P.ensure_balls(tmp_path, [BLUE], ring=(255, 0, 0, 255), body=WHITE_RGBA)
    assert (tmp_path / "ball_50_2563EB.png").read_bytes() != before


def test_ensure_balls_does_not_rerender_when_nothing_changed(tmp_path):
    """Assert on the PNG, not on the stamp: the stamp is a hash of the
    parameters, so it is identical whether or not the cache short-circuited.
    Only the file mtime shows whether the render was skipped."""
    import os
    import time
    P.ensure_balls(tmp_path, [BLUE], ring=(0, 0, 0, 255), body=WHITE_RGBA)
    png = tmp_path / "ball_50_2563EB.png"
    before = os.stat(png).st_mtime_ns
    time.sleep(0.01)
    P.ensure_balls(tmp_path, [BLUE], ring=(0, 0, 0, 255), body=WHITE_RGBA)
    assert os.stat(png).st_mtime_ns == before


def test_ensure_balls_rerenders_when_the_body_colour_changes(tmp_path):
    P.ensure_balls(tmp_path, [BLUE], ring=(0, 0, 0, 255), body=WHITE_RGBA)
    before = (tmp_path / "ball_25_2563EB.png").read_bytes()
    P.ensure_balls(tmp_path, [BLUE], ring=(0, 0, 0, 255), body=(0, 255, 0, 255))
    assert (tmp_path / "ball_25_2563EB.png").read_bytes() != before


def test_ball_raises_before_the_assets_are_rendered(slide, tmp_path):
    with pytest.raises(FileNotFoundError, match="ensure_balls"):
        P.ball(slide, tmp_path, 0, 0, 18, 50, BLUE)


def test_icon_raises_when_the_file_is_missing(slide, tmp_path):
    """The engagement's icon() failed silently: a typo rendered nothing and
    the build still succeeded."""
    with pytest.raises(FileNotFoundError, match="nope"):
        P.icon(slide, tmp_path, "nope", "blue", 0, 0, 18)


def test_lockup_places_two_pictures_and_a_divider(slide):
    before = len(slide.shapes)
    P.lockup(slide, LOGO, LOGO2, 1232, 32, BLUE)
    assert len(slide.shapes) - before == 3


def test_footer_draws_a_rule_a_page_number_and_the_left_text(slide):
    before = len(slide.shapes)
    P.footer(slide, "Deck title", 7, SANS, BLUE, BLUE, 8.0, width=1184, left=48, y=690)
    assert len(slide.shapes) - before >= 3
    texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
    assert "Deck title" in texts
    assert "7" in texts


# --- the doctrine, checked where it actually lands ------------------------

def test_settext_turns_autofit_off(slide):
    """python-pptx's add_textbox emits <a:spAutoFit/>, so every text box
    arrives with "resize shape to fit text" ON. Measure-then-place is
    defeated the moment PowerPoint is allowed to resize a box."""
    from pptx.oxml.ns import qn
    tb = P.textbox(slide, 0, 0, 100, 20, "text", 9.0, BLUE, SANS)
    bodyPr = tb.text_frame._txBody.bodyPr
    assert bodyPr.find(qn("a:spAutoFit")) is None
    assert bodyPr.find(qn("a:normAutofit")) is None
    assert bodyPr.find(qn("a:noAutofit")) is not None


def test_settext_clears_line_breaks_as_well_as_runs(slide):
    """Runs alone is not enough: a paragraph from a deck someone else
    authored routinely carries <a:br>, and leaving it makes the
    rewrite-the-first-paragraph contract quietly false."""
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import nsdecls, qn
    tb = P.textbox(slide, 0, 0, 100, 20, "before", 9.0, BLUE, SANS)
    para = tb.text_frame.paragraphs[0]._p
    para.append(parse_xml("<a:br %s/>" % nsdecls("a")))
    P.settext(tb, "after", 9.0, BLUE, SANS)
    assert para.findall(qn("a:br")) == []
    assert tb.text_frame.text == "after"


# --- the primitives the examples do not reach ----------------------------

def test_chip_is_a_pill_shaped_shape_carrying_its_text(slide):
    sp = P.chip(slide, 0, 0, 80, 20, "Tag", 8.0, BLUE, WHITE, BLUE, SANS)
    assert sp.text_frame.paragraphs[0].runs[0].text == "Tag"
    assert (sp.left, sp.width, sp.height) == (px(0), px(80), px(20))


def test_numcircle_writes_the_number_in_a_circle(slide):
    sp = P.numcircle(slide, 10, 10, 24, 7, BLUE, WHITE, SANS, 12.0)
    assert sp.text_frame.paragraphs[0].runs[0].text == "7"
    assert sp.width == sp.height == px(24)


def test_numcircle_requires_an_explicit_size(slide):
    with pytest.raises(TypeError):
        P.numcircle(slide, 0, 0, 24, 1, BLUE, WHITE, SANS)


def test_tri_is_centred_on_the_point_it_is_given(slide):
    sp = P.tri(slide, 100, 50, "down", BLUE, w=10, h=8)
    assert sp.left == px(95)
    assert sp.top == px(46)
    assert sp.rotation == 180


# --- footer geometry -----------------------------------------------------

def test_footer_keeps_every_part_inside_the_width_it_is_given(slide):
    """The left text box used to be a hardcoded 700px, so a narrow footer
    ran 300px past its own rule and collided with the page number."""
    P.footer(slide, "A title", 3, SANS, BLUE, BLUE, 8.0, width=400, left=48, y=600)
    right_edge = px(48 + 400)
    for sh in slide.shapes:
        assert sh.left + sh.width <= right_edge


def test_footer_does_not_overlap_the_page_number(slide):
    P.footer(slide, "A title", 3, SANS, BLUE, BLUE, 8.0, width=400, left=48, y=600)
    boxes = [sh for sh in slide.shapes if sh.has_text_frame]
    left_box = next(sh for sh in boxes if sh.text_frame.text == "A title")
    page_box = next(sh for sh in boxes if sh.text_frame.text == "3")
    assert left_box.left + left_box.width <= page_box.left


def test_chip_refuses_text_the_same_colour_as_its_fill(slide):
    """A chip inverted for a dark band whose caller forgot to invert the text
    exports as an empty pill; no count-based check sees it."""
    with pytest.raises(ValueError, match="invisible"):
        P.chip(slide, 0, 0, 80, 20, "HIGH", 9.0, WHITE, WHITE, BLUE, SANS)


def test_chip_allows_a_transparent_fill_whatever_the_text_colour(slide):
    sp = P.chip(slide, 0, 0, 80, 20, "Tag", 9.0, WHITE, None, BLUE, SANS)
    assert sp.text_frame.paragraphs[0].runs[0].text == "Tag"


# --- letter-spacing ------------------------------------------------------

def test_settext_writes_letter_spacing_in_hundredths_of_a_point(slide):
    """OOXML's a:rPr/@spc is hundredths of a point; the kit's argument is
    points, like every other size, and the conversion happens once."""
    sp = P.rrect(slide, 0, 0, 200, 20, fill=WHITE)
    P.settext(sp, [("KICKER", True)], 10.0, BLUE, SANS, spc=1.2)
    assert sp.text_frame.paragraphs[0].runs[0].font._rPr.get("spc") == "120"


def test_settext_writes_negative_letter_spacing(slide):
    sp = P.rrect(slide, 0, 0, 200, 40, fill=WHITE)
    P.settext(sp, [("Relay", True)], 30.0, BLUE, SANS, spc=-1.05)
    assert sp.text_frame.paragraphs[0].runs[0].font._rPr.get("spc") == "-105"


def test_no_spc_attribute_is_written_when_tracking_is_zero(slide):
    """The default must leave the run exactly as it was: an spc="0" on every
    run in the repository would be noise in every text dump and diff."""
    sp = P.rrect(slide, 0, 0, 200, 20, fill=WHITE)
    P.settext(sp, [("Plain", False)], 11.0, BLUE, SANS)
    assert sp.text_frame.paragraphs[0].runs[0].font._rPr.get("spc") is None


def test_textbox_and_add_para_carry_tracking_through(slide):
    tb = P.textbox(slide, 0, 0, 200, 40, [("Head", True)], 11.0, BLUE, SANS, spc=0.9)
    P.add_para(tb.text_frame, [("Sub", False)], 9.0, BLUE, SANS, spc=-0.4)
    paras = tb.text_frame.paragraphs
    assert paras[0].runs[0].font._rPr.get("spc") == "90"
    assert paras[1].runs[0].font._rPr.get("spc") == "-40"


def test_tracking_applies_to_every_run_in_a_paragraph(slide):
    tb = P.textbox(slide, 0, 0, 300, 20, [("A", True), ("B", False)], 10.0, BLUE, SANS, spc=1.2)
    assert [r.font._rPr.get("spc") for r in tb.text_frame.paragraphs[0].runs] == ["120", "120"]
