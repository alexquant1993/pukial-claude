"""Run attribution is tested through a saved package, not on an in-memory
paragraph: the point of preserving runs is what PowerPoint reads back."""

from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation

from deck_kit import merge, textedit
from deck_kit.brand import load_brand
from deck_kit.deck import add_slide, open_deck
from deck_kit.primitives import textbox

# TEMPLATE arrives from the session-scoped `template` fixture in conftest.py.
# Never a module constant pointing at out/: the gate runs the unit tests BEFORE
# it builds the template, and out/ is gitignored, so a hardcoded path fails on
# a clean checkout and, worse, silently uses a stale template on a dirty one.


@pytest.fixture(scope="module")
def brand():
    return load_brand("relay")


def mixed_deck(path, template, brand):
    """One slide, one paragraph split into three runs, the middle one bold.

    This is the shape PowerPoint produces on its own after a few edits, and
    the shape every naive retexting routine destroys.
    """
    prs = open_deck(template, brand)
    style = brand.style()
    slide = add_slide(prs, brand)
    textbox(slide, 80, 80, 900, 60,
            [("The gap is ", False), ("wide", True), (" in three domains.", False)],
            14, style.GREY, brand.sans)
    prs.save(str(path))
    return path


def only_paragraph(path):
    prs = Presentation(str(path))
    return prs, list(textedit.iter_paragraphs(prs.slides[0]))[0]


def test_paragraph_text_is_the_runs_not_the_frame(tmp_path, template, brand):
    _, para = only_paragraph(mixed_deck(tmp_path / "m.pptx", template, brand))
    assert textedit.paragraph_text(para) == "The gap is wide in three domains."
    assert len(para.runs) == 3


def test_an_edit_outside_the_bold_run_leaves_the_bold_intact(tmp_path, template, brand):
    """The whole point. A tail edit must not move, merge or unbold run 1."""
    src = mixed_deck(tmp_path / "m.pptx", template, brand)
    prs, para = only_paragraph(src)
    textedit.patch_paragraph(para, "The gap is wide in four domains.")
    out = tmp_path / "patched.pptx"
    prs.save(str(out))

    _, after = only_paragraph(out)
    assert textedit.paragraph_text(after) == "The gap is wide in four domains."
    bold = [r.text for r in after.runs if r.font.bold]
    assert bold == ["wide"], "mid-sentence bold was destroyed: %r" % (
        [(r.text, r.font.bold) for r in after.runs],)


def test_an_edit_inside_the_bold_run_stays_bold(tmp_path, template, brand):
    src = mixed_deck(tmp_path / "m.pptx", template, brand)
    prs, para = only_paragraph(src)
    textedit.patch_paragraph(para, "The gap is severe in three domains.")
    out = tmp_path / "patched.pptx"
    prs.save(str(out))

    _, after = only_paragraph(out)
    assert textedit.paragraph_text(after) == "The gap is severe in three domains."
    assert [r.text for r in after.runs if r.font.bold] == ["severe"]


def test_deleting_the_bold_word_empties_its_run_rather_than_reflowing(tmp_path, template, brand):
    """The delete spans the run0/run1 boundary: run0 loses its trailing
    space along with "wide", run1 is left empty rather than merged away or
    reused, and run2 is untouched. Checking only the concatenated text would
    pass under the naive body too, so this asserts the runs directly."""
    src = mixed_deck(tmp_path / "m.pptx", template, brand)
    prs, para = only_paragraph(src)
    textedit.patch_paragraph(para, "The gap is in three domains.")
    out = tmp_path / "patched.pptx"
    prs.save(str(out))
    _, after = only_paragraph(out)
    assert textedit.paragraph_text(after) == "The gap is in three domains."
    assert [(r.text, r.font.bold) for r in after.runs] == [
        ("The gap is", False), ("", True), (" in three domains.", False)]


def test_patch_paragraph_raises_on_a_paragraph_with_no_runs(tmp_path, template, brand):
    """The empty-paragraph guard from the ledger: an empty paragraph between
    two bullets is ordinary in a hand-edited deck and must not crash on
    owner[-1] of an empty list."""
    prs = open_deck(template, brand)
    style = brand.style()
    slide = add_slide(prs, brand)
    box = textbox(slide, 80, 80, 400, 40, "First line.", 14, style.GREY, brand.sans)
    empty_para = box.text_frame.add_paragraph()
    assert empty_para.runs == ()
    with pytest.raises(ValueError, match="no runs to attribute"):
        textedit.patch_paragraph(empty_para, "New text.")


def test_patch_paragraph_handles_a_run_that_holds_no_text(tmp_path, template, brand):
    """C2: runs can exist while holding no text - <a:r><a:t/></a:r> is
    ordinary in a hand-edited deck (an empty bullet the author typed into
    and then deleted, still leaving the run behind). `old == ""` then, but
    `runs` is non-empty, so the no-runs guard above does not fire and the
    replace/insert branch's `owner[-1]` on an empty `owner` list must not
    IndexError - the exact failure the no-runs guard's own message claims
    to have already fixed."""
    prs = open_deck(template, brand)
    style = brand.style()
    slide = add_slide(prs, brand)
    box = textbox(slide, 80, 80, 400, 40, "First line.", 14, style.GREY, brand.sans)
    para = box.text_frame.add_paragraph()
    run = para.add_run()
    assert len(para.runs) == 1
    assert textedit.paragraph_text(para) == ""

    textedit.patch_paragraph(para, "New text.")
    assert textedit.paragraph_text(para) == "New text."
    out = tmp_path / "empty_run_patched.pptx"
    prs.save(str(out))
    reopened = list(textedit.iter_paragraphs(Presentation(str(out)).slides[0]))
    assert textedit.paragraph_text(reopened[1]) == "New text."


def test_patch_paragraph_attaches_a_trailing_insert_to_the_last_run(tmp_path, template, brand):
    """The pure trailing-insert path: i1 == len(old), so the replace/insert
    branch falls to the owner[-1] fallback rather than owner[i1]. Deferred
    minor from Task 4, mis-triaged - writing this test is what surfaced the
    empty-run IndexError above, since both paths share the same fallback."""
    src = mixed_deck(tmp_path / "m.pptx", template, brand)
    prs, para = only_paragraph(src)
    textedit.patch_paragraph(para, "The gap is wide in three domains. Noted.")
    out = tmp_path / "patched_tail.pptx"
    prs.save(str(out))

    _, after = only_paragraph(out)
    assert textedit.paragraph_text(after) == "The gap is wide in three domains. Noted."
    runs = [(r.text, r.font.bold) for r in after.runs]
    assert runs[-1] == (" in three domains. Noted.", False), (
        "the trailing insert must attach to the last run: %r" % (runs,))
    assert runs[1] == ("wide", True), "the bold run must be untouched by a tail edit"


def test_patch_paragraph_asserts_the_reassembly(tmp_path, template, brand, monkeypatch):
    """A silent mis-assembly is worse than a crash: the deck ships wrong.

    The correct algorithm always reproduces `new` by construction, so this
    forces a mis-assembly by making the reassembly check itself lie, and
    confirms `patch_paragraph` raises rather than returning quietly."""
    src = mixed_deck(tmp_path / "m.pptx", template, brand)
    _, para = only_paragraph(src)

    real_paragraph_text = textedit.paragraph_text
    calls = []

    def lying_paragraph_text(p):
        calls.append(p)
        # First call inside patch_paragraph computes `old` and must be
        # truthful, or the routing itself breaks. Only the final
        # reassembly check gets lied to.
        if len(calls) > 1:
            return "not what was asked for"
        return real_paragraph_text(p)

    monkeypatch.setattr(textedit, "paragraph_text", lying_paragraph_text)
    with pytest.raises(AssertionError, match="run reassembly does not reproduce"):
        textedit.patch_paragraph(para, "Anything at all.")


def test_patch_text_requires_exactly_one_match(tmp_path, template, brand):
    src = mixed_deck(tmp_path / "m.pptx", template, brand)
    prs = Presentation(str(src))
    with pytest.raises(SystemExit) as e:
        textedit.patch_text(prs, [(1, "no such sentence", "x")])
    assert "hits=0" in str(e.value)


def test_patch_text_rejects_two_matches(tmp_path, template, brand):
    """Two matches means the patch is ambiguous and would silently change the
    wrong one on the next edit of the deck."""
    prs = open_deck(template, brand)
    style = brand.style()
    slide = add_slide(prs, brand)
    for y in (80, 200):
        textbox(slide, 80, y, 900, 60, "Repeated line.", 14, style.GREY, brand.sans)
    src = tmp_path / "twice.pptx"
    prs.save(str(src))

    prs = Presentation(str(src))
    with pytest.raises(SystemExit) as e:
        textedit.patch_text(prs, [(1, "Repeated line.", "Changed line.")])
    assert "hits=2" in str(e.value)
    # The exactly-once check must be a hard failure that has NOT already
    # written: nothing may be mutated before the count is known.
    texts = [textedit.paragraph_text(p)
             for p in textedit.iter_paragraphs(prs.slides[0])]
    assert texts.count("Repeated line.") == 2


def test_patch_text_is_atomic_across_the_patch_list(tmp_path, template, brand):
    """I1: a later patch's ambiguity must not leave an earlier one applied.

    Patch 1 (slide 1, unambiguous) would apply fine on its own; patch 2
    (slide 2) is ambiguous. A single pass that validates and mutates one
    patch at a time would already have rewritten slide 1 by the time patch
    2's SystemExit fires - resolve-all-then-apply must not."""
    prs = open_deck(template, brand)
    style = brand.style()
    slide1 = add_slide(prs, brand)
    textbox(slide1, 80, 80, 900, 60, "Unique first line.", 14, style.GREY, brand.sans)
    slide2 = add_slide(prs, brand)
    for y in (80, 200):
        textbox(slide2, 80, y, 900, 60, "Repeated line.", 14, style.GREY, brand.sans)
    src = tmp_path / "atomic.pptx"
    prs.save(str(src))

    prs = Presentation(str(src))
    with pytest.raises(SystemExit) as e:
        textedit.patch_text(prs, [
            (1, "Unique first line.", "Changed first line."),
            (2, "Repeated line.", "Changed line."),
        ])
    assert "hits=2" in str(e.value)
    texts = [textedit.paragraph_text(p)
             for p in textedit.iter_paragraphs(prs.slides[0])]
    assert texts == ["Unique first line."], (
        "the earlier, unambiguous patch must not have been applied: %r" % texts)


def test_patch_text_reaches_paragraphs_inside_groups(tmp_path, template, brand):
    """iter_paragraphs must recurse: a grouped shape is the normal state of a
    deck a human has been editing."""
    prs = open_deck(template, brand)
    style = brand.style()
    slide = add_slide(prs, brand)
    a = textbox(slide, 80, 80, 400, 40, "Inside a group.", 14, style.GREY, brand.sans)
    b = textbox(slide, 80, 140, 400, 40, "Also inside.", 14, style.GREY, brand.sans)
    slide.shapes.add_group_shape([a, b])
    src = tmp_path / "grouped.pptx"
    prs.save(str(src))

    prs = Presentation(str(src))
    assert textedit.patch_text(prs, [(1, "Inside a group.", "Inside a set.")]) == 1
    out = tmp_path / "grouped_out.pptx"
    prs.save(str(out))
    texts = [textedit.paragraph_text(p)
             for p in textedit.iter_paragraphs(Presentation(str(out)).slides[0])]
    assert "Inside a set." in texts


def test_shapes_by_id_resolves_through_groups(tmp_path, template, brand):
    """Keyed on shape id, not name: names repeat, and not on position or
    text, because both change."""
    prs = open_deck(template, brand)
    style = brand.style()
    slide = add_slide(prs, brand)
    a = textbox(slide, 80, 80, 400, 40, "A", 14, style.GREY, brand.sans)
    b = textbox(slide, 80, 140, 400, 40, "B", 14, style.GREY, brand.sans)
    ids = {a.shape_id, b.shape_id}
    slide.shapes.add_group_shape([a, b])
    src = tmp_path / "g.pptx"
    prs.save(str(src))

    slide = Presentation(str(src)).slides[0]
    assert ids <= set(textedit.shapes_by_id(slide))


def test_a_string_value_collapses_the_paragraph_to_one_run(tmp_path, template, brand):
    src = mixed_deck(tmp_path / "m.pptx", template, brand)
    prs, para = only_paragraph(src)
    sid = list(textedit.shapes_by_id(prs.slides[0]))[0]
    assert textedit.apply_text_map(prs, {1: {sid: {0: "One run now."}}}) == []
    out = tmp_path / "o.pptx"
    prs.save(str(out))
    _, after = only_paragraph(out)
    assert len(after.runs) == 1
    assert textedit.paragraph_text(after) == "One run now."


def test_a_list_value_assigns_run_by_run_and_keeps_each_format(tmp_path, template, brand):
    src = mixed_deck(tmp_path / "m.pptx", template, brand)
    prs, para = only_paragraph(src)
    sid = list(textedit.shapes_by_id(prs.slides[0]))[0]
    assert textedit.apply_text_map(
        prs, {1: {sid: {0: ["Risk is ", "acute", " here."]}}}) == []
    out = tmp_path / "o.pptx"
    prs.save(str(out))
    _, after = only_paragraph(out)
    assert [r.text for r in after.runs] == ["Risk is ", "acute", " here."]
    assert [r.text for r in after.runs if r.font.bold] == ["acute"]


def test_a_longer_list_deep_copies_the_last_run(tmp_path, template, brand):
    src = mixed_deck(tmp_path / "m.pptx", template, brand)
    prs, _ = only_paragraph(src)
    sid = list(textedit.shapes_by_id(prs.slides[0]))[0]
    assert textedit.apply_text_map(
        prs, {1: {sid: {0: ["a", "b", "c", "d", "e"]}}}) == []
    out = tmp_path / "o.pptx"
    prs.save(str(out))
    _, after = only_paragraph(out)
    assert [r.text for r in after.runs] == ["a", "b", "c", "d", "e"]


def test_none_deletes_the_paragraph(tmp_path, template, brand):
    prs = open_deck(template, brand)
    style = brand.style()
    slide = add_slide(prs, brand)
    box = textbox(slide, 80, 80, 400, 80, "keep", 14, style.GREY, brand.sans)
    from deck_kit.primitives import add_para
    add_para(box.text_frame, "drop", 14, style.GREY, brand.sans)
    src = tmp_path / "two.pptx"
    prs.save(str(src))

    prs = Presentation(str(src))
    sid = list(textedit.shapes_by_id(prs.slides[0]))[0]
    assert textedit.apply_text_map(prs, {1: {sid: {1: None}}}) == []
    out = tmp_path / "one.pptx"
    prs.save(str(out))
    paras = list(textedit.iter_paragraphs(Presentation(str(out)).slides[0]))
    assert [textedit.paragraph_text(p) for p in paras] == ["keep"]


def test_apply_text_map_accumulates_every_mismatch_instead_of_raising(tmp_path, template, brand):
    """The map applier never raises. One missing id must not hide the other
    nine problems in the same map - that is a whole round trip per defect."""
    src = mixed_deck(tmp_path / "m.pptx", template, brand)
    prs, _ = only_paragraph(src)
    problems = textedit.apply_text_map(
        prs, {1: {999999: {0: "x"}, 999998: {0: "y"}}})
    assert len(problems) == 2
    assert all(p[0] == 1 for p in problems)


def test_apply_text_map_reports_a_paragraph_index_that_does_not_exist(tmp_path, template, brand):
    src = mixed_deck(tmp_path / "m.pptx", template, brand)
    prs, _ = only_paragraph(src)
    sid = list(textedit.shapes_by_id(prs.slides[0]))[0]
    problems = textedit.apply_text_map(prs, {1: {sid: {7: "x"}}})
    assert len(problems) == 1 and "paragraph 7" in problems[0][2]


def _make_it_look_hand_edited(para):
    """Give the paragraph what a PowerPoint-authored one has and a
    python-pptx-authored one does not.

    Two facts, both measured rather than assumed. (a) `settext` builds through
    `add_textbox()` + `add_run()`, and neither emits an <a:endParaRPr> - so a
    fixture built by this repository has no endParaRPr at all, and a test that
    merely asserts "the new run comes before it" passes vacuously on a
    paragraph that has none. (b) The spellcheck attribute lives on <a:rPr>,
    not on <a:r>: across three human-edited decks in the engagement there are
    198 occurrences on rPr and zero on r, and <a:r> has no attributes in the
    schema at all.
    """
    from copy import deepcopy
    end = etree.SubElement(para._p, textedit.q("endParaRPr"))
    proto_rpr = para.runs[0]._r.find(textedit.q("rPr"))
    if proto_rpr is not None:
        end.attrib.update(proto_rpr.attrib)
        for child in proto_rpr:
            end.append(deepcopy(child))
    proto_rpr.set("err", "1")
    return end


def test_appended_runs_precede_endParaRPr_and_carry_no_err(tmp_path, template, brand):
    """Catalogue item 10. A run placed after <a:endParaRPr> is in the file and
    never appears on the slide, and a deep-copied run brings the spellcheck
    attribute with it."""
    src = mixed_deck(tmp_path / "m.pptx", template, brand)
    prs, para = only_paragraph(src)
    end = _make_it_look_hand_edited(para)

    made = textedit.append_runs(para, [" And more."], proto_index=0)

    # Asserted on the elements, not on a byte window of the slide part: a
    # fixed-width slice around the text can silently land on the wrong run.
    children = list(para._p)
    assert children.index(made[0]._r) < children.index(end), \
        "the appended run is after endParaRPr and PowerPoint will ignore it"
    appended_rpr = made[0]._r.find(textedit.q("rPr"))
    assert appended_rpr is not None
    assert appended_rpr.get("err") is None, "the spellcheck attribute was cloned"
    assert appended_rpr.get("dirty") is None

    out = tmp_path / "appended.pptx"
    prs.save(str(out))
    _, after = only_paragraph(out)
    assert textedit.paragraph_text(after).endswith("And more.")
    # And once in the saved package, so the doctrine holds in the artifact.
    parts = merge.read_package(out)
    body = parts[merge.slide_order(parts)[0]].decode("utf-8")
    assert body.index("And more.") < body.index("<a:endParaRPr")


def test_delete_shape_refuses_a_shape_the_timing_tree_targets(tmp_path, template, brand):
    """Catalogue item 8. A dangling p:spTgt makes automation-mode Open() fail
    outright - 'PowerPoint could not open the file' - rather than offer
    repair, so this cannot be left to the COM check to discover."""
    prs = open_deck(template, brand)
    style = brand.style()
    slide = add_slide(prs, brand)
    box = textbox(slide, 80, 80, 400, 40, "animated", 14, style.GREY, brand.sans)
    P = "http://schemas.openxmlformats.org/presentationml/2006/main"
    timing = etree.SubElement(slide._element, "{%s}timing" % P)
    tgt = etree.SubElement(timing, "{%s}spTgt" % P)
    tgt.set("spid", str(box.shape_id))
    src = tmp_path / "anim.pptx"
    prs.save(str(src))

    slide = Presentation(str(src)).slides[0]
    target = textedit.shapes_by_id(slide)[
        sorted(textedit.anim_targets(slide))[0]]
    assert textedit.anim_targets(slide)
    with pytest.raises(textedit.AnimationTargetError):
        textedit.delete_shape(slide, target)


def test_delete_shape_refuses_a_group_whose_child_is_an_animation_target(tmp_path, template, brand):
    """The animation guard checks the group's own id AND every descendant's.

    Deleting the group would delete the animated child with it, leaving the
    same dangling p:spTgt the guard exists to prevent - the guard must not
    look only at the group's own shape id, which the timing tree never
    mentions here."""
    prs = open_deck(template, brand)
    style = brand.style()
    slide = add_slide(prs, brand)
    a = textbox(slide, 80, 80, 200, 40, "animated child", 14, style.GREY, brand.sans)
    b = textbox(slide, 80, 140, 200, 40, "plain sibling", 14, style.GREY, brand.sans)
    group = slide.shapes.add_group_shape([a, b])
    P = "http://schemas.openxmlformats.org/presentationml/2006/main"
    timing = etree.SubElement(slide._element, "{%s}timing" % P)
    tgt = etree.SubElement(timing, "{%s}spTgt" % P)
    tgt.set("spid", str(a.shape_id))
    src = tmp_path / "anim_group.pptx"
    prs.save(str(src))

    slide = Presentation(str(src)).slides[0]
    group_shape = next(iter(slide.shapes))
    assert group_shape.shape_id == group.shape_id
    with pytest.raises(textedit.AnimationTargetError):
        textedit.delete_shape(slide, group_shape)


def test_delete_shape_removes_a_shape_nothing_animates(tmp_path, template, brand):
    prs = open_deck(template, brand)
    style = brand.style()
    slide = add_slide(prs, brand)
    box = textbox(slide, 80, 80, 400, 40, "free", 14, style.GREY, brand.sans)
    textedit.delete_shape(slide, box)
    out = tmp_path / "gone.pptx"
    prs.save(str(out))
    assert not list(textedit.iter_paragraphs(Presentation(str(out)).slides[0]))


def test_set_wrap_turns_wrap_on_in_the_saved_package(tmp_path, template, brand):
    """Catalogue item 12: text boxes inherited from a template may arrive with
    word_wrap off, so longer replacement copy overflows the right edge
    silently. Asserted in the XML, because python-pptx reports None for
    'inherit'."""
    prs = open_deck(template, brand)
    style = brand.style()
    slide = add_slide(prs, brand)
    box = textbox(slide, 80, 80, 200, 40, "long copy here", 14, style.GREY, brand.sans)
    box.text_frame.word_wrap = False
    src = tmp_path / "nowrap.pptx"
    prs.save(str(src))

    prs = Presentation(str(src))
    sid = list(textedit.shapes_by_id(prs.slides[0]))[0]
    assert textedit.set_wrap(prs, [(1, sid)]) == []
    out = tmp_path / "wrap.pptx"
    prs.save(str(out))

    parts = merge.read_package(out)
    xml = parts[merge.slide_order(parts)[0]].decode("utf-8")
    assert 'wrap="square"' in xml and 'wrap="none"' not in xml


def test_widen_sets_the_width_and_turns_wrap_on(tmp_path, template, brand):
    """The wrap half is pinned by setting word_wrap False before widen()
    runs and asserting the saved XML - textbox() defaults wrap on, so an
    in-memory `is True` check would pass even with widen()'s wrap line
    removed. Only the package settles catalogue item 12, because
    python-pptx reports None for an inherited word_wrap."""
    from deck_kit.geometry import EMU

    prs = open_deck(template, brand)
    style = brand.style()
    slide = add_slide(prs, brand)
    box = textbox(slide, 80, 80, 200, 40, "copy", 14, style.GREY, brand.sans)
    box.text_frame.word_wrap = False
    src = tmp_path / "narrow.pptx"
    prs.save(str(src))

    prs = Presentation(str(src))
    sid = list(textedit.shapes_by_id(prs.slides[0]))[0]
    assert textedit.widen(prs, [(1, sid, 420)]) == []
    out = tmp_path / "wide.pptx"
    prs.save(str(out))
    shape = list(textedit.shapes_by_id(Presentation(str(out)).slides[0]).values())[0]
    assert shape.width == 420 * EMU

    parts = merge.read_package(out)
    xml = parts[merge.slide_order(parts)[0]].decode("utf-8")
    assert 'wrap="square"' in xml and 'wrap="none"' not in xml


def test_widen_reports_a_shape_with_no_text_frame_as_a_miss(tmp_path, template, brand):
    """set_wrap reports a found-but-no-text-frame shape as a miss; widen must
    too, or half the contract silently no-ops with the caller none the
    wiser. A picture, not an autoshape: autoshapes always answer
    has_text_frame True even when empty."""
    prs = open_deck(template, brand)
    slide = add_slide(prs, brand)
    picture = slide.shapes.add_picture(str(brand.logo_primary), 100000, 100000)
    src = tmp_path / "noframe.pptx"
    prs.save(str(src))

    prs = Presentation(str(src))
    sid = list(textedit.shapes_by_id(prs.slides[0]))[0]
    assert textedit.widen(prs, [(1, sid, 420)]) == [(1, sid, 420)]


def test_set_wrap_reports_a_target_it_could_not_find(tmp_path, template, brand):
    src = mixed_deck(tmp_path / "m.pptx", template, brand)
    prs = Presentation(str(src))
    assert textedit.set_wrap(prs, [(1, 999999)]) == [(1, 999999)]


def test_backing_shape_finds_the_pill_and_delete_takes_both(tmp_path, template, brand):
    """Catalogue item 13: OOXML records no association between a text box and
    its pill background. It has to be recovered geometrically, and deleting
    one half leaves an orphan rectangle behind."""
    from deck_kit.primitives import rrect

    prs = open_deck(template, brand)
    style = brand.style()
    slide = add_slide(prs, brand)
    pill = rrect(slide, 100, 100, 300, 40, fill=style.CELL_BD)
    box = textbox(slide, 100, 100, 300, 40, "on a pill", 12, style.GREY, brand.sans)
    src = tmp_path / "pill.pptx"
    prs.save(str(src))

    slide = Presentation(str(src)).slides[0]
    by_id = textedit.shapes_by_id(slide)
    text = [s for s in by_id.values() if s.has_text_frame and s.text_frame.text == "on a pill"][0]
    found = textedit.backing_shape(slide, text)
    assert found is not None and found.shape_id != text.shape_id
    assert (found.left, found.top, found.width, found.height) == \
           (text.left, text.top, text.width, text.height)

    # The second half of item 13: deleting one leaves an orphan rectangle.
    prs = Presentation(str(src))
    slide = prs.slides[0]
    by_id = textedit.shapes_by_id(slide)
    text = [x for x in by_id.values()
            if x.has_text_frame and x.text_frame.text == "on a pill"][0]
    textedit.delete_with_backing(slide, text)
    out = tmp_path / "gone.pptx"
    prs.save(str(out))

    left = list(textedit.shapes_by_id(Presentation(str(out)).slides[0]).values())
    assert not [x for x in left if x.has_text_frame
                and x.text_frame.text == "on a pill"]
    assert not [x for x in left
                if (x.left, x.top, x.width, x.height) == (pill.left, pill.top,
                                                          pill.width, pill.height)], \
        "the pill was orphaned when its text box was deleted"


def test_backing_shape_returns_none_when_the_text_box_stands_alone(tmp_path, template, brand):
    src = mixed_deck(tmp_path / "m.pptx", template, brand)
    slide = Presentation(str(src)).slides[0]
    text = list(textedit.shapes_by_id(slide).values())[0]
    assert textedit.backing_shape(slide, text) is None


def test_delete_with_backing_refuses_a_card_enclosing_other_labels(tmp_path, template, brand):
    """A large empty auto-shape enclosing two bare labels, neither with its
    own pill, is a card frame. backing_shape() still finds it - it is the
    smallest empty auto-shape containing the centre, exactly as spec section
    4 stage 4 asks - but delete_with_backing() must refuse to take it: taking
    it deletes the card and every other label still inside it. Reproduced
    data loss, not a theoretical case."""
    from deck_kit.primitives import rrect

    prs = open_deck(template, brand)
    style = brand.style()
    slide = add_slide(prs, brand)
    card = rrect(slide, 80, 80, 400, 200, fill=style.CELL_BD)
    left_label = textbox(slide, 100, 100, 150, 40, "left label", 12, style.GREY, brand.sans)
    right_label = textbox(slide, 300, 220, 150, 40, "right label", 12, style.GREY, brand.sans)
    src = tmp_path / "card.pptx"
    prs.save(str(src))

    prs = Presentation(str(src))
    slide = prs.slides[0]
    by_id = textedit.shapes_by_id(slide)
    left = [x for x in by_id.values() if x.has_text_frame
            and x.text_frame.text == "left label"][0]
    result = textedit.delete_with_backing(slide, left)
    assert result is None
    out = tmp_path / "card_out.pptx"
    prs.save(str(out))

    left_after = list(textedit.shapes_by_id(Presentation(str(out)).slides[0]).values())
    assert not [x for x in left_after if x.has_text_frame
                and x.text_frame.text == "left label"]
    assert [x for x in left_after if x.has_text_frame
            and x.text_frame.text == "right label"], \
        "the card took the other label with it"
    assert [x for x in left_after
            if (x.left, x.top, x.width, x.height) == (card.left, card.top,
                                                       card.width, card.height)], \
        "the card itself was deleted along with one label"


def test_ensure_notes_body_clones_a_donor_placeholder(tmp_path, template, brand):
    """Catalogue item 11: a notes master without placeholders makes note
    insertion fail silently - the text is written and never appears.

    out/template.pptx carries no ppt/notesMasters/ part at all, and
    python-pptx's default notes master DOES carry a BODY placeholder - so a
    naive fixture that merely writes a note to a target slide never drives
    notes_text_frame to None, the donor-cloning branch never runs, and the
    test pins nothing. The fixture below constructs the precondition itself:
    materialise the master through a donor slide (which keeps its own,
    unaffected clone of the BODY placeholder), let the target clone one too,
    then strip the BODY placeholder from both the master and the target's own
    notes slide so ensure_notes_body genuinely has nothing to return without
    cloning from the donor.
    """
    prs = open_deck(template, brand)
    style = brand.style()
    donor_slide = add_slide(prs, brand)
    textbox(donor_slide, 80, 80, 400, 40, "one", 14, style.GREY, brand.sans)
    donor_slide.notes_slide.notes_text_frame.text = "donor note"

    target = add_slide(prs, brand)
    textbox(target, 80, 80, 400, 40, "two", 14, style.GREY, brand.sans)
    target.notes_slide  # materialise target's own notes slide while the
                         # master still has a BODY placeholder to clone

    def _strip_body(shapes):
        from pptx.enum.shapes import PP_PLACEHOLDER
        for shape in list(shapes):
            if (shape.is_placeholder
                    and shape.placeholder_format.type == PP_PLACEHOLDER.BODY):
                shape._element.getparent().remove(shape._element)

    _strip_body(prs.notes_master.placeholders)
    _strip_body(target.notes_slide.shapes)
    assert target.notes_slide.notes_text_frame is None, (
        "the fixture did not strip the BODY placeholder from the target's "
        "own notes slide; item 11 would be unpinned")

    frame = textedit.ensure_notes_body(target, prs)
    assert frame is not None
    frame.text = "target note"
    out = tmp_path / "notes.pptx"
    prs.save(str(out))

    got = Presentation(str(out))
    assert got.slides[1].notes_slide.notes_text_frame.text == "target note"
    assert got.slides[0].notes_slide.notes_text_frame.text == "donor note"
