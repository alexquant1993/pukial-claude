"""The version driver's contract: never write over its input, write every
intermediate, and refuse to continue when the deck is not the file the build
was written against."""

import subprocess
import sys
from pathlib import Path

import pytest
from pptx import Presentation

from deck_kit.brand import load_brand
from deck_kit.deck import add_slide, open_deck
from deck_kit.merge import read_package
from deck_kit.primitives import textbox

# `template` arrives from the session-scoped fixture in conftest.py, passed
# as a normal pytest fixture argument to each test that needs it. Never a
# module constant pointing at out/: the gate runs the unit tests BEFORE it
# builds the template, and out/ is gitignored, so a hardcoded path fails on
# a clean checkout and, worse, silently uses a stale template on a dirty one.
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import new_version  # noqa: E402


@pytest.fixture(scope="module")
def brand():
    return load_brand("relay")


def _deck(path, template, brand, labels):
    prs = open_deck(template, brand)
    style = brand.style()
    for label in labels:
        slide = add_slide(prs, brand)
        textbox(slide, 80, 80, 600, 40, label, 20, style.GREY, brand.sans)
    prs.save(str(path))
    return path


def test_new_version_refuses_to_write_over_its_input(tmp_path, template, brand):
    """Catalogue item 21: the author's file is frequently open in PowerPoint
    and therefore locked. Nothing is overwritten, ever."""
    src = _deck(tmp_path / "author.pptx", template, brand, ["A", "B"])
    with pytest.raises(ValueError) as e:
        new_version.Version(source=src, work=tmp_path, final=src)
    assert "overwrite" in str(e.value).lower()


def test_new_version_writes_every_intermediate(tmp_path, template, brand):
    src = _deck(tmp_path / "author.pptx", template, brand, ["A", "B"])
    work = tmp_path / "work"
    v = new_version.Version(source=src, work=work, final=tmp_path / "out.pptx")

    v.step("noop", lambda prs: None)
    v.step("retitle", lambda prs: setattr(
        prs.slides[0].shapes[0].text_frame, "text", "A2"))

    assert (work / "step0_noop.pptx").exists()
    assert (work / "step1_retitle.pptx").exists()
    assert Presentation(str(work / "step1_retitle.pptx")).slides[0].shapes[0]\
        .text_frame.text == "A2"
    assert Presentation(str(work / "step0_noop.pptx")).slides[0].shapes[0]\
        .text_frame.text == "A", "an earlier step must not be rewritten"

    # Each stepN file is a genuine, independently openable OOXML package, and
    # "retitle" changed exactly the part it should have and nothing else -
    # a claim `"[Content_Types].xml" in parts` cannot make, since that key is
    # present in every package python-pptx (or this module) ever writes and
    # so cannot fail against a broken step. Compared part-by-part instead:
    # only the slide's own XML differs; ppt/presentation.xml, which encodes
    # no slide text, is byte-for-byte the same across both steps.
    before = read_package(work / "step0_noop.pptx")
    after = read_package(work / "step1_retitle.pptx")
    changed = sorted(k for k in before if before[k] != after[k])
    assert changed == ["ppt/slides/slide1.xml"], changed
    assert before["ppt/presentation.xml"] == after["ppt/presentation.xml"]


def test_new_version_asserts_the_file_it_was_written_against(tmp_path, template, brand):
    """Catalogue item 22: regenerating from a stale assumption destroys the
    author's manual edits. The build stops instead."""
    src = _deck(tmp_path / "author.pptx", template, brand, ["Overview", "Roadmap"])
    v = new_version.Version(source=src, work=tmp_path / "w",
                            final=tmp_path / "out.pptx")
    with pytest.raises(AssertionError):
        v.expect(slide_count=2, fingerprint={2: "Findings"})
    v.expect(slide_count=2, fingerprint={2: "Roadmap"})


def test_new_version_runs_a_subprocess_step_and_checks_its_output(tmp_path, template, brand):
    """The zip-surgery stages run out of process: a crash inside them must not
    be able to corrupt an in-memory Presentation."""
    src = _deck(tmp_path / "author.pptx", template, brand, ["A"])
    work = tmp_path / "w"
    v = new_version.Version(source=src, work=work, final=tmp_path / "out.pptx")
    v.subprocess_step("copy", [sys.executable, "-c",
                               "import shutil,sys; shutil.copyfile(sys.argv[1], sys.argv[2])"])
    assert (work / "step0_copy.pptx").exists()
    assert len(Presentation(str(work / "step0_copy.pptx")).slides) == 1


def test_new_version_step_checks_declared_slide_count(tmp_path, template, brand):
    """A `step` that changes the slide count without declaring it must fail -
    half of the controller ruling that overrides the brief: `step` and
    `subprocess_step` take `slides=`, asserted after the step runs, because
    `delete_positions` and `merge` both change the count and a silent
    off-by-one is invisible until someone presents the deck."""
    src = _deck(tmp_path / "author.pptx", template, brand, ["A", "B"])
    work = tmp_path / "w"
    v = new_version.Version(source=src, work=work, final=tmp_path / "out.pptx")
    with pytest.raises(AssertionError):
        v.step("drop_last", lambda prs: prs.slides._sldIdLst.remove(
            prs.slides._sldIdLst[-1]), slides=2)


def test_new_version_subprocess_step_checks_declared_slide_count(tmp_path, template, brand):
    """The other half of the same ruling, for `subprocess_step`: it re-opens
    the child's output file and counts slides itself rather than trusting the
    exit code, because Task 9 drives both zip-surgery stages through this
    method with `slides=` at every call site. A regression in the re-open-
    and-assert would pass silently here and surface as a wrong slide count in
    a generated deck.

    The wrapper is a real child process - two positional paths, no options -
    so this exercises the actual subprocess path rather than a stub."""
    src = _deck(tmp_path / "author.pptx", template, brand, ["A", "B"])
    work = tmp_path / "w"
    wrapper = tmp_path / "_drop_last_slide.py"
    wrapper.write_text(
        "import sys\n"
        "from pptx import Presentation\n"
        "prs = Presentation(sys.argv[1])\n"
        "prs.slides._sldIdLst.remove(prs.slides._sldIdLst[-1])\n"
        "prs.save(sys.argv[2])\n",
        encoding="utf-8")
    v = new_version.Version(source=src, work=work, final=tmp_path / "out.pptx")
    with pytest.raises(AssertionError) as e:
        v.subprocess_step("drop_last", [sys.executable, str(wrapper)], slides=2)
    assert "drop_last" in str(e.value)


def test_new_version_finalise_writes_the_named_output(tmp_path, template, brand):
    src = _deck(tmp_path / "author.pptx", template, brand, ["A"])
    final = tmp_path / "v2.pptx"
    v = new_version.Version(source=src, work=tmp_path / "w", final=final)
    v.step("noop", lambda prs: None)
    v.finalise()
    assert final.exists() and len(Presentation(str(final)).slides) == 1


def test_dump_text_round_trips_and_diff_reports_the_change(tmp_path, template, brand):
    """Catalogue item 22, the other half: dump and diff the author's current
    file before touching it, to discover what they changed by hand."""
    a = _deck(tmp_path / "a.pptx", template, brand, ["Overview", "Roadmap"])
    prs = Presentation(str(a))
    prs.slides[1].shapes[0].text_frame.text = "Roadmap 2027"
    b = tmp_path / "b.pptx"
    prs.save(str(b))

    def dump(path, out):
        subprocess.check_call(
            [sys.executable, str(ROOT / "scripts" / "dump_text.py"),
             "--deck", str(path), "--out", str(out)])
        return out.read_text(encoding="utf-8").splitlines()

    da = dump(a, tmp_path / "a.txt")
    db = dump(b, tmp_path / "b.txt")
    assert len(da) == len(db)
    changed = [(x, y) for x, y in zip(da, db) if x != y]
    assert len(changed) == 1 and "Roadmap 2027" in changed[0][1]

    rc = subprocess.call([sys.executable, str(ROOT / "scripts" / "diff_package.py"),
                          str(a), str(b)])
    assert rc == 1, "diff_package must exit 1 when the packages differ"
    rc = subprocess.call([sys.executable, str(ROOT / "scripts" / "diff_package.py"),
                          str(a), str(a)])
    assert rc == 0


def test_dump_text_refuses_out_equal_to_deck(tmp_path, template, brand):
    """I2: `--deck deck.pptx --out deck.pptx` used to silently replace the
    deck with a text dump - and references/04 instructs the reader to run
    this against the author's live file."""
    src = _deck(tmp_path / "author.pptx", template, brand, ["A"])
    before = read_package(src)
    rc = subprocess.call(
        [sys.executable, str(ROOT / "scripts" / "dump_text.py"),
         "--deck", str(src), "--out", str(src)],
        stderr=subprocess.DEVNULL)
    assert rc != 0
    assert read_package(src) == before, "the deck must be untouched"


def test_dump_text_expect_paragraphs_is_exact(tmp_path, template, brand):
    """The gate's `walkthrough text dump` step asserted nothing beyond exit 0,
    so a deck that dumped nothing passed the step whose purpose is to prove the
    dump works. Derived, not observed: `_deck` writes one text box per label
    and `add_slide` adds no other text frame, so two labels are two
    paragraphs."""
    src = _deck(tmp_path / "two.pptx", template, brand, ["A", "B"])

    def run(expect):
        return subprocess.run(
            [sys.executable, str(ROOT / "scripts" / "dump_text.py"),
             "--deck", str(src), "--out", str(tmp_path / "out.txt"),
             "--expect-paragraphs", str(expect)],
            capture_output=True, text=True)

    wrong = run(3)
    assert wrong.returncode == 1
    assert "FAIL: expected 3 paragraph(s), found 2" in wrong.stdout
    right = run(2)
    assert right.returncode == 0, right.stdout + right.stderr
    assert "FAIL" not in right.stdout
