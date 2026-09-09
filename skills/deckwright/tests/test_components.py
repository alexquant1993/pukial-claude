import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE

from deck_kit.brand import load_brand
from deck_kit.components import (Pill, PillStyle, draw_pill, explain_box,
                                 flow_pills, loop_line, matrix_frame, pill_w)
from deck_kit.geometry import px
from deck_kit.metrics import TextMetrics, WarnRegister

# Literal swatches, not the brand's: what these tests assert is the kit's
# geometry and text, and a colour it does not read from a style module cannot
# drift when the brand's palette does.
BLUE = RGBColor(0x25, 0x63, 0xEB)
NAVY = RGBColor(0x0F, 0x17, 0x2A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SANS = "Public Sans"
STYLE = PillStyle(fill=WHITE, line=BLUE, lw=1.2, fg=NAVY)

NL = chr(10)
COLS = [("Deferred", "T-1"), ("Near", "SECONDS"), ("In flight", "MS")]
# Icon names brands/relay actually ships, in its ink colour key.
ROWS = [("activity", "Signals &" + NL + "evidence"),
        ("terminal", "Models &" + NL + "graphs")]


@pytest.fixture
def slide():
    prs = Presentation()
    return prs.slides.add_slide(prs.slide_layouts[6])


@pytest.fixture
def m():
    return TextMetrics("brands/relay/fonts", "PublicSans-Regular.ttf",
                       "PublicSans-Bold.ttf", 1.11, 1.06)


@pytest.fixture
def brand():
    return load_brand("relay")


@pytest.fixture
def style(brand):
    return brand.style()


# --- pills ---------------------------------------------------------------

def test_pill_width_grows_with_the_title(m):
    assert (pill_w(m, Pill("abcdefghij", STYLE), 9.5, 9.5)
            > pill_w(m, Pill("ab", STYLE), 9.5, 9.5))


def test_pill_width_accounts_for_a_longer_sublabel(m):
    plain = pill_w(m, Pill("Feature Store", STYLE), 9.5, 9.5)
    subbed = pill_w(m, Pill("Feature Store", STYLE,
                            sub="today: 2 of 6 channels, nowhere near complete"),
                    9.5, 9.5)
    assert subbed > plain


def test_draw_pill_writes_the_title(slide):
    sp = draw_pill(slide, 0, 0, 120, 23, Pill("Labelling", STYLE), SANS, 9.5, 9.5)
    assert sp.text_frame.paragraphs[0].runs[0].text == "Labelling"


def test_draw_pill_with_a_sublabel_writes_two_paragraphs(slide):
    st = PillStyle(fill=WHITE, line=BLUE, lw=1.2, fg=NAVY, sub_fg=BLUE)
    sp = draw_pill(slide, 0, 0, 160, 32,
                   Pill("Feature Store", st, sub="today: 2 of 6"), SANS, 9.5, 9.5)
    assert len(sp.text_frame.paragraphs) == 2
    assert sp.text_frame.paragraphs[1].runs[0].text == "today: 2 of 6"


def test_draw_pill_adds_an_underline_bar_when_the_style_asks(slide):
    st = PillStyle(fill=WHITE, line=BLUE, lw=1.2, fg=NAVY, bar=BLUE)
    before = len(slide.shapes)
    draw_pill(slide, 0, 0, 120, 27, Pill("Base", st), SANS, 9.5, 9.5)
    assert len(slide.shapes) - before == 2


# --- flow ----------------------------------------------------------------

def test_flow_pills_fits_a_short_row_without_warning(slide, m):
    warn = WarnRegister()
    before = len(slide.shapes)
    flow_pills(slide, (0, 0, 400, 60), [Pill("A", STYLE), Pill("B", STYLE)],
               m, SANS, warn)
    assert warn.clean
    assert len(slide.shapes) - before == 2, "a clean register is not proof it drew"


def test_flow_pills_wraps_onto_a_second_row(slide, m):
    """Counting shapes proves nothing about rows. Distinct tops do."""
    warn = WarnRegister()
    pills = [Pill("Capability %d" % i, STYLE) for i in range(4)]
    flow_pills(slide, (0, 0, 220, 120), pills, m, SANS, warn)
    assert len({sh.top for sh in slide.shapes}) >= 2


def test_flow_pills_puts_everything_on_one_row_when_it_fits(slide, m):
    warn = WarnRegister()
    pills = [Pill("A", STYLE), Pill("B", STYLE), Pill("C", STYLE)]
    flow_pills(slide, (0, 0, 600, 60), pills, m, SANS, warn)
    assert len({sh.top for sh in slide.shapes}) == 1


def test_height_overflow_does_not_also_report_a_width_overflow(slide, m):
    """The two failures are different and must not be conflated: these pills
    fit the width, and only the stack is too tall."""
    warn = WarnRegister()
    pills = [Pill("A somewhat long capability name", STYLE) for _ in range(3)]
    flow_pills(slide, (0, 0, 300, 40), pills, m, SANS, warn)
    assert "HEIGHT" in warn.dump()
    assert "OVERFLOW" not in warn.dump()


def test_width_overflow_does_not_also_report_a_height_overflow(slide, m):
    warn = WarnRegister()
    long_title = "An impossibly long capability title for any cell this narrow"
    flow_pills(slide, (0, 0, 200, 400), [Pill(long_title, STYLE)], m, SANS, warn)
    assert "OVERFLOW" in warn.dump()
    assert "HEIGHT" not in warn.dump()


def test_flow_pills_catches_a_pill_wider_than_its_cell(slide, m):
    """A single over-wide pill wraps onto a row of its own and still fits
    vertically, so a height-only check reports a clean build while the text
    runs past the cell edge. That was the hole in the source repository, and
    the whole zero-warning gate depends on it not being there."""
    warn = WarnRegister()
    long_title = ("An impossibly long capability title that no cell of this "
                  "width could contain at any rung of the ladder")
    flow_pills(slide, (0, 0, 300, 400), [Pill(long_title, STYLE)], m, SANS, warn)
    assert not warn.clean
    assert "OVERFLOW" in warn.dump()
    assert "wide" in warn.dump()


def test_the_ladder_steps_down_for_width_not_only_for_height(slide, m):
    """A cell roomy in height but tight in width used to keep the first rung
    and then report an overflow the next rung would have fixed."""
    warn = WarnRegister()
    pill = Pill("Feature store, every channel", STYLE)
    widest = pill_w(m, pill, 9.5, 9.5)
    narrowest = pill_w(m, pill, 8.5, 9.5)
    cell_w = (widest + narrowest) / 2.0 + 12     # fits at 8.5, not at 9.5
    flow_pills(slide, (0, 0, cell_w, 400), [pill], m, SANS, warn)
    assert warn.clean, warn.dump()
    assert slide.shapes[0].width < px(widest)


def test_flow_pills_labels_the_cell_in_its_warning(slide, m):
    """The register is a set, so two cells overflowing by the same amount
    collapse into one entry unless the cell is named."""
    warn = WarnRegister()
    pills = [Pill("A somewhat long capability name", STYLE) for _ in range(3)]
    flow_pills(slide, (0, 0, 300, 40), pills, m, SANS, warn, label="r0c0")
    flow_pills(slide, (0, 0, 300, 40), pills, m, SANS, warn, label="r1c2")
    assert len(warn.dump().splitlines()) == 2


def test_flow_pills_with_no_pills_draws_nothing_and_stays_clean(slide, m):
    warn = WarnRegister()
    before = len(slide.shapes)
    flow_pills(slide, (0, 0, 300, 40), [], m, SANS, warn)
    assert len(slide.shapes) == before
    assert warn.clean


# --- the components that previously had no tests at all ------------------

def test_matrix_frame_returns_a_cell_per_row_and_column(slide, style, brand):
    cells = matrix_frame(slide, 140, 96, COLS, ROWS, style.GEOM, style,
                         style.SANS, brand.icon_dir)
    assert sorted(cells) == [(0, 0), (0, 1), (0, 2), (1, 0), (1, 1), (1, 2)]


def test_matrix_frame_cells_match_the_declared_grid(slide, style, brand):
    cells = matrix_frame(slide, 140, 96, COLS, ROWS, style.GEOM, style,
                         style.SANS, brand.icon_dir)
    for (r, c), (x, y, w, h) in cells.items():
        assert x == style.GEOM["col_x"][c]
        assert w == style.GEOM["col_w"]
        assert h == 96
    assert cells[(1, 0)][1] > cells[(0, 0)][1]


def test_matrix_frame_accepts_a_single_line_row_label(slide, style, brand):
    """Unpacking a fixed pair turned a one-line label into a ValueError
    raised from inside the frame."""
    cells = matrix_frame(slide, 140, 96, COLS, [("activity", "Signals")],
                         style.GEOM, style, style.SANS, brand.icon_dir)
    assert len(cells) == 3


def test_matrix_frame_accepts_a_three_line_row_label(slide, style, brand):
    label = NL.join(["Signals", "and", "evidence"])
    cells = matrix_frame(slide, 140, 96, COLS, [("activity", label)],
                         style.GEOM, style, style.SANS, brand.icon_dir)
    assert len(cells) == 3


def test_matrix_frame_marks_an_empty_cell(slide, style, brand):
    matrix_frame(slide, 140, 96, COLS, ROWS, style.GEOM, style, style.SANS,
                 brand.icon_dir)
    plain = len(slide.shapes)
    prs = Presentation()
    other = prs.slides.add_slide(prs.slide_layouts[6])
    matrix_frame(other, 140, 96, COLS, ROWS, style.GEOM, style, style.SANS,
                 brand.icon_dir, empty={(1, 2)})
    assert len(other.shapes) > plain, "the empty cell adds its dash"


def test_matrix_frame_raises_on_an_unknown_icon(slide, style, brand):
    with pytest.raises(FileNotFoundError):
        matrix_frame(slide, 140, 96, COLS, [("no-such-icon", "A")], style.GEOM,
                     style, style.SANS, brand.icon_dir)


def test_draw_pill_takes_the_brands_control_radius(slide, m):
    """A radius is a design-system value. It used to be the literal 7 in the
    kit, which put one brand's furniture in src/ - the same mistake as the
    hardcoded icon colour key. brands/relay is square by rule, and
    examples/ai-horizon's S09 passes RAD_CONTROL."""
    sharp = draw_pill(slide, 0, 0, 120, 23, Pill("A", STYLE), SANS, 9.5, 9.5, rad=0)
    default = draw_pill(slide, 0, 40, 120, 23, Pill("A", STYLE), SANS, 9.5, 9.5)
    assert sharp.auto_shape_type == MSO_SHAPE.RECTANGLE
    assert default.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE
    control = draw_pill(slide, 0, 80, 120, 23, Pill("A", STYLE), SANS, 9.5, 9.5, rad=2)
    assert control.adjustments[0] < default.adjustments[0]


def test_flow_pills_hands_its_radius_to_every_pill_it_draws(slide, m):
    """The radius has to reach the flow, not only draw_pill: a builder that
    lays its pills out with flow_pills never calls draw_pill itself."""
    warn = WarnRegister()
    before = len(slide.shapes)
    flow_pills(slide, (0, 0, 300, 60), [Pill("A", STYLE), Pill("B", STYLE)],
               m, SANS, warn, rad=0)
    drawn = list(slide.shapes)[before:]
    assert drawn, "flow_pills drew nothing"
    assert all(sh.auto_shape_type == MSO_SHAPE.RECTANGLE for sh in drawn)


def test_matrix_frame_takes_the_brands_surface_radius(slide, style, brand):
    """Same argument one level up: the cell is a surface, and Relay's
    surfaces are square."""
    prs = Presentation()
    other = prs.slides.add_slide(prs.slide_layouts[6])
    matrix_frame(slide, 140, 96, COLS, ROWS, style.GEOM, style, style.SANS,
                 brand.icon_dir, cell_rad=0)
    matrix_frame(other, 140, 96, COLS, ROWS, style.GEOM, style, style.SANS,
                 brand.icon_dir)
    kinds = {sh.auto_shape_type for sh in slide.shapes
             if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE}
    assert MSO_SHAPE.ROUNDED_RECTANGLE not in kinds
    assert any(sh.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE
               for sh in other.shapes
               if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE)


def test_explain_box_writes_title_and_description(slide, style, m):
    warn = WarnRegister()
    explain_box(slide, (0, 0, 300, 90), "Labelling",
                "Every fraud tagged by its modus operandi.",
                style, style.SANS, m, warn)
    texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
    assert any("Labelling" in t for t in texts)
    assert warn.clean


def test_explain_box_warns_when_the_copy_wraps_past_the_line_budget(slide, style, m):
    warn = WarnRegister()
    explain_box(slide, (0, 0, 120, 90), "A fairly long title for this box",
                "And a description considerably longer than the box can take "
                "at this width without wrapping many times over.",
                style, style.SANS, m, warn)
    assert "WRAP" in warn.dump()


def test_loop_line_measures_what_it_places(slide, style, m):
    """It sets wrap off, so an overrun runs out of the box with nothing to
    stop it. It was the one component that placed without measuring."""
    warn = WarnRegister()
    loop_line(slide, 600, [("short note", False)], style.GEOM, style,
              style.SANS, m, warn)
    assert warn.clean

    over = WarnRegister()
    loop_line(slide, 600, [("a very long transversal note " * 12, False)],
              style.GEOM, style, style.SANS, m, over)
    assert "OVERFLOW" in over.dump()
