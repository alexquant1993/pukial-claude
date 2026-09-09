"""Build a template deck satisfying a brand's contract.

The builders open a template and purge its slides, inheriting its masters,
layouts and background. This produces such a template from nothing, so the
contract is expressed by code rather than by a binary nobody can regenerate.

python-pptx cannot add shapes to a layout, so the background is built on a
scratch slide and its element is moved into the layout's shape tree.

`--variant NAME` builds a SECOND master from the same brand's contract: the
master and layout names get " (NAME)" appended and the background is painted
in the brand's own PANEL rather than its BG_TOP/BG_BOTTOM. It exists for the
Phase 2 round trip, which needs a host and a donor whose masters are genuinely
different parts - `variant_spec` below says what "genuinely different" means
and what would silently stop being tested without it.

Run: uv run --offline --no-project --with python-pptx --with pillow \
       python scripts/make_template.py --brand relay --out out/template.pptx
     uv run --offline --no-project --with python-pptx --with pillow \
       python scripts/make_template.py --brand relay --variant donor \
         --out out/template2.pptx
"""

import argparse
import copy
import sys
from dataclasses import replace
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn
from pptx.util import Emu

from deck_kit.brand import load_brand
from deck_kit.deck import _drop
from deck_kit.geometry import px
from deck_kit.xmlfill import grad

BLANK = "Blank"


def variant_spec(spec, variant):
    """The same brand's contract under a second master and layout name.

    `dataclasses.replace` rather than mutation: a loaded brand is a
    module-level singleton, so renaming it in place would rename it for every
    other caller in the process. replace() re-runs `__post_init__`, so the
    by-name-never-by-index rule still refuses an index-shaped selector on the
    variant - tests/test_template.py::
    test_a_variant_spec_still_refuses_an_index_shaped_selector.

    Why a variant rather than a second brand: the Phase 2 round trip has to
    transplant slides ACROSS masters, which is catalogue item 5 and the whole
    reason `deck_kit.merge` exists. Two templates generated from one brand
    without this flag are the same master twice - the copied master is
    byte-identical to the destination's, and the master-transplant path is
    exercised against the easiest possible input. With it, and still generated
    by one code path, the two templates carry:

      - the same part name (`ppt/slideMasters/slideMaster1.xml` in both), so
        merge has to mint a fresh, non-colliding destination part name;
      - the same `p:sldLayoutId` value (2147483655, python-pptx's own), so
        merge has to re-mint it above the destination's maximum - a genuine
        collision, not a no-op;
      - DIFFERENT master and layout names, so `deck._layout` can still
        resolve a layout by name in the merged file. Two masters carrying one
        layout name is the failure tests/test_brand.py::
        test_no_two_brands_share_a_master_or_layout_name forbids between
        brands, and a variant must not reintroduce it.

    tests/test_template.py::
    test_a_variant_template_is_a_second_master_of_the_same_brand asserts all
    three against the two generated files.
    """
    return replace(spec,
                   master_name="%s (%s)" % (spec.master_name, variant),
                   layout_name="%s (%s)" % (spec.layout_name, variant))


def _background_stops(style, variant):
    """The two gradient stops the layout background is painted with.

    A variant's ground is the brand's own PANEL, flat. It is not a colour
    this script chose: every value still comes from the brand's style module,
    which is where colour lives. A visibly different ground is what makes the
    two masters tell apart in an exported PNG - the names alone are invisible
    to anyone looking at the deck.
    """
    if variant is None:
        return [(0, style.BG_TOP, 100), (100, style.BG_BOTTOM, 100)]
    panel = getattr(style, "PANEL", None)
    if panel is None:
        raise SystemExit("--variant needs the brand's style module to name "
                         "PANEL, the quiet ground a variant master is painted "
                         "in; %s does not" % style.__name__)
    return [(0, panel, 100), (100, panel, 100)]


def _prune_layouts(master, keep):
    layouts = master.slide_layouts
    for layout in list(layouts):
        if layout.name != keep:
            layouts.remove(layout)


def _position_placeholder(ph, x, y, w, h):
    """Give a layout placeholder an explicit `p:spPr/a:xfrm`.

    A layout placeholder normally has no `xfrm` of its own and inherits its
    box from the master. Task 7 needs the surviving SLIDE_NUMBER placeholder
    positioned from the brand's PAGE_X/PAGE_Y/PAGE_W/PAGE_H rather than
    wherever the built-in template happens to put it.
    """
    spPr = ph._element.find(qn("p:spPr"))
    xfrm = etree.SubElement(spPr, qn("a:xfrm"))
    off = etree.SubElement(xfrm, qn("a:off"))
    off.set("x", str(int(px(x))))
    off.set("y", str(int(px(y))))
    ext = etree.SubElement(xfrm, qn("a:ext"))
    ext.set("cx", str(int(px(w))))
    ext.set("cy", str(int(px(h))))


def _background_element(prs, style, w, h, variant=None):
    """Draw the background on a scratch slide and hand back its element."""
    scratch = prs.slides.add_slide(next(iter(prs.slide_masters[0].slide_layouts)))
    sp = scratch.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Emu(w), Emu(h))
    grad(sp, _background_stops(style, variant), ang=90)
    sp.line.fill.background()
    sp.shadow.inherit = False
    element = copy.deepcopy(sp._element)
    _drop(prs, list(prs.slides._sldIdLst))
    return element


def build(spec, out_path, variant=None):
    """Write the template.

    `spec` is already the variant's spec when one was asked for - the caller
    renames it through `variant_spec` - because `inspect_template.py` has to
    be able to check the result against the same contract object.
    """
    style = spec.style()
    prs = Presentation()
    w, h = spec.slide_size_emu
    prs.slide_width, prs.slide_height = Emu(w), Emu(h)

    master = prs.slide_masters[0]
    master.name = spec.master_name
    _prune_layouts(master, BLANK)
    layout = master.slide_layouts[0]

    background = _background_element(prs, style, w, h, variant)

    layout.name = spec.layout_name
    for ph in list(layout.placeholders):
        # Every placeholder is stripped except SLIDE_NUMBER, which Task 7
        # needs kept as a real inherited placeholder holding an <a:fld
        # type="slidenum"> - the only artifact in the repository that can
        # exercise that flavour of pagenums.find_marker, and the reason
        # add_slide's keep_placeholders parameter has anything to preserve.
        if ph.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
            _position_placeholder(ph, style.PAGE_X, style.PAGE_Y,
                                  style.PAGE_W, style.PAGE_H)
            continue
        ph._element.getparent().remove(ph._element)
    layout.shapes._spTree.append(background)

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    print("template:", out_path)
    print("  master:", spec.master_name)
    print("  layout:", spec.layout_name)
    print("  slide size:", w, "x", h, "EMU")
    return out_path


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--brand", required=True)
    ap.add_argument("--out", required=True, type=Path)
    ap.add_argument("--variant", default=None,
                    help="build a SECOND master from this brand's contract, "
                         "with ' (NAME)' appended to the master and layout "
                         "names and the ground painted in the brand's PANEL")
    args = ap.parse_args()
    spec = load_brand(args.brand)
    if args.variant:
        spec = variant_spec(spec, args.variant)
    build(spec, args.out, args.variant)
