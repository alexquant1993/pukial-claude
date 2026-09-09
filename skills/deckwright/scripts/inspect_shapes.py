"""Report a deck's shapes: what a text map is keyed on.

The map is keyed on shape id resolved through groups, so the ids have to come
from somewhere. This is that somewhere.

Run: PYTHONIOENCODING=utf-8 uv run --offline --no-project --with python-pptx \
       --with pillow python scripts/inspect_shapes.py --deck deck.pptx --slide 14
"""

import argparse
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from deck_kit.geometry import EMU

A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def walk(shapes, prefix=""):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for row in walk(shape.shapes, prefix + shape.name + "/"):
                yield row
        else:
            yield prefix, shape


def describe(shape):
    text = ""
    field = False
    if shape.has_text_frame:
        text = shape.text_frame.text.strip().replace("\n", " ")[:60]
        field = shape.text_frame._txBody.find(".//{%s}fld" % A) is not None
    ph = shape.placeholder_format.type if shape.is_placeholder else None
    return ("id=%-6d %-28s %-14s x=%-5s y=%-5s w=%-5s h=%-5s ph=%-14s fld=%-5s %r"
            % (shape.shape_id, shape.name[:28], str(shape.shape_type)[:14],
               shape.left // EMU if shape.left is not None else "-",
               shape.top // EMU if shape.top is not None else "-",
               shape.width // EMU if shape.width is not None else "-",
               shape.height // EMU if shape.height is not None else "-",
               str(ph)[:14], field, text))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--deck", type=Path, required=True)
    ap.add_argument("--slide", type=int, help="1-based; omit for all")
    args = ap.parse_args()
    prs = Presentation(str(args.deck))
    for number, slide in enumerate(prs.slides, 1):
        if args.slide and number != args.slide:
            continue
        print("=== slide %d%s" % (number,
              " (hidden)" if slide._element.get("show") == "0" else ""))
        for prefix, shape in walk(slide.shapes):
            print("  %s%s" % (prefix, describe(shape)))


if __name__ == "__main__":
    main()
