"""Dump every paragraph of a deck as stable, diffable text.

Before regenerating anything, dump the text of the author's current file and
diff it against the last known version, to discover what they changed by hand.
Regenerating from a stale assumption destroys those edits, and there is no
undo.

Run: PYTHONIOENCODING=utf-8 uv run --offline --no-project --with python-pptx \
       --with pillow python scripts/dump_text.py --deck deck.pptx --out v3.txt
"""

import argparse
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from pptx import Presentation

from deck_kit.merge import refuse_in_place
from deck_kit.pagenums import is_hidden
from deck_kit.textedit import iter_shapes, paragraph_text


def dump(deck):
    prs = Presentation(str(deck))
    lines = []
    for number, slide in enumerate(prs.slides, 1):
        hidden = is_hidden(slide)
        for shape in iter_shapes(slide.shapes):
            if not shape.has_text_frame:
                continue
            for index, para in enumerate(shape.text_frame.paragraphs):
                text = paragraph_text(para).replace("\t", " ").replace("\n", " ")
                lines.append("%d%s\t%d\t%d\t%s"
                             % (number, " (hidden)" if hidden else "",
                                shape.shape_id, index, text))
    return lines


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--deck", type=Path, required=True)
    ap.add_argument("--out", type=Path)
    # Exact, not a floor: a step that only proves the dump exited 0 passes on a
    # deck that dumps nothing, which is the one thing it exists to disprove.
    ap.add_argument("--expect-paragraphs", type=int, default=None,
                    help="exact number of paragraphs the dump must contain")
    args = ap.parse_args()
    if args.out is not None:
        refuse_in_place(args.out, args.deck)
    lines = dump(args.deck)
    body = "\n".join(lines) + "\n"
    if args.out:
        args.out.write_text(body, encoding="utf-8")
        print("%d paragraph(s) -> %s" % (len(lines), args.out))
    else:
        print(body, end="")
    if args.expect_paragraphs is not None and len(lines) != args.expect_paragraphs:
        print("FAIL: expected %d paragraph(s), found %d"
              % (args.expect_paragraphs, len(lines)))
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
