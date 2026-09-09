"""Content lint for a deck: what must never reach the client, what must always be there.

A different tool class from file QA (check_pptx.ps1) and visual QA
(compose_qa.py): this reads the text. Every check is driven by a TOML rules
file, reports how many rules it ran with, and a check with no rules is
SKIPPED, never passed - `--expect-checks N` makes the gate assert that
every check it relies on actually ran.

Checks, in order: typos, codes, forbidden, required, kickers, notes_only,
exact, assets. See references/05-qa.md for the rules file.

`--deck` takes **either a built .pptx or a directory of slide drafts**. An
HTML deck is a deliverable in its own right (the spec's `**Output:**` line),
and a deliverable nobody lints is a deliverable nobody checked. On a drafts
directory the text is harvested from the markup with `html.parser`, one
"slide" per `sNN-*.html` in slide order, and every check runs except
`notes_only`, which is skipped with a reason: a drafts directory carries no
speaker notes for a string to be in. `assets` is unchanged - it checks the
brand, not the deck.

`--brand NAME` with no `--deck` runs the `[assets]` check alone, which is
what the "Set up a brand" route needs: the manifest is written two stages
before any deck exists.

Run: uv run --offline --no-project --with python-pptx --with pillow python scripts/lint_deck.py --deck out/walk/deck.pptx --rules examples/walkthrough/lint.toml --expect-checks 8
     ... python scripts/lint_deck.py --deck examples/ai-horizon/drafts --rules examples/ai-horizon/lint.toml --expect-checks 7
     ... python scripts/lint_deck.py --brand relay
"""

import argparse
import hashlib
import re
import sys
import tomllib
from collections import namedtuple
from html.parser import HTMLParser
from pathlib import Path

from pptx import Presentation
from pptx.exc import PackageNotFoundError

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_html_deck import DraftError, slide_files  # noqa: E402
from deck_kit.brand import brands_dir  # noqa: E402
from deck_kit.notes import read_notes  # noqa: E402
from deck_kit.textedit import iter_paragraphs, paragraph_text  # noqa: E402

Finding = namedtuple("Finding", "check position detail")
CheckResult = namedtuple("CheckResult", "name rules findings skip_reason")
# A check is skipped for exactly two reasons: it has no rules, or the deck
# under test cannot carry what it checks. The second needs saying out loud -
# "skipped (no rules)" on a drafts directory would read as a rules-file bug.
CheckResult.__new__.__defaults__ = (None,)


def slide_texts(prs):
    """{position: joined paragraph text} for every slide, hidden ones included."""
    return {
        position: "\n".join(paragraph_text(p) for p in iter_paragraphs(slide))
        for position, slide in enumerate(prs.slides, 1)
    }


# Elements whose text runs on within its parent's line. Everything else ends
# a line, so a kicker written as its own <div> stays a paragraph of its own -
# which is what the [kickers] pattern matches against.
INLINE_TAGS = {"a", "abbr", "b", "code", "del", "em", "i", "ins", "kbd",
               "mark", "q", "s", "samp", "small", "span", "strong", "sub",
               "sup", "time", "u", "var"}
# Never read: <title> is chrome, and <style>/<script> are not what the
# audience sees. A draft carries no <script> at all (tests/test_html_stage.py).
SILENT_TAGS = {"title", "style", "script", "head"}


class _DraftText(HTMLParser):
    """The visible text of one slide draft, one line per block element.

    Deliberately literal: it reads what the markup says and applies no CSS.
    A design system sets its labels uppercase in a stylesheet, so the source
    of a kicker reads `01. How to read this` where the slide reads
    `01. HOW TO READ THIS` - which is why the [kickers] check case-folds on a
    drafts directory and nothing else does. Resolving `text-transform` here
    would mean implementing selector matching and cascade order, and getting
    it subtly wrong would silently change what `[exact]` compares.
    """

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.lines = []
        self._silent = 0
        self._current = []

    def _break(self):
        line = " ".join("".join(self._current).split())
        if line:
            self.lines.append(line)
        self._current = []

    def handle_starttag(self, tag, attrs):
        low = tag.lower()
        if low in SILENT_TAGS:
            self._silent += 1
        elif low not in INLINE_TAGS:
            self._break()

    def handle_startendtag(self, tag, attrs):
        if tag.lower() == "br":
            self._break()

    def handle_endtag(self, tag):
        low = tag.lower()
        if low in SILENT_TAGS:
            self._silent = max(0, self._silent - 1)
        elif low not in INLINE_TAGS:
            self._break()

    def handle_data(self, data):
        if not self._silent:
            self._current.append(data)

    def text(self):
        self._break()
        return "\n".join(self.lines)


def draft_text(path):
    parser = _DraftText()
    parser.feed(Path(path).read_text(encoding="utf-8"))
    parser.close()
    return parser.text()


def draft_texts(directory):
    """{position: text} for a directory of drafts, one position per file.

    Positions are the slide numbers `sNN-*.html` declares, renumbered 1..N in
    file order, so `slide=3` in a finding means the third slide of the deck
    exactly as it does for a .pptx.
    """
    return {position: draft_text(path)
            for position, path in enumerate(slide_files(directory), 1)}


def _flat(text):
    return " ".join(text.split())


def _rule_list(name, rules, key, item_type=str):
    """A list of `item_type` from the rules table, or [] when absent.

    A scalar (e.g. `needles = "ASSESSEMENT"`) is a rules-file bug: it would
    otherwise be silently iterated character by character and inflate the
    reported rule count.
    """
    value = rules.get(key, [])
    ok = isinstance(value, list) and all(
        isinstance(v, item_type) and not isinstance(v, bool) for v in value)
    if not ok:
        kind = "strings" if item_type is str else "ints"
        raise ValueError("rules table %s: %s must be a list of %s" % (name, key, kind))
    return value


def _compile(check_name, pattern):
    try:
        return re.compile(pattern)
    except re.error as e:
        raise ValueError("invalid regex in %s: %r (%s)" % (check_name, pattern, e))


# --- checks: each takes (texts, notes, rules_table, ctx) and returns (rules_count, findings)

def check_typos(texts, notes, rules, ctx):
    needles = _rule_list("typos", rules, "needles")
    found = [Finding("typos", pos, n) for pos, t in texts.items()
             for n in needles if n in _flat(t)]
    return len(needles), found


def check_codes(texts, notes, rules, ctx):
    patterns = [_compile("codes", p) for p in _rule_list("codes", rules, "patterns")]
    found = []
    for pos, t in texts.items():
        flat = _flat(t)
        for pat in patterns:
            for hit in sorted(set(pat.findall(flat))):
                found.append(Finding("codes", pos, hit))
    return len(patterns), found


def check_forbidden(texts, notes, rules, ctx):
    # `tokens_file` is the half this check could not carry. The rules file
    # travels with the deck, so listing a client's own name in it writes that
    # name into the one file it must not be in - which meant the check could
    # enforce "no vendor names" and not "no client names", the half that
    # matters. A path lets the list live somewhere the deck directory does
    # not: a home-directory file, an environment-supplied path. One token per
    # line; blank lines and lines starting with # are ignored. Resolved
    # against the rules file's own directory when relative, so an absolute
    # path outside the tree is the normal case rather than a special one.
    tokens = list(_rule_list("forbidden", rules, "tokens"))
    path = _rule_str("forbidden", rules, "tokens_file")
    if path:
        resolved = Path(path).expanduser()
        if not resolved.is_absolute():
            resolved = ctx["rules_dir"] / resolved
        tokens += [line.strip() for line in
                   resolved.read_text(encoding="utf-8").splitlines()
                   if line.strip() and not line.startswith("#")]
    found = [Finding("forbidden", pos, tok) for pos, t in texts.items()
             for tok in tokens if tok.lower() in _flat(t).lower()]
    return len(tokens), found


def check_required(texts, notes, rules, ctx):
    strings = _rule_list("required", rules, "strings")
    exempt = set(_rule_list("required", rules, "exempt", item_type=int))
    found = [Finding("required", pos, s) for pos, t in texts.items() if pos not in exempt
             for s in strings if _flat(s) not in _flat(t)]
    return len(strings), found


def check_kickers(texts, notes, rules, ctx):
    # _rule_str, not rules.get: a list here (`pattern = ["a"]`) reached
    # re.compile and came out as a raw TypeError, which is the one thing
    # references/05-qa.md promises no bad input does.
    pattern = _rule_str("kickers", rules, "pattern")
    allowed = _rule_list("kickers", rules, "allowed")
    if not pattern or not allowed:
        return 0, []
    pat = _compile("kickers", pattern)
    exempt = set(_rule_list("kickers", rules, "exempt", item_type=int))
    # A kicker is a label, and a design system sets its labels uppercase in
    # CSS: `<div class="kicker">01. How to read this</div>` reads
    # `01. HOW TO READ THIS` on the slide and is written uppercase into the
    # PPTX. The markup carries the source case, so on a drafts directory the
    # pattern and the allowlist are matched against an uppercased copy. This
    # is the only check that folds case on drafts; `exact` and `required`
    # must not, because they compare strings the spec fixed.
    # tests/test_lint_deck.py::test_a_draft_kicker_uppercased_by_css_is_still_matched.
    fold = ctx.get("drafts", False)
    found = []
    for pos, t in texts.items():
        if pos in exempt:
            continue
        # Kickers are paragraph-level and the pattern relies on the line
        # boundary, so match against the newline-joined text, not _flat(t).
        kickers = [k.strip() for k in pat.findall(t.upper() if fold else t)]
        if not kickers:
            found.append(Finding("kickers", pos, "no kicker"))
            continue
        if len(kickers) > 1:
            found.append(Finding("kickers", pos, "%d kickers on one slide: %r" % (len(kickers), kickers)))
        for k in kickers:
            if k not in allowed:
                found.append(Finding("kickers", pos, "kicker not allowed: %r" % k))
    return len(allowed), found


def check_notes_only(texts, notes, rules, ctx):
    strings = _rule_list("notes_only", rules, "strings")
    found = []
    for s in strings:
        on_slides = [pos for pos, t in texts.items() if _flat(s).lower() in _flat(t).lower()]
        in_notes = any(_flat(s).lower() in _flat(" ".join(paras)).lower()
                       for paras in notes.values())
        for pos in on_slides:
            found.append(Finding("notes_only", pos, s))
        if not in_notes:
            found.append(Finding("notes_only", None, "%s: not in any note" % s))
    return len(strings), found


TEXT_SUFFIXES = {".py", ".md", ".toml", ".css", ".txt", ".json", ".html"}


def _rule_str(name, rules, key):
    """A scalar string from the rules table, or None when absent.

    A list where a scalar is expected (e.g. `spec = ["spec.md"]`) is a
    rules-file bug rather than something to silently misinterpret.
    """
    value = rules.get(key)
    if value is None:
        return None
    if not isinstance(value, str):
        raise ValueError("rules table %s: %s must be a string" % (name, key))
    return value


def exact_strings_from_spec(path):
    """Column `Exact string` of the first markdown table whose header names it."""
    lines = Path(path).read_text(encoding="utf-8").splitlines()
    for i, line in enumerate(lines):
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if "Exact string" not in cells:
            continue
        col = cells.index("Exact string")
        out = []
        for row in lines[i + 2:]:            # skip the |---| separator
            if not row.strip().startswith("|"):
                break
            rc = [c.strip() for c in row.strip().strip("|").split("|")]
            if len(rc) > col and rc[col]:
                out.append(rc[col].strip("`"))
        return out
    raise ValueError("%s has no table with an 'Exact string' column" % path)


def check_exact(texts, notes, rules, ctx):
    strings = list(_rule_list("exact", rules, "strings"))
    spec = _rule_str("exact", rules, "spec")
    if spec:
        strings += exact_strings_from_spec(ctx["rules_dir"] / spec)
    flat = {pos: _flat(t) for pos, t in texts.items()}
    found = [Finding("exact", None, "not found verbatim on any slide: %r" % s)
             for s in strings if not any(_flat(s) in t for t in flat.values())]
    return len(strings), found


def _assets(brand_dir):
    """Every file that is not text and not a bytecode cache. Extension-agnostic."""
    return sorted(p for p in brand_dir.rglob("*")
                  if p.is_file() and p.suffix.lower() not in TEXT_SUFFIXES
                  and "__pycache__" not in p.parts)


def _sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _validated_manifest(entries, name):
    """{path: entry} from the manifest's `[[asset]]` blocks, or raise.

    Every block must be a table with string `path`, `sha256` and
    `provenance` - a missing or wrong-typed field would otherwise crash the
    linter with a raw traceback (`KeyError` from `a["path"]`,
    `AttributeError` from `.get("provenance", "").strip()` on a non-string).
    A `path` listed twice is also refused: last-wins would let a second,
    well-provenanced block mask a first block with none.
    """
    listed = {}
    for i, a in enumerate(entries):
        ok = (isinstance(a, dict)
              and isinstance(a.get("path"), str)
              and isinstance(a.get("sha256"), str)
              and isinstance(a.get("provenance"), str))
        if not ok:
            raise ValueError(
                "assets.toml in %s: entry %d needs string path, sha256 and provenance" % (name, i))
        if a["path"] in listed:
            raise ValueError("assets.toml in %s: duplicate path %s" % (name, a["path"]))
        listed[a["path"]] = a
    return listed


def check_assets(texts, notes, rules, ctx):
    name = _rule_str("assets", rules, "brand")
    if not name:
        return 0, []
    brand_dir = brands_dir() / name
    manifest = brand_dir / "assets.toml"
    if not manifest.exists():
        return 1, [Finding("assets", None, "no assets.toml in %s" % name)]
    with open(manifest, "rb") as fh:
        listed = _validated_manifest(tomllib.load(fh).get("asset", []), name)
    found = []
    on_disk = {p.relative_to(brand_dir).as_posix(): p for p in _assets(brand_dir)}
    for rel, p in on_disk.items():
        if rel not in listed:
            found.append(Finding("assets", None, "not in assets.toml: %s" % rel))
            continue
        if listed[rel].get("sha256") != _sha256(p):
            found.append(Finding("assets", None, "sha256 mismatch: %s" % rel))
        if not listed[rel].get("provenance", "").strip():
            found.append(Finding("assets", None, "empty provenance: %s" % rel))
    for rel in listed:
        if rel not in on_disk:
            found.append(Finding("assets", None, "listed but missing: %s" % rel))
    return 1, found


def manifest_text(name):
    brand_dir = brands_dir() / name
    blocks = ['[[asset]]\npath = "%s"\nsha256 = "%s"\nprovenance = ""\n'
              % (p.relative_to(brand_dir).as_posix(), _sha256(p)) for p in _assets(brand_dir)]
    return "\n".join(blocks)


CHECKS = {
    "typos": check_typos,
    "codes": check_codes,
    "forbidden": check_forbidden,
    "required": check_required,
    "kickers": check_kickers,
    "notes_only": check_notes_only,
}
CHECKS["exact"] = check_exact
CHECKS["assets"] = check_assets


def lint_brand(name):
    """The `[assets]` check alone, against a brand and no deck.

    "Set up a brand" step 4 says to record every binary's provenance and run
    the content lint with `[assets]` set - but `--deck` and `--rules` were
    both required, so the manifest could not be verified until a deck and a
    rules file existed, two stages later. The check never read the deck: it
    walks the brand directory and its assets.toml.
    tests/test_lint_deck.py::test_the_assets_check_runs_against_a_brand_with_no_deck.
    """
    count, findings = check_assets({}, {}, {"brand": name}, {"rules_dir": Path.cwd()})
    return [CheckResult("assets", count, findings)]


# What a drafts directory cannot carry, and why. Named here rather than
# inside the check so the reason reaches the report.
DRAFTS_SKIP = {
    "notes_only": "a drafts directory carries no speaker notes",
}


def lint(deck_path, rules_path):
    rules_path = Path(rules_path)
    with open(rules_path, "rb") as fh:
        rules = tomllib.load(fh)
    drafts = deck_path is not None and Path(deck_path).is_dir()
    if deck_path is None:
        # --brand with no --deck: the [assets] check alone, which is what
        # "Set up a brand" needs two stages before a deck exists.
        texts, notes = {}, {}
    elif drafts:
        texts, notes = draft_texts(deck_path), {}
    else:
        prs = Presentation(str(deck_path))
        texts, notes = slide_texts(prs), read_notes(prs)
    ctx = {"rules_dir": rules_path.parent,
           "deck": None if deck_path is None else Path(deck_path),
           "drafts": drafts}
    results = []
    for name, fn in CHECKS.items():
        table = rules.get(name, {})
        if not isinstance(table, dict):
            raise ValueError("rules table %s must be a table" % name)
        count, findings = fn(texts, notes, table, ctx)
        reason = None
        if drafts and count and name in DRAFTS_SKIP:
            count, findings, reason = 0, [], DRAFTS_SKIP[name]
        results.append(CheckResult(name, count, findings, reason))
    return results


def report(results, expect_checks=None):
    """Print the per-check report and return the exit code."""
    total = 0
    ran = 0
    for r in results:
        if r.rules == 0:
            print("check %s: skipped (%s)" % (r.name, r.skip_reason or "no rules"))
            continue
        ran += 1
        print("check %s: rules=%d findings=%d" % (r.name, r.rules, len(r.findings)))
        for f in r.findings:
            where = "deck" if f.position is None else "slide=%d" % f.position
            print("FAIL %s %s: %s" % (f.check, where, f.detail))
        total += len(r.findings)
    skipped = len(results) - ran
    code = 0
    if expect_checks is not None and ran != expect_checks:
        print("FAIL: expected %d check(s) to run, ran=%d skipped=%d" % (expect_checks, ran, skipped))
        code = 1
    if total:
        print("LINT FAILED findings=%d ran=%d skipped=%d" % (total, ran, skipped))
        code = 1
    if code == 0:
        print("LINT PASSED ran=%d skipped=%d" % (ran, skipped))
    return code


def main(argv=None):
    argv = sys.argv[1:] if argv is None else list(argv)
    if argv[:1] == ["manifest"]:
        ap = argparse.ArgumentParser(prog="lint_deck.py manifest")
        ap.add_argument("--brand", required=True)
        args = ap.parse_args(argv[1:])
        sys.stdout.write(manifest_text(args.brand))
        return 0
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--deck", type=Path, default=None,
                    help="a built .pptx, or a directory of sNN-*.html drafts")
    ap.add_argument("--rules", type=Path, default=None)
    ap.add_argument("--brand", default=None,
                    help="run the [assets] check alone against this brand, "
                         "with no deck and no rules file")
    ap.add_argument("--expect-checks", type=int, default=None,
                    help="exact number of checks that must run with at least one rule")
    args = ap.parse_args(argv)
    if args.brand and args.deck is None:
        try:
            results = lint_brand(args.brand)
        except (ValueError, OSError, tomllib.TOMLDecodeError) as e:
            print("FAIL: %s" % e)
            return 1
        return report(results, args.expect_checks)
    if args.deck is None or args.rules is None:
        print("FAIL: --deck and --rules are both required "
              "(or --brand alone for the [assets] check)")
        return 1
    try:
        results = lint(args.deck, args.rules)
    # PackageNotFoundError is python-pptx's own, raised by Presentation() on a
    # deck that is missing or is not a readable package. It does not derive
    # from OSError, so without it a mistyped --deck path comes out as a raw
    # traceback rather than the FAIL: line every other bad input produces.
    # DraftError is build_html_deck's, for a --deck directory that holds no
    # sNN-*.html drafts; it is a ValueError, and named here so it reads.
    except (DraftError, ValueError, FileNotFoundError, OSError,
            PackageNotFoundError, tomllib.TOMLDecodeError) as e:
        print("FAIL: %s" % e)
        return 1
    return report(results, args.expect_checks)


if __name__ == "__main__":
    sys.exit(main())
