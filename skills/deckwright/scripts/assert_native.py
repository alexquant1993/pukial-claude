"""Assert what the built package actually contains.

Doctrines have to be checked against the artifact, not against the source. A
grep of *.py cannot see autofit that arrives from a library default, and it
cannot see a picture that a builder placed. Everything here opens the package
and looks.

  --autoshapes N     exact count of autoshapes plus freeforms
  --min-freeforms N  at least N freeform shapes, so a curve cannot silently
                     vanish inside the autoshape bucket
  --max-pictures N   the raster boundary: PNG only where OOXML cannot express
                     the mark, so the count is known and bounded

Any slide-sized picture fails outright: that deck is an image of a deck.
Any autofit on any text frame fails: measure-then-place is defeated the moment
PowerPoint is allowed to resize a box.

Run: PYTHONIOENCODING=utf-8 uv run --offline --no-project --with python-pptx \
       --with pillow python scripts/assert_native.py out/ai-horizon/deck.pptx \
       --autoshapes 145 --min-pictures 32 --max-pictures 32 --no-dangling

Those are examples/ai-horizon/run.ps1's own numbers, derived in
examples/ai-horizon/build_ai_horizon.py's docstring before the deck is built.
Every count here is exact on purpose, and `rule` is an `rrect`, so hairlines
are most of the autoshape count on a Relay deck - references/03-pptx-stage.md,
"The vocabulary", says which primitive lands in which bucket.
"""

import argparse
import re
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

AUTOFIT = re.compile(r"<a:(spAutoFit|normAutofit)\b")


def walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            for inner in walk(sh.shapes):
                yield inner


def count_autofit(path):
    """Autofit anywhere in any slide part, counted in the XML."""
    hits = {}
    with zipfile.ZipFile(str(path)) as z:
        for name in z.namelist():
            if name.startswith("ppt/slides/slide") and name.endswith(".xml"):
                n = len(AUTOFIT.findall(z.read(name).decode("utf-8")))
                if n:
                    hits[name] = n
    return hits


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("deck", type=Path)
    ap.add_argument("--autoshapes", type=int, required=True,
                    help="exact expected count of autoshapes plus freeforms")
    ap.add_argument("--min-freeforms", type=int, default=0)
    ap.add_argument("--max-pictures", type=int, required=True)
    ap.add_argument("--expect-fonts", type=int, default=None,
                    help="exact count of embedded .fntdata parts")
    ap.add_argument("--min-pictures", type=int, default=0,
                    help="at least N pictures - the floor that stops the "
                         "transplanted media vanishing under a ceiling")
    ap.add_argument("--no-dangling", action="store_true",
                    help="fail if any relationship or r:id does not resolve")
    args = ap.parse_args()

    prs = Presentation(str(args.deck))
    autoshapes = freeforms = pictures = full_bleed = 0
    for slide in prs.slides:
        for sh in walk(slide.shapes):
            if sh.shape_type == MSO_SHAPE_TYPE.FREEFORM:
                freeforms += 1
                autoshapes += 1
            elif sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                autoshapes += 1
            elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures += 1
                if (sh.width >= prs.slide_width * 0.98
                        and sh.height >= prs.slide_height * 0.98):
                    full_bleed += 1

    autofit = count_autofit(args.deck)
    print("slides=%d autoshapes=%d (freeforms=%d) pictures=%d full-bleed=%d autofit=%d"
          % (len(prs.slides), autoshapes, freeforms, pictures, full_bleed,
             sum(autofit.values())))

    failures = []
    if full_bleed:
        failures.append("%d slide-sized picture(s) - this deck is an image" % full_bleed)
    if autofit:
        failures.append("autofit present in %s - measure-then-place is defeated"
                        % ", ".join("%s x%d" % (k.rsplit("/", 1)[-1], v)
                                    for k, v in sorted(autofit.items())))
    if autoshapes != args.autoshapes:
        failures.append("%d autoshapes, expected exactly %d" % (autoshapes, args.autoshapes))
    if freeforms < args.min_freeforms:
        failures.append("%d freeforms, expected at least %d" % (freeforms, args.min_freeforms))
    if pictures > args.max_pictures:
        failures.append("%d pictures, expected at most %d - the raster boundary is "
                        "conic fills and icons only" % (pictures, args.max_pictures))

    from deck_kit.fonts import font_parts
    from deck_kit.merge import dangling_refs, read_package
    parts = read_package(args.deck)
    if args.expect_fonts is not None:
        n = len(font_parts(parts))
        print("embedded font parts: %d" % n)
        if n != args.expect_fonts:
            failures.append("%d embedded font part(s), expected exactly %d"
                            % (n, args.expect_fonts))
    if pictures < args.min_pictures:
        failures.append("%d pictures, expected at least %d - the donor's media "
                        "did not survive the transplant" % (pictures, args.min_pictures))
    if args.no_dangling:
        problems = dangling_refs(parts)
        if problems:
            failures.append("%d unresolved reference(s): %s"
                            % (len(problems), "; ".join(problems[:5])))

    for f in failures:
        print("FAIL:", f)
    if failures:
        return 1
    print("PASS")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
