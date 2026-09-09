"""The versioned build driver.

The author edits the PPTX in PowerPoint between versions, so a generated deck
is never the deliverable and the author's file is never overwritten. Each
version is a numbered, resumable pipeline: copy the author's file into a work
directory, assert it is the file the build was written against, then run a
sequence of steps. Every step is written to disk as stepN_<name>.pptx, so a
failure is diagnosable at the step that caused it rather than at the end.

The zip-surgery stages run as separate subprocesses: a crash inside zip
surgery cannot then corrupt an in-memory Presentation, and the intermediate
on disk is still openable.

    v = Version(source=author_file, work=Path("work/v3"), final=Path("v3.pptx"))
    v.expect(slide_count=30, fingerprint={14: "Fraud", 23: "Market"})
    v.step("patch text", lambda prs: patch_text(prs, PATCHES))
    v.step("delete", lambda prs: delete_positions(prs, [23, 14]), slides=28)
    v.subprocess_step("merge", [sys.executable, "-m", "...", ...], slides=30)
    v.finalise()
"""

import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from pptx import Presentation

from deck_kit.deck import assert_fingerprint


class Version:
    def __init__(self, source, work, final):
        source, work, final = Path(source), Path(work), Path(final)
        if final.resolve() == source.resolve():
            raise ValueError(
                "refusing to overwrite the source deck %s - the author's file "
                "is frequently open in PowerPoint and locked, and a rebuild "
                "from a stale assumption would destroy their manual edits"
                % source)
        work.mkdir(parents=True, exist_ok=True)
        self.source, self.work, self.final = source, work, final
        self.n = 0
        self.current = work / "step_input.pptx"
        shutil.copyfile(str(source), str(self.current))
        self.log = []

    def _next(self, name):
        path = self.work / ("step%d_%s.pptx" % (self.n, name.replace(" ", "_")))
        self.n += 1
        return path

    def expect(self, slide_count, fingerprint):
        """Assert this is the file the build was written against.

        Not a checksum - the author edits legitimately between versions. What
        must hold is the slide count and what sits at each position the build
        is about to touch.
        """
        assert_fingerprint(Presentation(str(self.current)), slide_count, fingerprint)
        self.log.append(("expect", slide_count, sorted(fingerprint)))

    def step(self, name, fn, slides=None):
        """One in-process step: open, mutate, save to the next stepN file.

        `slides` is the slide count this step must leave behind, asserted
        after `fn` runs and before the result is saved. It is optional in the
        signature but not in practice for any step that can change the slide
        count - `delete_positions` and `merge` both do, and a silent
        off-by-one there is invisible until someone presents the deck.
        """
        out = self._next(name)
        prs = Presentation(str(self.current))
        fn(prs)
        if slides is not None:
            if len(prs.slides) != slides:
                raise AssertionError(
                    "step %r left %d slides, declared %d"
                    % (name, len(prs.slides), slides))
        prs.save(str(out))
        self.current = out
        self.log.append((name, str(out)))
        print("  step %-24s -> %s" % (name, out.name))
        return out

    def subprocess_step(self, name, argv, slides=None):
        """One out-of-process step.

        The contract every wrapper honours: `argv` gets the input path and the
        output path appended, in that order, after whatever options it
        already carries. So `[python, _merge_step.py, "--donor", D]` becomes
        `python _merge_step.py --donor D <in> <out>` - argparse parses an
        optional followed by two positionals cleanly.

        `slides` is the slide count this step must leave behind. The output
        is re-opened to count it, since the mutation ran out of process.
        """
        out = self._next(name)
        subprocess.check_call([str(a) for a in argv] + [str(self.current), str(out)])
        if not out.exists():
            raise AssertionError("%s produced no output" % name)
        if slides is not None:
            got = len(Presentation(str(out)).slides)
            if got != slides:
                raise AssertionError(
                    "step %r left %d slides, declared %d" % (name, got, slides))
        self.current = out
        self.log.append((name, str(out)))
        print("  step %-24s -> %s (subprocess)" % (name, out.name))
        return out

    def finalise(self):
        shutil.copyfile(str(self.current), str(self.final))
        print("  final -> %s" % self.final)
        return self.final
