"""Editing text in a deck someone else has been editing by hand.

Two complementary mechanisms, and they are for different jobs.

Character-level run attribution - `patch_paragraph` - is for surgical edits to
a file the author has since hand-edited. PowerPoint fragments a paragraph into
arbitrary runs by spellcheck state, language tags and edit history. The naive
fix, writing the new string into run 0 and blanking the rest, silently
destroys mid-sentence bold and colour. The real fix builds a character-to-run
ownership array from the old text, runs a sequence matcher, routes unchanged
characters back to their owning run, drops deleted ones, attaches replacements
and insertions to the run owning the first affected old character, and asserts
the reassembly.

The text map - `apply_text_map` - is for wholesale retexting of an approved
template, and lives in the second half of this module.
"""

import copy
import difflib

from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def q(tag, ns=A):
    return "{%s}%s" % (ns, tag)


def iter_shapes(shapes):
    """Every leaf shape, recursing through groups.

    Groups are the normal state of a deck a human has edited, and a walker
    that stops at the group is a walker that silently misses half the text.
    """
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            for inner in iter_shapes(sh.shapes):
                yield inner
        else:
            yield sh


def iter_paragraphs(slide):
    """Every paragraph on the slide, including the ones inside table cells."""
    for sh in iter_shapes(slide.shapes):
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                yield p
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for cell in row.cells:
                    for p in cell.text_frame.paragraphs:
                        yield p


def paragraph_text(para):
    """The paragraph as its runs see it.

    Not `text_frame.text` and not `paragraph.text`: those interpolate line
    breaks and other non-run content, so a patch matched against them would
    be matched against a string the runs cannot reproduce.
    """
    return "".join(r.text for r in para.runs)


def patch_paragraph(para, new):
    """Rewrite the paragraph to `new`, preserving per-run formatting.

    Characters that survive go back to the run that owned them. Insertions and
    replacements attach to the run owning the first affected old character,
    which is what keeps a tail edit out of a bold run in the middle.
    """
    runs = para.runs
    if not runs:
        raise ValueError(
            "paragraph has no runs to attribute characters to - an empty "
            "paragraph between two bullets is ordinary in a hand-edited deck, "
            "and the engagement's version indexed owner[-1] on an empty list")
    old = paragraph_text(para)
    if old == "":
        # Runs exist but hold no text - <a:r><a:t/></a:r> is ordinary in a
        # hand-edited deck. `owner` built from run lengths below would be
        # empty, and the replace/insert branch's `owner[-1]` fallback would
        # then IndexError on an empty list - the exact failure the guard
        # above's message claims to have already fixed.
        runs[0].text = new
        if paragraph_text(para) != new:
            raise AssertionError(
                "run reassembly does not reproduce the target text: %r != %r"
                % (paragraph_text(para), new))
        return
    owner = []
    for i, r in enumerate(runs):
        owner += [i] * len(r.text)
    out = [[] for _ in runs]
    matcher = difflib.SequenceMatcher(None, old, new, autojunk=False)
    for op, i1, i2, j1, j2 in matcher.get_opcodes():
        if op == "equal":
            for i in range(i1, i2):
                out[owner[i]].append(old[i])
        elif op == "delete":
            continue
        else:                      # replace / insert
            idx = owner[i1] if i1 < len(old) else owner[-1]
            out[idx].extend(new[j1:j2])
    for r, chars in zip(runs, out):
        r.text = "".join(chars)
    if paragraph_text(para) != new:
        raise AssertionError(
            "run reassembly does not reproduce the target text: %r != %r"
            % (paragraph_text(para), new))


def patch_text(prs, patches):
    """Apply (slide_number, old, new) patches, each matching exactly once.

    Zero matches and two matches are both hard failures. Zero means the deck
    moved under the patch - the author edited that sentence - and applying
    nothing would ship the old text silently. Two means the patch is
    ambiguous and would change the wrong one as soon as either diverges.

    Resolve-all-then-apply: every patch's single match is found first, and
    the first ambiguity raises before ANY patch is mutated. A single pass
    that validates and mutates one patch at a time leaves the deck
    half-patched the moment a LATER patch turns out to be ambiguous - the
    earlier ones are already rewritten by the time the SystemExit fires, and
    nothing but the caller dying before prs.save() stood between that and a
    corrupted deck on disk.
    """
    slides = list(prs.slides)
    resolved = []
    for number, old, new in patches:
        matches = [p for p in iter_paragraphs(slides[number - 1])
                   if paragraph_text(p).strip() == old.strip()]
        if len(matches) != 1:
            raise SystemExit("patch slide %d hits=%d: %s"
                              % (number, len(matches), old[:70]))
        resolved.append((matches[0], old, new))
    for para, old, new in resolved:
        text = paragraph_text(para)
        patch_paragraph(para, text.replace(old.strip(), new.strip()))
    return len(resolved)


# --------------------------------------------------------------- the text map
#
# A pure-data map against a pure-mechanism applier. The map is keyed on shape
# id resolved recursively through groups - not on name, because names repeat,
# and not on position or text, because both change under the author's hand.
#
# Three-valued protocol per paragraph:
#   str        collapse the paragraph to a single run
#   list[str]  assign run by run, preserving each run's formatting
#   None       delete the paragraph
#
# The applier never raises. It accumulates every mismatch and returns them
# together, because a map has dozens of entries and one hard failure per round
# trip is a day of round trips.


def shapes_by_id(slide):
    """Every leaf shape on the slide, keyed by shape id, recursing through
    groups. The id is the only stable handle: names repeat, and position and
    text both change under the author's hand."""
    return {sh.shape_id: sh for sh in iter_shapes(slide.shapes)}


def _strip_err(r_element):
    """A deep-copied run carries the spellcheck attributes with it, and
    PowerPoint then underlines text nobody has checked.

    `err` and `dirty` are attributes of <a:rPr>, not of <a:r> - <a:r> has no
    attributes at all. Measured across three human-edited decks in the
    engagement: 198 occurrences on rPr, zero on r. Stripping the wrong
    element gives a doctrine whose test can only pass.
    """
    rpr = r_element.find(q("rPr"))
    if rpr is not None:
        rpr.attrib.pop("err", None)
        rpr.attrib.pop("dirty", None)
    return r_element


def append_runs(para, texts, proto_index=-1):
    """Clone a prototype run and append one run per string.

    New runs must be inserted before <a:endParaRPr> or PowerPoint ignores
    them entirely - the text is in the file and never appears on the slide.
    """
    runs = para.runs
    if not runs:
        raise AssertionError("append_runs needs at least one run to clone")
    proto = runs[proto_index]._r
    end = para._p.find(q("endParaRPr"))
    made = []
    for text in texts:
        element = _strip_err(copy.deepcopy(proto))
        if end is not None:
            end.addprevious(element)
        else:
            para._p.append(element)
        run = para.runs[-1]
        run.text = text
        made.append(run)
    return made


def set_paragraph(para, value):
    """The str / list branch of the three-valued protocol."""
    runs = para.runs
    if not runs:
        return False
    if isinstance(value, str):
        runs[0].text = value
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        return True
    for i, text in enumerate(value):
        if i < len(runs):
            runs[i].text = text
        else:
            append_runs(para, [text], proto_index=-1)
            runs = para.runs
    for r in para.runs[len(value):]:
        r._r.getparent().remove(r._r)
    return True


def apply_text_map(prs, mapping):
    """Apply {slide: {shape_id: {paragraph_index: value}}}.

    Returns a list of (slide_number, shape_id, problem) and never raises. A
    caller that wants a hard failure asserts the list is empty - which is
    what the version driver does.
    """
    problems = []
    for number, shape_map in mapping.items():
        slide = prs.slides[number - 1]
        by_id = shapes_by_id(slide)
        for shape_id, paragraphs in shape_map.items():
            shape = by_id.get(shape_id)
            if shape is None:
                problems.append((number, shape_id, "shape not found"))
                continue
            if not shape.has_text_frame:
                problems.append((number, shape_id, "shape has no text frame"))
                continue
            frame = shape.text_frame
            to_drop = []
            for index, value in paragraphs.items():
                if index >= len(frame.paragraphs):
                    problems.append((number, shape_id, "paragraph %d does not exist" % index))
                    continue
                if value is None:
                    to_drop.append(index)
                    continue
                if not set_paragraph(frame.paragraphs[index], value):
                    problems.append((number, shape_id, "paragraph %d has no runs" % index))
            for index in sorted(to_drop, reverse=True):
                element = frame.paragraphs[index]._p
                element.getparent().remove(element)
    return problems


# ------------------------------------------------------- deck-repair helpers
#
# Things a deck built by someone else needs done to it before generated
# content can live in it. Each one is a catalogue item: it cost the
# engagement a debugging session.

P = "http://schemas.openxmlformats.org/presentationml/2006/main"


class AnimationTargetError(Exception):
    """Raised rather than let a deleted shape leave a dangling p:spTgt.

    Catalogue item 8: a dangling `p:spTgt` makes automation-mode Open() fail
    outright - 'PowerPoint could not open the file' - rather than offer
    repair. There is no recovery downstream, so the deletion is refused at
    the point it is asked for.
    """


def anim_targets(slide):
    """Shape ids the slide's timing tree targets."""
    ids = set()
    for tgt in slide._element.iter(q("spTgt", P)):
        spid = tgt.get("spid")
        if spid and spid.isdigit():
            ids.add(int(spid))
    return ids


def _shape_and_descendant_ids(shape):
    """`shape`'s own id, plus every id nested inside it if it is a group.

    Deleting a group deletes its children with it, so the animation guard
    has to check the union, not just the group's own id: a group whose own
    id the timing tree never mentions can still contain the actual
    animation target as a child, and removing the group removes that child
    too, leaving the same dangling `p:spTgt` the guard exists to prevent.
    """
    ids = {shape.shape_id}
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for sub in shape.shapes:
            ids |= _shape_and_descendant_ids(sub)
    return ids


def delete_shape(slide, shape):
    """Remove a shape from the slide, refusing an animated one.

    Every deletion in this module goes through here, so the animation guard
    (catalogue item 8) applies uniformly: a shape the timing tree targets -
    or, for a group, any shape nested inside it - is never removed, because
    there is no repair for the dangling `p:spTgt` that would leave behind.
    """
    hit = _shape_and_descendant_ids(shape) & anim_targets(slide)
    if hit:
        raise AnimationTargetError(
            "shape %d is, or contains, an animation target (id(s) %s); "
            "deleting it makes the file unopenable under COM automation. "
            "Remove the animation first."
            % (shape.shape_id, sorted(hit)))
    shape._element.getparent().remove(shape._element)


def set_wrap(prs, targets):
    """Turn word wrap on. Returns the targets it could not find.

    Catalogue item 12: boxes inherited from a template may arrive with wrap
    off, so replacement copy longer than the original runs past the right
    edge with no warning anywhere.
    """
    missing = []
    for number, shape_id in targets:
        shape = shapes_by_id(prs.slides[number - 1]).get(shape_id)
        if shape is None or not shape.has_text_frame:
            missing.append((number, shape_id))
            continue
        shape.text_frame.word_wrap = True
    return missing


def widen(prs, targets):
    """Set width in design pixels and turn wrap on. Returns misses.

    The other half of catalogue item 12: a box whose width was auto-fitted to
    the original copy has to be widened explicitly, because wrap alone just
    makes the overflow vertical.
    """
    from .geometry import px
    missing = []
    for number, shape_id, width in targets:
        shape = shapes_by_id(prs.slides[number - 1]).get(shape_id)
        if shape is None or not shape.has_text_frame:
            missing.append((number, shape_id, width))
            continue
        shape.width = px(width)
        shape.text_frame.word_wrap = True
    return missing


def _contains_other_text(slide, candidate, exclude):
    """True when `candidate`'s bounding box also contains the centre of some
    other text-bearing shape on the slide, besides `exclude`.

    A pill sized to one label never overlaps a neighbour's centre; a card
    frame enclosing several labels does. `backing_shape` has to answer the
    geometry question honestly - it is a query, and the smallest empty
    auto-shape containing the centre is exactly what spec section 4 stage 4
    asks for - so the container/pill distinction is enforced here, at the
    point that actually destroys shapes, not by narrowing the query.
    """
    for shape in iter_shapes(slide.shapes):
        if shape.shape_id in (candidate.shape_id, exclude.shape_id):
            continue
        if not (shape.has_text_frame and shape.text_frame.text.strip()):
            continue
        cx = shape.left + shape.width // 2
        cy = shape.top + shape.height // 2
        if (candidate.left <= cx <= candidate.left + candidate.width
                and candidate.top <= cy <= candidate.top + candidate.height):
            return True
    return False


def delete_with_backing(slide, text_shape):
    """Delete a text box and the pill behind it, together.

    The second half of catalogue item 13. Deleting one half leaves an orphan
    rectangle on the slide, which nobody notices until the deck is presented.
    Both deletions go through `delete_shape`, so the animation guard applies
    to each.

    A backing that also encloses some other text shape's centre is a card,
    not a pill - refuse to take it. `backing_shape` still returns it; this
    is the point where deleting it would cost more than an orphan.
    """
    backing = backing_shape(slide, text_shape)
    if backing is not None and _contains_other_text(slide, backing, text_shape):
        backing = None      # a card enclosing several labels is not this label's pill
    delete_shape(slide, text_shape)
    if backing is not None:
        delete_shape(slide, backing)
    return backing


def backing_shape(slide, text_shape):
    """The auto-shape sitting behind a text box.

    Catalogue item 13: OOXML records no association between a text box and
    its pill background, so it is recovered geometrically - exact bounding-box
    equality first, then the smallest empty auto-shape containing the text
    box's centre. Deleting one half without the other leaves an orphan
    rectangle on the slide.
    """
    box = (text_shape.left, text_shape.top, text_shape.width, text_shape.height)
    cx = text_shape.left + text_shape.width // 2
    cy = text_shape.top + text_shape.height // 2
    candidates = []
    for shape in iter_shapes(slide.shapes):
        if shape.shape_id == text_shape.shape_id:
            continue
        if shape.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE:
            continue
        if shape.has_text_frame and shape.text_frame.text.strip():
            continue
        if (shape.left, shape.top, shape.width, shape.height) == box:
            return shape
        if (shape.left <= cx <= shape.left + shape.width
                and shape.top <= cy <= shape.top + shape.height):
            candidates.append(shape)
    if not candidates:
        return None
    return min(candidates, key=lambda s: s.width * s.height)


def ensure_notes_body(slide, prs):
    """The slide's notes text frame, creating the placeholder if needed.

    Catalogue item 11: a notes master without placeholders makes note
    insertion fail silently - the text goes into the file and never appears
    in the notes pane. The fix is to clone a BODY placeholder from a notes
    slide that has one.
    """
    notes = slide.notes_slide
    frame = notes.notes_text_frame
    if frame is not None:
        return frame
    donor = None
    for other in prs.slides:
        if other is slide or not other.has_notes_slide:
            continue
        for shape in other.notes_slide.shapes:
            if (shape.is_placeholder
                    and shape.placeholder_format.type == PP_PLACEHOLDER.BODY):
                donor = shape
                break
        if donor is not None:
            break
    if donor is None:
        raise LookupError(
            "no notes BODY placeholder anywhere in this deck to clone; add a "
            "note to one slide in PowerPoint first")
    notes.shapes._spTree.append(copy.deepcopy(donor._element))
    frame = notes.notes_text_frame
    frame.clear()
    return frame
