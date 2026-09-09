"""Shapes and text.

Every function takes its colours and its font as arguments: the core
library holds no palette. Geometry arguments are CSS pixels.
"""

import hashlib
import os

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Pt

from .geometry import px
# One definition of "a bare string is one regular run", shared with
# TextMetrics.lines so that measuring and writing take the same argument.
from .metrics import norm_runs as _norm


def _track(run_font, spc):
    """Letter-spacing, in points, written straight onto the run properties.

    python-pptx has no API for `a:rPr/@spc`, and the attribute is in
    hundredths of a point rather than points - so a caller reaching for the
    XML itself has two chances to get the unit wrong. This takes points, the
    same unit as every other size in the kit, and converts once.

    Tracking is not decoration for every brand: a design system whose display
    type is set at -0.04em and whose labels are set at +0.12em is a different
    system without it. It carries no colour and no brand knowledge - the
    amount is the caller's - which is why it belongs here rather than in a
    brand module. Pinned by
    tests/test_primitives.py::test_settext_writes_letter_spacing_in_hundredths_of_a_point.
    """
    if not spc:
        return
    run_font._rPr.set("spc", str(int(round(spc * 100))))


def _paint(sp, fill, line, lw, shadow):
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(lw)
    sp.shadow.inherit = shadow


def rrect(slide, x, y, w, h, fill=None, line=None, lw=1.0, rad=8, shadow=False):
    rounded = bool(rad) and rad > 1
    sp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        px(x), px(y), px(w), px(h))
    if rounded:
        try:
            sp.adjustments[0] = min(0.5, rad / float(min(w, h)))
        except Exception:
            pass
    _paint(sp, fill, line, lw, shadow)
    return sp


def oval(slide, x, y, w, h, fill=None, line=None, lw=1.0):
    sp = slide.shapes.add_shape(MSO_SHAPE.OVAL, px(x), px(y), px(w), px(h))
    _paint(sp, fill, line, lw, False)
    return sp


def tri(slide, cx, cy, direction, color, w=11, h=9):
    sp = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
                                px(cx - w / 2.0), px(cy - h / 2.0), px(w), px(h))
    sp.rotation = {"up": 0, "right": 90, "down": 180, "left": 270}[direction]
    _paint(sp, color, None, 1.0, False)
    return sp


def rule(slide, x, y, w, h, color):
    return rrect(slide, x, y, w, max(0.75, h), fill=color, rad=0)


def _no_autofit(tf):
    """Turn PowerPoint's autofit off on this text frame, in the XML.

    python-pptx's add_textbox() emits <a:bodyPr><a:spAutoFit/></a:bodyPr>, so
    every text box arrives with "resize shape to fit text" ON unless something
    clears it. That silently defeats measure-then-place: the builder positions
    what it measured, and PowerPoint then resizes the box the moment anyone
    edits the text.

    Done at the XML level on purpose. Reaching for MSO_AUTO_SIZE would import
    the very enum the doctrine forbids, and a source-level grep cannot see
    autofit that arrives by default anyway - so the real check lives in
    scripts/assert_native.py, against the built package.
    """
    bodyPr = tf._txBody.bodyPr
    for tag in ("a:spAutoFit", "a:normAutofit"):
        for e in bodyPr.findall(qn(tag)):
            bodyPr.remove(e)
    if bodyPr.find(qn("a:noAutofit")) is None:
        bodyPr.append(parse_xml("<a:noAutofit %s/>" % nsdecls("a")))


def settext(shape_or_tf, runs, size, color, font, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.MIDDLE, ls=1.05, wrap=True, upper=False, ml=0, mr=0,
            spc=0):
    """Write the FIRST paragraph. Later paragraphs are add_para's job.

    `spc` is letter-spacing in points, positive or negative, applied to
    every run. Keyword-only in practice: it sits after `mr` so no existing
    positional call moves.
    """
    tf = shape_or_tf.text_frame if hasattr(shape_or_tf, "text_frame") else shape_or_tf
    tf.word_wrap = wrap
    _no_autofit(tf)
    tf.margin_left, tf.margin_right = px(ml), px(mr)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    # Clear runs, line breaks and fields alike. Runs alone is not enough: a
    # paragraph from a deck someone else authored routinely carries <a:br>
    # and <a:fld>, and leaving them makes settext's "rewrite the first
    # paragraph" contract quietly false.
    for tag in ("a:r", "a:br", "a:fld"):
        for el in p._p.findall(qn(tag)):
            p._p.remove(el)
    p.alignment = align
    p.line_spacing = ls
    for run in _norm(runs):
        text, bold = run[0], run[1]
        rc = run[2] if len(run) > 2 else color
        r = p.add_run()
        r.text = text.upper() if upper else text
        f = r.font
        f.size, f.name, f.bold = Pt(size), font, bold
        f.color.rgb = rc
        _track(f, spc)
    return tf


def textbox(slide, x, y, w, h, runs, size, color, font, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.TOP, ls=1.05, wrap=True, upper=False, spc=0):
    """Parameter order matches settext deliberately: the two used to swap
    `wrap` and `upper`, so a positional call transposed them silently."""
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    settext(tb, runs, size, color, font, align, anchor, ls, wrap, upper, spc=spc)
    return tb


def add_para(tf, runs, size, color, font, align=PP_ALIGN.LEFT, ls=1.05,
             space_before=0, upper=False, spc=0):
    p = tf.add_paragraph()
    p.alignment, p.line_spacing = align, ls
    if space_before:
        p.space_before = Pt(space_before)
    for run in _norm(runs):
        text, bold = run[0], run[1]
        rc = run[2] if len(run) > 2 else color
        r = p.add_run()
        r.text = text.upper() if upper else text
        f = r.font
        f.size, f.name, f.bold = Pt(size), font, bold
        f.color.rgb = rc
        _track(f, spc)
    return p


def chip(slide, x, y, w, h, text, size, fg, fill, border, font, lw=0.75):
    """A pill-shaped label. Refuses text the same colour as its ground: a
    chip inverted for a dark band whose caller forgot to invert the text
    exports as an empty pill, and no count-based check can see it."""
    if fill is not None and fg == fill:
        raise ValueError("chip %r: text colour equals fill colour, so the label is "
                         "invisible" % (text,))
    sp = rrect(slide, x, y, w, h, fill=fill, line=border, lw=lw, rad=h / 2.0)
    settext(sp, text, size, fg, font, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, wrap=False)
    return sp


def numcircle(slide, x, y, sz, n, fill, fg, font, size):
    sp = oval(slide, x, y, sz, sz, fill=fill)
    settext(sp, str(n), size, fg, font, PP_ALIGN.CENTER,
            MSO_ANCHOR.MIDDLE, ls=1.0, wrap=False)
    return sp


def icon(slide, icon_dir, name, colorkey, x, y, sz):
    """Place a rasterised icon.

    A missing file raises. The engagement's version returned silently, so a
    typo'd name rendered nothing and the build still reported success.
    """
    path = os.path.join(str(icon_dir), "ic_%s_%s.png" % (name, colorkey))
    if not os.path.exists(path):
        raise FileNotFoundError("no icon %r in %s" % (name, icon_dir))
    return slide.shapes.add_picture(path, px(x), px(y), px(sz), px(sz))


def _hexname(c):
    return "%02X%02X%02X" % (c[0], c[1], c[2])


def ensure_balls(ball_dir, colors, ring, body, fractions=(0, 25, 50, 75, 100), size=18, ss=4):
    """Pre-render conic fills to transparent PNGs.

    DrawingML has no angular gradient, so this is one of the only two
    places raster is allowed. The cache keys on a hash of the rendering
    parameters rather than on file existence: the engagement's version
    regenerated nothing when a colour changed.

    Every colour - the filled arc, the ring, the unfilled body - is an
    argument. The core decides none of them.
    """
    from PIL import Image, ImageDraw

    ball_dir = str(ball_dir)
    os.makedirs(ball_dir, exist_ok=True)
    sz = size * ss
    for color in colors:
        rgba = (color[0], color[1], color[2], 255)
        for frac in fractions:
            path = os.path.join(ball_dir, "ball_%d_%s.png" % (frac, _hexname(color)))
            stamp = hashlib.sha1(
                repr((rgba, tuple(ring), tuple(body), frac, size, ss)).encode()).hexdigest()[:12]
            stamp_path = path + ".stamp"
            if os.path.exists(path) and os.path.exists(stamp_path):
                with open(stamp_path) as fh:
                    if fh.read() == stamp:
                        continue
            img = Image.new("RGBA", (sz, sz), (0, 0, 0, 0))
            d = ImageDraw.Draw(img)
            bb = [ss, ss, sz - ss, sz - ss]
            d.ellipse(bb, fill=tuple(body), outline=tuple(ring), width=int(1.5 * ss))
            if frac == 100:
                d.ellipse(bb, fill=rgba, outline=tuple(ring), width=int(1.5 * ss))
            elif frac > 0:
                d.pieslice(bb, -90, -90 + 360 * frac / 100.0, fill=rgba)
                d.ellipse(bb, outline=tuple(ring), width=int(1.5 * ss))
            img.save(path)
            with open(stamp_path, "w") as fh:
                fh.write(stamp)


def ball(slide, ball_dir, x, y, sz, frac, color):
    path = os.path.join(str(ball_dir), "ball_%d_%s.png" % (frac, _hexname(color)))
    if not os.path.exists(path):
        raise FileNotFoundError("no ball asset %s; call ensure_balls first" % path)
    return slide.shapes.add_picture(path, px(x), px(y), px(sz), px(sz))


def lockup(slide, primary, secondary, right_x, y, divider_color,
           primary_h=24, secondary_h=20, gap=14):
    """Two marks separated by a hairline, right edge at right_x."""
    from PIL import Image

    def aspect(path, h):
        with Image.open(str(path)) as im:
            return h * im.width / float(im.height)

    pw = aspect(primary, primary_h)
    sw = aspect(secondary, secondary_h)
    total = pw + gap + 1 + gap + sw
    x = right_x - total
    slide.shapes.add_picture(str(primary), px(x), px(y - primary_h / 2.0),
                             px(pw), px(primary_h))
    sx = x + pw + gap
    rule(slide, sx, y - 11, 1, 22, divider_color)
    slide.shapes.add_picture(str(secondary), px(sx + 1 + gap),
                             px(y - secondary_h / 2.0), px(sw), px(secondary_h))


def footer(slide, left_text, page, font, color, rule_color, size,
           width=1184, left=48, y=690, page_w=60, gap=12):
    """A rule, left text and a page number, all inside `width`."""
    rule(slide, left, y, width, 1, rule_color)
    textbox(slide, left, y + 6, max(0, width - page_w - gap), 16, left_text,
            size, color, font, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE, ls=1.0)
    textbox(slide, left + width - page_w, y + 6, page_w, 16, str(page), size,
            color, font, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE, ls=1.0)
