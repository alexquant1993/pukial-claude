"""Speaker notes: the model, the PPTX writer and the teleprompter emitter.

The note model carries only what the author writes - position, title, spoken
text, a time budget and an optional section label. Whether a slide is hidden
and what number is printed on it are properties of the deck, read from the
presentation when the teleprompter is emitted (see `numbering`, Task 2),
never typed twice.

The PPTX writer goes through `textedit.ensure_notes_body` for catalogue
item 11 (a notes master without placeholders makes note insertion fail
silently). It does not reimplement the placeholder clone.

Notes are opt-in. A deck carries no notes and no teleprompter unless its
author asks for them, because most decks are read, not presented, and a
notes pane full of generated text is noise to the person who edits the file
by hand. `emit_notes` is the one entry point a builder routes through, and
its `mode` is one of `NOTES_MODES`:

  "none"          write nothing; the note list is not even checked for
                  coverage, so a builder can keep its notes_data.py and
                  still ship a deck without them
  "pptx"          the notes pane in the PPTX (`write_notes`)
  "teleprompter"  the standalone HTML page (`teleprompter_html`)
  "both"          generated once, into both (design decision D5)

`write_notes` and `teleprompter_html` stay public for callers that want one
output without the file handling.
"""

import html
import os
import re
from dataclasses import dataclass
from pathlib import Path

from pptx.util import Pt

from deck_kit.pagenums import is_hidden, visible_numbers
from deck_kit.textedit import ensure_notes_body

ASIDE_PREFIX = ">"

NOTES_MODES = ("none", "pptx", "teleprompter", "both")


@dataclass(frozen=True)
class Note:
    position: int          # 1-based file position of the slide
    title: str             # short label shown in the teleprompter
    text: str              # paragraphs separated by blank lines
    minutes: int = 0       # spoken-time budget; 0 for separators
    section: str = ""      # optional label, e.g. "Annex"


def paragraphs(text):
    """Blank-line separated paragraphs, stripped, with empties dropped."""
    return [p.strip() for p in text.strip().split("\n\n") if p.strip()]


def is_aside(paragraph):
    """A paragraph spoken only if asked. Marked by a leading '>'."""
    return paragraph.startswith(ASIDE_PREFIX)


def _check_coverage(prs, notes):
    count = len(prs.slides)
    positions = [n.position for n in notes]
    outside = sorted(p for p in positions if not 1 <= p <= count)
    if outside:
        raise ValueError("note position(s) %s outside 1..%d" % (outside, count))
    seen, duplicates = set(), []
    for p in positions:
        if p in seen:
            duplicates.append(p)
        seen.add(p)
    if duplicates:
        raise ValueError("duplicate note position(s) %s" % sorted(set(duplicates)))
    missing = sorted(set(range(1, count + 1)) - seen)
    if missing:
        raise ValueError("missing note(s) for slide position(s) %s" % missing)


def write_notes(prs, notes, *, font, size_pt, gap_pt):
    """Write every note into its slide's notes body. Returns the positions written.

    `font`, `size_pt` and `gap_pt` are required: the core library carries no
    measurement defaults, and a note size belongs to the brand's type scale.

    Every slide must have exactly one note: a teleprompter with a gap is a
    teleprompter the presenter discovers is missing on stage.
    """
    _check_coverage(prs, notes)
    written = []
    for note in sorted(notes, key=lambda n: n.position):
        slide = prs.slides[note.position - 1]
        frame = ensure_notes_body(slide, prs)
        frame.clear()
        for k, ptxt in enumerate(paragraphs(note.text)):
            para = frame.paragraphs[0] if k == 0 else frame.add_paragraph()
            run = para.add_run()
            run.text = ptxt
            run.font.size = Pt(size_pt)
            run.font.name = font
            if k:
                para.space_before = Pt(gap_pt)
        written.append(note.position)
    return written


def read_notes(prs):
    """{position: paragraphs} for every slide that carries a notes body."""
    out = {}
    for position, slide in enumerate(prs.slides, 1):
        if not slide.has_notes_slide:
            continue
        frame = slide.notes_slide.notes_text_frame
        if frame is None:
            continue
        paras = [p.text for p in frame.paragraphs if p.text.strip()]
        if paras:
            out[position] = paras
    return out


_REPO_TEMPLATES = Path(__file__).resolve().parents[2] / "templates"


def templates_dir():
    """Resolved per call: DECKWRIGHT_TEMPLATES if set, else the source checkout.

    The same rule as `brand.brands_dir()`; a wheel install has no
    `templates/` two levels above this file.
    """
    env = os.environ.get("DECKWRIGHT_TEMPLATES")
    return Path(env) if env else _REPO_TEMPLATES


def numbering(prs, skip=(1,)):
    """{position: (printed_number or None, hidden)} straight from the deck."""
    numbers = visible_numbers(prs, skip)
    return {
        position: (numbers[position], is_hidden(slide))
        for position, slide in enumerate(prs.slides, 1)
    }


def _label(note, printed, hidden):
    label = "Slide %d" % note.position
    if printed is not None:
        label += " [%d]" % printed
    if hidden:
        label += " · hidden"
    return label


PLACEHOLDERS = ("title", "total_minutes", "nav", "sections")
_PLACEHOLDER = re.compile(r"\{\{(\w+)\}\}")


def teleprompter_html(notes, numbers, *, title, template=None):
    """Fill templates/teleprompter.html. `numbers` is what `numbering(prs)` returns.

    The skeleton is checked for its placeholder set BEFORE substitution and
    filled in a single regex pass: checking the filled output for a stray
    "{{" (the earlier approach) both raises on legitimate note/title/section
    text that happens to contain "{{" (html.escape does not touch braces)
    and, worse, substitutes into that content when a later .replace pass
    matches text a previous pass just inserted.
    """
    positions = sorted(n.position for n in notes)
    if positions != sorted(numbers):
        raise ValueError("notes cover positions %s but numbering covers %s"
                         % (positions, sorted(numbers)))
    path = Path(template) if template else templates_dir() / "teleprompter.html"
    skeleton = path.read_text(encoding="utf-8")
    found = _PLACEHOLDER.findall(skeleton)
    unknown = sorted(set(found) - set(PLACEHOLDERS))
    missing = sorted(set(PLACEHOLDERS) - set(found))
    if unknown or missing:
        raise ValueError("template %s: unknown placeholder(s) %s, missing %s"
                         % (path, unknown, missing))

    nav, sections = [], []
    for note in sorted(notes, key=lambda n: n.position):
        printed, hidden = numbers[note.position]
        cls = ' class="hidden"' if hidden else ""
        nav.append('<a href="#s%d"%s>%d</a>' % (note.position, cls, note.position))
        head = ['<span class="n">%s</span>' % html.escape(_label(note, printed, hidden))]
        if note.minutes:
            head.append('<span class="min">%d min</span>' % note.minutes)
        if note.section:
            head.append('<span class="section-label">%s</span>' % html.escape(note.section))
        head.append("<h2>%s</h2>" % html.escape(note.title))
        body = []
        for ptxt in paragraphs(note.text):
            if is_aside(ptxt):
                body.append('<p class="aside">%s</p>'
                            % html.escape(ptxt[len(ASIDE_PREFIX):].strip()))
            else:
                body.append("<p>%s</p>" % html.escape(ptxt))
        sections.append('<section id="s%d"%s>\n<div class="hd">%s</div>\n%s\n</section>'
                        % (note.position, cls, "\n".join(head), "\n".join(body)))

    values = {
        "title": html.escape(title),
        "total_minutes": str(sum(n.minutes for n in notes)),
        "nav": "\n".join(nav),
        "sections": "\n".join(sections),
    }
    return _PLACEHOLDER.sub(lambda m: values[m.group(1)], skeleton)


def emit_notes(prs, notes, mode, *, font, size_pt, gap_pt, title, teleprompter=None,
               skip=(1,)):
    """Emit the notes the author opted into. Returns {"pptx": ..., "teleprompter": ...}.

    `mode` is one of NOTES_MODES. "pptx" fills the notes pane and reports
    the positions written under "pptx"; "teleprompter" fills the HTML page,
    writes it to `teleprompter` (parents created, utf-8) and reports that
    path under "teleprompter"; "both" does both; "none" does nothing and
    reports None for each - without validating coverage, because a note list
    nobody asked to emit cannot fail a build. A key that was not asked for
    is None, so a builder's check() can tell "not requested" from "written".

    `font`, `size_pt` and `gap_pt` are required even for the modes that do
    not use them: the caller supplies the brand's type scale once, and the
    mode - a runtime choice - decides whether it is spent. `skip` is the
    page-number skip list `numbering` reads with, so the teleprompter agrees
    with the footer.

    An unknown mode is a ValueError naming the modes rather than a silent
    "none": a typo in a CLI choice must not ship a deck without its notes.
    """
    if mode not in NOTES_MODES:
        raise ValueError("unknown notes mode %r; choose one of %s"
                         % (mode, ", ".join(NOTES_MODES)))
    emitted = {"pptx": None, "teleprompter": None}
    if mode == "none":
        return emitted
    wants_teleprompter = mode in ("teleprompter", "both")
    if wants_teleprompter and teleprompter is None:
        raise ValueError("notes mode %r writes a teleprompter and needs its path" % mode)
    if mode in ("pptx", "both"):
        emitted["pptx"] = write_notes(prs, notes, font=font, size_pt=size_pt, gap_pt=gap_pt)
    if wants_teleprompter:
        html_text = teleprompter_html(notes, numbering(prs, skip=skip), title=title)
        path = Path(teleprompter)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(html_text, encoding="utf-8")
        emitted["teleprompter"] = path
    return emitted
