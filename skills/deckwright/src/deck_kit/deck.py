"""Deck lifecycle. This module owns slide-list mutation; nothing else in
the library removes or reorders slides.

A builder opens the brand's template and purges its slides, so generated
slides inherit the master, the layout and the background without the
builder drawing any of it.
"""

from pptx import Presentation

RT_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"


def open_deck(template_path, spec):
    """Open the template and drop every slide, keeping masters and layouts.

    The brand's layout is resolved here rather than at the first add_slide,
    so a template/brand mismatch fails before any content is built.
    """
    prs = Presentation(str(template_path))
    _layout(prs, spec)
    _drop(prs, list(prs.slides._sldIdLst))
    return prs


def _drop(prs, sld_ids):
    lst = prs.slides._sldIdLst
    for sldId in sld_ids:
        prs.part.drop_rel(sldId.get(RT_ID))
        lst.remove(sldId)


def _layout(prs, spec):
    for master in prs.slide_masters:
        if master.name != spec.master_name:
            continue
        for layout in master.slide_layouts:
            if layout.name == spec.layout_name:
                return layout
    raise LookupError(
        "layout %r not found under master %r; run inspect_template.py against "
        "this deck to see what it does have" % (spec.layout_name, spec.master_name))


def add_slide(prs, spec, keep_placeholders=()):
    """New slide on the brand's layout, with inherited placeholders removed.

    `keep_placeholders` names placeholder types (members of
    `pptx.enum.shapes.PP_PLACEHOLDER`) to leave in place rather than purge -
    Task 7 needs one real SLIDE_NUMBER placeholder to survive somewhere in
    the repository, and this function's unconditional purge is why nothing
    in it has one. Defaults to `()` so existing callers are unaffected.

    python-pptx's own `Slides.add_slide()` clones the layout's placeholders
    onto the new slide EXCEPT the "latent" ones - date, footer and slide
    number - because those are normally left to render from the layout
    untouched. SLIDE_NUMBER is exactly the placeholder Task 7 needs kept, so
    a `keep_placeholders` request for it is cloned here explicitly, from the
    layout's own placeholder, before the purge loop runs; otherwise the loop
    below would never see it to keep.

    Catalogue item 9: python-pptx picks the next slideN part name naively,
    so after any deletion it reuses a name still occupied by another slide
    and one overwrites the other inside the package. The catalogue names
    `add_slide()`, not `duplicate_slide`, so the rename happens here too.
    """
    layout = _layout(prs, spec)
    slide = prs.slides.add_slide(layout)
    _rename_above_max(prs, slide)
    present = {ph.placeholder_format.type for ph in slide.placeholders}
    for lph in layout.placeholders:
        wanted = lph.placeholder_format.type
        if wanted in keep_placeholders and wanted not in present:
            slide.shapes.clone_placeholder(lph)
            present.add(wanted)
    for ph in list(slide.placeholders):
        if ph.placeholder_format.type in keep_placeholders:
            continue
        ph._element.getparent().remove(ph._element)
    return slide


def delete_positions(prs, positions):
    """Remove slides by 1-based position.

    Positions are validated and de-duplicated. `positions` being 1-based is
    the kind of thing a caller forgets, and an unvalidated `ids[pos - 1]`
    turns a 0 into "delete the last slide" and a repeat into a KeyError
    halfway through the mutation.
    """
    ids = list(prs.slides._sldIdLst)
    unique = sorted(set(positions))
    bad = [p for p in unique if not 1 <= p <= len(ids)]
    if bad:
        raise IndexError(
            "slide position(s) %s outside 1..%d" % (bad, len(ids)))
    _drop(prs, [ids[p - 1] for p in unique])


import copy

from pptx.opc.packuri import PackURI

RT_SLIDE_LAYOUT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout"
RT_NOTES = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/notesSlide"
_SHAPE_TAGS = ("}sp", "}pic", "}graphicFrame", "}grpSp", "}cxnSp")


def _rename_above_max(prs, slide):
    """Give `slide` a part name above every existing slideN.

    Catalogue item 9. python-pptx picks the next slideN number naively, so
    after any deletion it reuses a name that is still occupied and one slide
    overwrites another inside the package. The catalogue names `add_slide()`,
    so both callers go through here - fixing only `duplicate_slide` leaves the
    documented failure live in the function every builder calls.
    """
    used = [s.part.partname for s in prs.slides if s.part is not slide.part]
    numbers = [int(str(u).rsplit("slide", 1)[1].split(".")[0]) for u in used]
    slide.part.partname = PackURI(
        "/ppt/slides/slide%d.xml" % ((max(numbers) + 1) if numbers else 1))
    return slide


def duplicate_slide(prs, index0):
    """Duplicate slide `index0` (0-based) and place the copy right after it.

    add_slide() picks the next slideN part name naively, so after any deletion
    it reuses a name that is still occupied and one slide overwrites another
    inside the package. The maximum over the existing names is computed and
    the part name set explicitly.
    """
    src = prs.slides[index0]
    dest = prs.slides.add_slide(src.slide_layout)
    _rename_above_max(prs, dest)

    tree = dest.shapes._spTree
    for element in list(tree):
        if any(element.tag.endswith(t) for t in _SHAPE_TAGS):
            tree.remove(element)

    rmap = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype in (RT_SLIDE_LAYOUT, RT_NOTES):
            continue
        if rel.is_external:
            rmap[rid] = dest.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            rmap[rid] = dest.part.relate_to(rel.target_part, rel.reltype)

    for shape in src.shapes:
        element = copy.deepcopy(shape._element)
        for node in element.iter():
            for attr in list(node.attrib):
                if attr.startswith(RT_ID.rsplit("}", 1)[0] + "}"):
                    if node.get(attr) in rmap:
                        node.set(attr, rmap[node.get(attr)])
        tree.append(element)

    lst = prs.slides._sldIdLst
    ids = list(lst)
    moved = ids[-1]
    lst.remove(moved)
    lst.insert(index0 + 1, moved)
    return dest


def fingerprint(prs, positions):
    """The joined text of each 1-based position.

    Walks through groups. Everything else in the phase does, and on a real
    author's deck the title is usually inside one - a non-recursive walk
    returns an empty fingerprint and `assert_fingerprint` then fails on a deck
    that is entirely correct. A guard that cries wolf is a guard people stop
    passing arguments to.
    """
    from .textedit import iter_shapes
    out = {}
    for pos in positions:
        slide = prs.slides[pos - 1]
        out[pos] = " | ".join(sh.text_frame.text.strip()
                              for sh in iter_shapes(slide.shapes)
                              if sh.has_text_frame)
    return out


def assert_fingerprint(prs, slide_count, expected):
    """Assert this is the file the build was written against.

    Not a checksum: the author edits the deck between versions, so the file
    changes legitimately. What must not change is the slide count and what
    sits at each position the build is about to touch. `expected` maps a
    1-based position to a substring that must appear in that slide's text.
    """
    if len(prs.slides) != slide_count:
        raise AssertionError(
            "deck has %d slides, expected exactly %d - the author added or "
            "removed slides; re-read the deck before rebuilding"
            % (len(prs.slides), slide_count))
    got = fingerprint(prs, sorted(expected))
    for pos, needle in expected.items():
        # Matched against the untruncated join; only the message truncates, so
        # a needle beyond an arbitrary limit cannot fail for the wrong reason.
        if needle not in got[pos]:
            raise AssertionError(
                "position %d does not contain %r; it holds %r - the author moved "
                "slides and this build would replace the wrong one"
                % (pos, needle, got[pos][:300]))
